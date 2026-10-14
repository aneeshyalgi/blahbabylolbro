"""Content Lineage AI Agent: the forward business ("content") lineage of one field.

Technical lineage answers how the code moves data from column to column. Content lineage answers what happens to the
content of a field — which fields it flows into next and further downstream, up to the end of the calculation, and in
which business role:

- the next fields: every field whose business rule reads the selected field, with the cases that use it as a value,
  as a condition that decides whether a case applies, or as the key of a parameter table;
- the fields further downstream, reached through the next fields, up to the end fields of the calculation, and the
  routes the content takes;
- for every record, whether and how each downstream value depends on the selected field: attributed to the case that
  produced it and confirmed by an impact test that re-runs the code with other values of the field;
- the business meaning of every field and rule, and how the selected field reaches each downstream field.

Every run follows one field of one execution (the code of a cluster applied to its dataset) in a fixed pipeline:

1. trace     - static data-flow analysis of the code (shared with the technical lineage agent) and a statement-by-
               statement replay of the execution; every value of every column is attributed to the case that wrote it;
2. follow    - the fields the selected field's content flows into, computed on the column versions of the static
               analysis (a write nothing uses carries nothing); the business rules and cases of every downstream field;
               per record, the case behind every downstream value and whether it uses the selected field; an impact
               test re-runs the code with other values of the field and compares every value of the result;
3. describe  - the LLM names and defines the business terms, states every rule and case in business language and
               explains how the selected field reaches each downstream field (closed sets of terms and case ids,
               numbers checked against the constants of the code, one correction);
4. translate - every AI text is available in English and German.

Runs execute in a background thread and are polled by the browser; finished runs are stored as JSON and can be
exported as an Excel workbook.
"""
from __future__ import annotations

import ast
import copy
import json
import math
import os
import re
import threading
import time
import traceback
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

import pandas as pd
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

import agent_runtime as runtime
import database as db
import lineage_agent as tla
import lineage_analysis as la

router = APIRouter(prefix="/api/content-lineage-agent", tags=["content-lineage-agent"])

RUNS_DIR = runtime.RESULTS_DIR / "content_lineage_agent"
RESULT_MODE = "forward"  # results of earlier versions (backward from reported figures) carry no mode
REPLAY_ROW_LIMIT = 20000
RECORD_ROWS = 400  # records whose forward trace (case, values, what-if) is kept
MAX_ROUTES = 30
ROUTE_CANDIDATES = 400
MAX_CATEGORIES = 40  # a field with at most this many distinct values is shown value by value
MAX_OUTCOMES = 6
MAX_ROTATIONS = 11  # impact test: other values of the field tried per record
IMPACT_BUDGET_SECONDS = 90
PREVIEW_ROWS = 300  # setup preview: quick impact test of every field on the first records
PREVIEW_BUDGET_SECONDS = 8
RULE_BATCH = 8
WORKERS = 3
MAX_ACTIVE_RUNS = 2
TRANSLATE_BATCH = 40
DEFAULT_OPENAI_MODEL = "gpt-4.1"  # CONTENT_LINEAGE_AGENT_MODEL overrides it
LANGUAGE_NAMES = {"en": "English", "de": "German"}
TEXT_LANGUAGES = ("en", "de")
_ID = re.compile(r"^[0-9a-f-]{8,64}$")
_MISSING = object()


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _configured_model() -> str:
    return os.environ.get("CONTENT_LINEAGE_AGENT_MODEL", "").strip() or (
        DEFAULT_OPENAI_MODEL if os.environ.get("OPENAI_API_KEY") else os.environ.get("AZURE_OPENAI_DEPLOYMENT_NAME", "")
    )


def _plain(text: Any) -> str:
    """A formula of the static analysis with ⟦column⟧ markers written as [column] (for the language model)."""
    return re.sub(r"⟦([^⟦⟧]+)⟧", r"[\1]", str(text or ""))


def _number(value: Any) -> Optional[float]:
    if la._is_null(value) or isinstance(value, bool):
        return None
    if la._is_number(value):
        number = float(value)
        return number if math.isfinite(number) else None
    return None


def _key_text(value: Any) -> str:
    item = la.json_value(value)
    if item is None:
        return "∅"
    if isinstance(item, float) and item.is_integer():
        return str(int(item))
    return str(item)


def _round(value: Optional[float]) -> Optional[float]:
    return None if value is None else round(value, 6)


# --------------------------------------------------------------------------------------------------
# Job state
# --------------------------------------------------------------------------------------------------

class _Cancelled(Exception):
    pass


class _Job:
    """Run state shared with the polling browser; the interface matches what the technical agent's LLM wrapper uses."""

    def __init__(self, state: Dict[str, Any]):
        self.state = state
        self.lock = threading.Lock()
        self.cancel = threading.Event()
        self.started = time.monotonic()
        self.last_save = 0.0

    def log(self, level: str, code: str, message: str, /, **params: Any) -> None:
        # Positional-only: the parameters of an entry may themselves be named code, message or level.
        with self.lock:
            self.state["log"].append({
                "t": round(time.monotonic() - self.started, 2), "level": level, "code": code, "params": params, "message": message,
            })

    def stage(self, stage: str) -> None:
        with self.lock:
            self.state["stage"] = stage

    def progress(self, **changes: Any) -> None:
        with self.lock:
            self.state["progress"].update(changes)

    def increment(self, key: str, by: int = 1) -> None:
        with self.lock:
            self.state["progress"][key] = self.state["progress"].get(key, 0) + by

    def count_call(self, prompt: int, completion: int, tools: int = 0) -> None:
        with self.lock:
            usage = self.state["usage"]
            usage["llm_calls"] += 1
            usage["prompt_tokens"] += prompt
            usage["completion_tokens"] += completion
            usage["tool_calls"] += tools

    def set_model(self, model: str) -> None:
        with self.lock:
            self.state["model"] = model

    def check(self) -> None:
        if self.cancel.is_set():
            raise _Cancelled()

    def snapshot(self, lite: bool = False) -> Dict[str, Any]:
        with self.lock:
            if lite:
                return copy.deepcopy({key: value for key, value in self.state.items() if key != "result"})
            return copy.deepcopy(self.state)

    def save(self, force: bool = False) -> None:
        now = time.monotonic()
        if not force and now - self.last_save < 2.0:
            return
        self.last_save = now
        _write_run(self.snapshot())


_JOBS: Dict[str, _Job] = {}
_JOBS_LOCK = threading.Lock()


def _run_path(run_id: str) -> Path:
    if not _ID.match(run_id or ""):
        raise HTTPException(404, "Run not found")
    return RUNS_DIR / f"{run_id}.json"


def _write_run(state: Dict[str, Any]) -> None:
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    path = _run_path(state["id"])
    temporary = path.with_name(f"{path.stem}.{uuid.uuid4().hex[:8]}.tmp")
    temporary.write_text(json.dumps(state, ensure_ascii=False, default=str), encoding="utf-8")
    for attempt in range(20):
        try:
            os.replace(temporary, path)
            return
        except PermissionError:
            if attempt == 19:
                temporary.unlink(missing_ok=True)
                raise
            time.sleep(0.05)


def _read_run(run_id: str) -> Dict[str, Any]:
    path = _run_path(run_id)
    if not path.exists():
        raise HTTPException(404, "Run not found")
    for attempt in range(10):
        try:
            state = json.loads(path.read_text(encoding="utf-8"))
            break
        except PermissionError:
            if attempt == 9:
                raise
            time.sleep(0.05)
    if state.get("status") == "running" and run_id not in _JOBS:
        state["status"] = "interrupted"  # the server restarted while it ran
    return state


class _LLM(tla._LLM):
    """The technical agent's JSON-schema client (retries, model fallback, usage counts) with this agent's model."""

    def __init__(self, job: _Job):
        super().__init__(job)  # type: ignore[arg-type]
        override = os.environ.get("CONTENT_LINEAGE_AGENT_MODEL", "").strip()
        if override:
            _, configured = runtime._llm_client()
            self.model = override
            self.fallback = configured if configured and configured != override else None
        self.seed = 20261008


def _strings() -> Dict[str, Any]:
    return {"type": "array", "items": {"type": "string"}}


def _enum(values: List[str]) -> Dict[str, Any]:
    return {"type": "string", "enum": values} if values else {"type": "string"}


# --------------------------------------------------------------------------------------------------
# Context
# --------------------------------------------------------------------------------------------------

class _Context:
    """Everything the run knows about the execution and its content."""

    def __init__(self) -> None:
        self.language = "en"
        self.stored: Dict[str, Any] = {}
        self.code = ""
        self.lines: List[str] = []
        self.prepared = ""
        self.frame: Any = None
        self.make_globals: Any = None
        self.analysis: Dict[str, Any] = {}
        self.nodes: List[Dict[str, Any]] = []
        self.columns: Dict[str, Dict[str, Any]] = {}
        self.steps_by_line: Dict[int, Dict[str, Any]] = {}
        self.lookups: Dict[str, Dict[str, Any]] = {}
        self.final: Any = None
        self.statements: List[Dict[str, Any]] = []
        # column -> row label -> statement lines that changed the cell, in order
        self.history: Dict[str, Dict[Any, List[int]]] = {}
        # (statement line, column) -> row label -> lookup key the statement mapped
        self.keys: Dict[Tuple[int, str], Dict[Any, Any]] = {}
        # (statement line, column) -> lookup key -> rows the table did not cover and that stayed empty
        self.unmapped: Dict[Tuple[int, str], Dict[str, int]] = {}
        # column -> row label -> [{line, before, after, inputs}] for the first RECORD_ROWS records
        self.details: Dict[str, Dict[Any, List[Dict[str, Any]]]] = {}
        self.record_labels: List[Any] = []
        # The selected field and what its content flows into (see forward_structure).
        self.field = ""
        self.structure: Dict[str, Any] = {}
        self.reached: List[str] = []  # downstream fields, in the order of the calculation
        self.context_fields: List[str] = []  # other inputs of the downstream rules
        self.scope: List[str] = []  # selected field, downstream fields, other inputs
        self.derived: Set[str] = set()  # downstream fields with a business rule
        self.rules: List[Dict[str, Any]] = []
        self.rule_by_column: Dict[str, Dict[str, Any]] = {}
        self.cases: Dict[str, Dict[str, Any]] = {}
        self.case_by_write: Dict[Tuple[int, str], Dict[str, Any]] = {}
        self.identifier: Optional[str] = None
        self.filters: List[Dict[str, Any]] = []
        self.replay: Dict[str, Any] = {}
        self.dead: Set[int] = set()
        # Impact test: per downstream field, the records whose value changed and the first what-if per record.
        self.impact: Dict[str, Any] = {}
        self.changed: Dict[str, Set[Any]] = {}
        self.whatif: Dict[Tuple[str, Any], Dict[str, Any]] = {}
        self.x_values: Dict[Any, Any] = {}  # the selected field's value per record (its last state in the calculation)

    def code_span(self, line: int, end: int) -> str:
        return "\n".join(self.lines[line - 1:end])


# --------------------------------------------------------------------------------------------------
# 1. Trace: static analysis and replay with per-cell attribution
# --------------------------------------------------------------------------------------------------

def _load(job: _Job, context: _Context) -> None:
    job.stage("load")
    stored = tla._stored_execution(job.state["source"]["execution_id"])
    context.stored = stored
    context.code = tla._code_text(stored["code_id"])
    context.lines = context.code.split("\n")
    if context.lines and context.lines[-1] == "":
        context.lines.pop()
    context.prepared = runtime._platform("prepare_user_code")(context.code)
    frame, _metadata, _table = runtime._platform("load_execution_input")(stored["dataset_id"], stored.get("table_id"))
    context.frame = frame
    context.make_globals = runtime._platform("execution_globals")
    job.log("info", "code_loaded", f"Loaded {job.state['source']['code_filename']} ({len(context.lines)} lines)", file=job.state["source"]["code_filename"], lines=len(context.lines))
    job.log("info", "data_loaded", f"Loaded the input table: {len(frame)} records × {len(frame.columns)} columns", rows=len(frame), columns=len(frame.columns))


def _analyse(job: _Job, context: _Context) -> None:
    job.stage("trace")
    try:
        analysis = la.analyse(context.prepared, [str(column) for column in context.frame.columns])
    except la.LineageError as exc:
        raise HTTPException(400, str(exc)) from exc
    context.analysis = analysis
    context.nodes = analysis["nodes"]
    context.columns = {item["name"]: item for item in analysis["columns"]}
    context.steps_by_line = {step["line"]: step for step in analysis["steps"]}
    context.lookups = {item["name"]: item for item in analysis["lookups"]}
    context.dead = {item["node"] for item in analysis.get("dead_writes", [])}
    writes = sum(1 for node in analysis["nodes"] if node["kind"] == "write")
    job.log(
        "success", "static_done",
        f"Static data-flow analysis: {len(analysis['steps'])} statements, {writes} column writes, {len(analysis['lookups'])} parameter table(s)",
        statements=len(analysis["steps"]), writes=writes, lookups=len(analysis["lookups"]),
    )


def _writes_of(context: _Context, line: int, column: str) -> List[Dict[str, Any]]:
    """The writes of a statement to a column that matter (a write no later statement and no output uses is dead)."""
    step = context.steps_by_line.get(line)
    if step is None:
        return []
    return [context.nodes[node] for node in step["writes"] if context.nodes[node]["column"] == column and node not in context.dead]


def _lookup_use(context: _Context, line: int, column: str) -> Optional[Dict[str, Any]]:
    """The mapping table a statement applies to write a column, with the column holding its keys."""
    step = context.steps_by_line.get(line)
    if step is None:
        return None
    for node in _writes_of(context, line, column):
        if node.get("operation") == "cast":
            continue
        for name in node.get("lookups", []):
            lookup = context.lookups.get(name)
            if not lookup or lookup["kind"] != "mapping":
                continue
            use = next(
                (item for item in lookup.get("uses", []) if item.get("method") == "map" and item.get("column") and step["line"] <= item["line"] <= step["end_line"]),
                None,
            )
            if use:
                return {"name": name, "key_column": use["column"], "entries": lookup["entries"], "line": lookup["line"], "end_line": lookup["end_line"]}
    return None


def _reads(context: _Context) -> Dict[Tuple[int, str], List[str]]:
    """Per (statement, written column): the columns the statement reads to write it."""
    reads: Dict[Tuple[int, str], List[str]] = {}
    analysis = context.analysis
    for node in context.nodes:
        if node["kind"] != "write" or node.get("step") is None:
            continue
        line = analysis["steps"][node["step"]]["line"]
        used = reads.setdefault((line, node["column"]), [])
        for dependency in [*[item["node"] for item in node["data"]], *node["control"], *node["group"], *node.get("join", [])]:
            name = context.nodes[dependency]["column"]
            if name not in used:
                used.append(name)
    return reads


def _trace(job: _Job, context: _Context) -> None:
    """Replay the code statement by statement; record which statement wrote every cell, the lookup keys it mapped and,
    for the first records, the values before and after every write together with the values it read."""
    frame = context.frame
    partial = len(frame) > REPLAY_ROW_LIMIT
    if partial:
        frame = frame.iloc[:REPLAY_ROW_LIMIT]
    reads = _reads(context)
    lookups: Dict[Tuple[int, str], Dict[str, Any]] = {}
    for step in context.analysis["steps"]:
        for node_id in step["writes"]:
            column = context.nodes[node_id]["column"]
            use = _lookup_use(context, step["line"], column)
            if use:
                lookups[(step["line"], column)] = use
    started = time.perf_counter()
    tree = ast.parse(context.prepared)
    namespace = context.make_globals(frame.copy(deep=True))
    previous = namespace["df"].copy(deep=True)
    context.record_labels = list(frame.index)[:RECORD_ROWS]
    recorded = set(context.record_labels)
    error: Optional[Dict[str, Any]] = None
    for statement in tree.body:
        job.check()
        line = statement.lineno
        step_started = time.perf_counter()
        try:
            exec(compile(ast.Module(body=[statement], type_ignores=[]), "<script>", "exec"), namespace)
        except Exception as exc:
            error = {"line": line, "message": f"{type(exc).__name__}: {exc}"[:400]}
            context.statements.append({"line": line, "end_line": statement.end_lineno or line, "error": error["message"], "changed": {}, "rows_removed": 0, "nulled": {}})
            break
        current = namespace.get("df")
        if not isinstance(current, pd.DataFrame):
            error = {"line": line, "message": "df is no longer a DataFrame after this statement"}
            context.statements.append({"line": line, "end_line": statement.end_lineno or line, "error": error["message"], "changed": {}, "rows_removed": 0, "nulled": {}})
            break
        diff = la.compare_frames(previous, current)
        record = {
            "line": line, "end_line": statement.end_lineno or line, "error": None,
            "changed": {str(column): len(labels) for column, labels in diff["changes"].items() if labels},
            "rows_removed": diff["rows_removed"], "rows_added": diff["rows_added"], "nulled": {},
            "ms": round((time.perf_counter() - step_started) * 1000, 2),
        }
        for column, labels in diff["changes"].items():
            name = str(column)
            if name in diff["added"]:
                labels = list(current.index)  # a new column is written in every record, also where it stays empty
            if not labels:
                continue
            use = lookups.get((line, name))
            key_column = use["key_column"] if use else None
            history = context.history.setdefault(name, {})
            reading = reads.get((line, name), [])
            for label in labels:
                history.setdefault(label, []).append(line)
                before = previous.at[label, column] if column in previous.columns and label in previous.index else None
                after = current.at[label, column]
                if not la._is_null(before) and la._is_null(after):
                    entry = record["nulled"].setdefault(name, {"count": 0, "examples": []})
                    entry["count"] += 1
                    example = la.json_value(before)
                    if example not in entry["examples"] and len(entry["examples"]) < 5:
                        entry["examples"].append(example)
                if key_column and key_column in previous.columns and label in previous.index:
                    context.keys.setdefault((line, name), {})[label] = previous.at[label, key_column]
                if label in recorded:
                    inputs = {}
                    for source in reading:
                        if source != name and source in previous.columns and label in previous.index:
                            inputs[source] = la.json_value(previous.at[label, source])
                    context.details.setdefault(name, {}).setdefault(label, []).append({
                        "line": line, "before": la.json_value(before), "after": la.json_value(after), "inputs": inputs,
                    })
        # Categories the parameter table does not cover: the mapping yields nothing and the value stays empty.
        for (lookup_line, name), use in lookups.items():
            if lookup_line != line or name not in current.columns or use["key_column"] not in previous.columns:
                continue
            keys = [key for key, _ in use["entries"]]
            for label in current.index:
                if label not in previous.index or not la._is_null(current.at[label, name]):
                    continue
                value = previous.at[label, use["key_column"]]
                if not la._is_null(value) and any(la.same_value(value, key) for key in keys):
                    continue
                missing = context.unmapped.setdefault((line, name), {})
                missing[_key_text(value)] = missing.get(_key_text(value), 0) + 1
        context.statements.append(record)
        previous = current.copy(deep=True)
    final = namespace.get("df") if isinstance(namespace.get("df"), pd.DataFrame) else previous
    context.final = final
    elapsed = round((time.perf_counter() - started) * 1000)
    stored_rows = context.stored.get("data") or []
    stored_columns = (context.stored.get("summary") or {}).get("columns") or []
    if partial:
        comparison = la.compare_with_execution(final, stored_rows[:REPLAY_ROW_LIMIT], stored_columns)
        comparison["rows_match"] = True
        comparison["ok"] = comparison["columns_match"] and comparison["mismatches"] == 0
    else:
        comparison = la.compare_with_execution(final, stored_rows, stored_columns)
    comparison.update({"ms": elapsed, "partial": partial, "rows": len(frame), "error": error})
    comparison["examples"] = comparison.get("examples", [])[:10]
    context.replay = comparison
    for statement in context.statements:
        if statement.get("rows_removed"):
            step = context.steps_by_line.get(statement["line"])
            context.filters.append({
                "id": f"F{len(context.filters) + 1}", "line": statement["line"], "end_line": statement["end_line"],
                "code": (step or {}).get("code") or context.code_span(statement["line"], statement["end_line"]), "rows_removed": statement["rows_removed"],
            })
    if error:
        job.log("error", "replay_error", f"Replay stopped at line {error['line']}: {error['message']}", line=error["line"], message=error["message"])
        raise HTTPException(400, f"The code fails at line {error['line']} when it is replayed: {error['message']}")
    cells = sum(sum(item["changed"].values()) for item in context.statements)
    if comparison["ok"]:
        job.log("success", "replay_done", f"Replay reproduces the stored execution: {comparison['cells_compared']} values identical; {cells} value changes attributed to statements ({elapsed} ms)", cells=comparison["cells_compared"], changes=cells, ms=elapsed)
    elif not comparison["columns_match"]:
        job.log("warn", "replay_columns", "The code now produces other columns than the stored execution; the content is traced on the replayed result.", extra=comparison["extra_columns"], missing=comparison["missing_columns"])
    else:
        job.log("warn", "replay_mismatch", f"Replay differs from the stored execution in {comparison['mismatches']} of {comparison['cells_compared']} values (the input data may have changed since)", mismatches=comparison["mismatches"], cells=comparison["cells_compared"])


# --------------------------------------------------------------------------------------------------
# 2. Follow: where the selected field's content flows, the rules on the way, record attribution, impact test
# --------------------------------------------------------------------------------------------------

def _dependencies(node: Dict[str, Any]) -> List[Tuple[int, str]]:
    """(version node, role) of everything a write reads: the values it computes with, the conditions that select its
    rows, grouping keys and join keys."""
    found = [(int(item["node"]), "value") for item in node.get("data", [])]
    found += [(int(item), "condition") for item in node.get("control", [])]
    found += [(int(item), "group") for item in node.get("group", [])]
    found += [(int(item), "join") for item in node.get("join", [])]
    return found


def _first_line(analysis: Dict[str, Any], dead: Set[int], column: str) -> int:
    """The first statement that derives a column (more than a type conversion of itself)."""
    lines = [
        analysis["steps"][node["step"]]["line"] for node in analysis["nodes"]
        if node["column"] == column and node["kind"] == "write" and node.get("step") is not None and node["id"] not in dead
        and (node.get("operation") != "cast" or any(analysis["nodes"][dependency]["column"] != column for dependency, _ in _dependencies(node)))
    ]
    return min(lines) if lines else 0


def forward_structure(analysis: Dict[str, Any], field: str) -> Dict[str, Any]:
    """Every field the content of `field` flows into, computed on column versions.

    A write that reads a version carrying the field's content creates a version that carries it too; a field is
    downstream when one of its versions carries the content. Writes that no later statement and no output uses are dead
    and carry nothing, so a value that is overwritten before anything reads it creates no link. Dataset operations
    (filters, sorts, joins) that read a carrying version decide which records the result keeps."""
    nodes = analysis["nodes"]
    steps = analysis["steps"]
    dead = {item["node"] for item in analysis.get("dead_writes", [])}
    dependents: Dict[int, List[Tuple[int, str]]] = {}
    for node in nodes:
        if node["kind"] != "write" or node["id"] in dead:
            continue
        for dependency, role in _dependencies(node):
            dependents.setdefault(dependency, []).append((node["id"], role))
    carrying = {node["id"] for node in nodes if node["column"] == field}
    queue = sorted(carrying)
    edges: Dict[Tuple[str, str], Dict[str, Set[Any]]] = {}
    while queue:
        current = queue.pop(0)
        source = nodes[current]["column"]
        for dependent, role in dependents.get(current, []):
            target = nodes[dependent]["column"]
            if target != source and target != field:
                entry = edges.setdefault((source, target), {"lines": set(), "roles": set()})
                entry["lines"].add(steps[nodes[dependent]["step"]]["line"])
                entry["roles"].add(role)
            if dependent not in carrying:
                carrying.add(dependent)
                queue.append(dependent)
    reached = {nodes[node]["column"] for node in carrying} - {field}
    successors: Dict[str, Set[str]] = {}
    for source, target in edges:
        successors.setdefault(source, set()).add(target)
    distance = {field: 0}
    frontier = [field]
    while frontier:
        following = []
        for name in frontier:
            for target in sorted(successors.get(name, ())):
                if target not in distance:
                    distance[target] = distance[name] + 1
                    following.append(target)
        frontier = following
    # Position in the calculation: the longest route from the field (bounded, so a cycle between versions ends).
    level = dict(distance)
    for _ in range(len(reached) + 1):
        changed = False
        for source, target in edges:
            if source not in level:
                continue
            wanted = level[source] + 1
            if wanted > level.get(target, 0) and wanted <= len(reached):
                level[target] = wanted
                changed = True
        if not changed:
            break
    ends = {name for name in reached if not (successors.get(name, set()) - {name})}
    order = sorted(reached, key=lambda name: (level.get(name, 1), _first_line(analysis, dead, name), name))
    selection = []
    for op in analysis.get("dataset_ops", []):
        op_nodes = op.get("nodes")
        if op_nodes is not None:
            affected = any(int(item) in carrying for item in op_nodes)
        else:
            affected = any(column == field or column in reached for column in op.get("columns", []))
        if affected:
            selection.append({"kind": op["kind"], "line": op["line"], "code": op.get("code"), "columns": [column for column in op.get("columns", []) if column == field or column in reached]})
    return {
        "field": field,
        "fields": [
            {"column": name, "level": level.get(name, 1), "distance": distance.get(name, 1), "direct": (field, name) in edges, "end": name in ends}
            for name in order
        ],
        "edges": [{"from": source, "to": target, "lines": sorted(entry["lines"]), "roles": sorted(entry["roles"])} for (source, target), entry in sorted(edges.items())],
        "next": [name for name in order if (field, name) in edges],
        "ends": [name for name in order if name in ends],
        "selection": selection,
        "carrying": carrying,
    }


def field_names(analysis: Dict[str, Any], inputs: List[str]) -> List[str]:
    """Every field of an execution: the source fields in their order, then the ones the code creates, as it does."""
    dead = {item["node"] for item in analysis.get("dead_writes", [])}
    columns = {item["name"] for item in analysis["columns"]}
    names = [name for name in inputs if name in columns]
    return names + sorted((name for name in columns if name not in inputs), key=lambda name: (_first_line(analysis, dead, name), name))


def field_catalogue(analysis: Dict[str, Any], inputs: List[str], numeric: Set[str], probe: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    """Every field of an execution with what its content flows into: the same computation as an analysis, completed
    by the quick impact test of the preview (fields it finds that the code analysis does not link are marked)."""
    columns = {item["name"]: item for item in analysis["columns"]}
    probe = probe or {}
    catalogue = []
    for name in field_names(analysis, inputs):
        structure = forward_structure(analysis, name)
        fields = [{**{key: item[key] for key in ("column", "level", "direct", "end")}, "evidence": "code"} for item in structure["fields"]]
        found = (probe.get("found") or {}).get(name, [])
        known = {item["column"] for item in fields}
        extra = [column for column in found if column not in known and column != name and column in columns]
        deepest = max([item["level"] for item in fields] + [0])
        fields += [{"column": column, "level": deepest + 1, "direct": False, "end": True, "evidence": "impact"} for column in extra]
        catalogue.append({
            "column": name,
            "origin": "source" if name in inputs else "derived" if columns[name]["in_output"] else "intermediate",
            "in_output": bool(columns[name]["in_output"]),
            "numeric": name in numeric,
            "next": structure["next"],
            "ends": [*structure["ends"], *extra],
            "fields": fields,
            "edges": [{"from": edge["from"], "to": edge["to"], "lines": edge["lines"]} for edge in structure["edges"]],
            "selection": [{"kind": item["kind"], "line": item["line"]} for item in structure["selection"]],
            "tested": name in (probe.get("tested") or []),
            "confirmed": [column for column in known if column in found],
        })
    return catalogue


def _numeric_columns(frame: Any) -> Set[str]:
    numeric = set()
    if frame is None:
        return numeric
    for column in frame.columns:
        series = frame[column]
        if pd.api.types.is_bool_dtype(series):
            continue
        if pd.api.types.is_numeric_dtype(series):
            numeric.add(str(column))
            continue
        values = [value for value in series.tolist() if not la._is_null(value)]
        if values and all(_number(value) is not None for value in values):
            numeric.add(str(column))
    return numeric


def _factor_columns(frame: Any, numeric: Set[str]) -> Set[str]:
    """Numeric columns that hold factors or rates (every value between 0 and 12.5, some not whole): their sum means
    nothing, so no amounts are added up for them."""
    factors = set()
    for name in numeric:
        if name not in frame.columns:
            continue
        values = [_number(value) for value in frame[name].tolist()]
        values = [value for value in values if value is not None]
        if values and all(0 <= value <= 12.5 for value in values) and any(not float(value).is_integer() for value in values):
            factors.add(name)
    return factors


def _value_type(series: Any, numeric: bool) -> str:
    if numeric:
        return "number"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "date"
    values = [value for value in series.tolist() if not la._is_null(value)]
    if not values:
        return "empty"
    if all(isinstance(value, (bool,)) for value in values):
        return "boolean"
    if all(isinstance(value, (pd.Timestamp, datetime)) for value in values):
        return "date"
    return "text"


_IDENTIFIER_NAME = re.compile(r"(?:^|[\s_\-])(id|position|pos|key|nr|no|number|account|contract|deal|loan|darlehen|konto|vertrag|gesch[aä]ft)(?:$|[\s_\-])|id$", re.IGNORECASE)


def _identifier(frame: Any, exclude: Set[str]) -> Optional[str]:
    candidates = []
    for column in frame.columns:
        name = str(column)
        if name in exclude:
            continue
        series = frame[column]
        if len(series) == 0 or series.isna().any():
            continue
        if pd.api.types.is_numeric_dtype(series) and not _IDENTIFIER_NAME.search(name):
            continue
        if series.astype(str).nunique() != len(series):
            continue
        candidates.append(name)
    preferred = [name for name in candidates if _IDENTIFIER_NAME.search(name)]
    return (preferred or candidates or [None])[0]


def _constants(text: str) -> Set[float]:
    """Numbers written in code or a table (as a parameter may be cited)."""
    values: Set[float] = set()
    try:
        tree = ast.parse(text)
        for node in ast.walk(tree):
            if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) and not isinstance(node.value, bool):
                values.add(float(node.value))
            elif isinstance(node, ast.Constant) and isinstance(node.value, str):
                values |= set(_numbers_in(node.value, single=True))
    except SyntaxError:
        values |= set(_numbers_in(text, single=True))
    return values


_NUMBER = re.compile(r"(?<![\w])(\d+(?:[.,]\d+)*)(?![\w])")


def _numbers_in(text: str, single: bool = False) -> List[Any]:
    """Every number written in a text: one value, or (single=False) the set of its possible readings
    ('1.000' is 1.0 in English notation and 1000 in German notation)."""
    found: List[Any] = []
    for match in _NUMBER.finditer(text or ""):
        raw = match.group(1)
        readings: Set[float] = set()
        if re.fullmatch(r"\d+", raw):
            readings.add(float(raw))
        else:
            if re.fullmatch(r"\d+[.,]\d+", raw):
                readings.add(float(raw.replace(",", ".")))
            if re.fullmatch(r"\d{1,3}(?:[.,]\d{3})+", raw):
                readings.add(float(re.sub(r"[.,]", "", raw)))
            if not readings:
                parts = re.split(r"[.,]", raw)
                readings.add(float(parts[0]))
        if single:
            found.extend(readings)
        else:
            found.append(readings)
    return found


def _is_case_of(context: _Context, nodes: List[Dict[str, Any]], column: str) -> bool:
    """A statement is a case of a column's rule unless it only converts the column's own type."""
    return any(
        node.get("operation") != "cast" or any(context.nodes[dependency]["column"] != column for dependency, _ in _dependencies(node))
        for node in nodes
    )


def _add_rule(context: _Context, column: str) -> None:
    """The business rule of a downstream field: one case per statement that derives it — all of them, also the cases
    that do not use the selected field (they produce the other records' values). Rule ids follow the calculation."""
    lines = sorted({
        context.analysis["steps"][node["step"]]["line"] for node in context.nodes
        if node["column"] == column and node["kind"] == "write" and node.get("step") is not None and node["id"] not in context.dead
    })
    rule = {"id": f"R{len(context.rules) + 1}", "column": column, "cases": [], "inputs": [], "selectors": [], "roles": {}, "lines": []}
    for line in lines:
        nodes = _writes_of(context, line, column)
        if not nodes or not _is_case_of(context, nodes, column):
            continue
        step = context.steps_by_line[line]
        use = _lookup_use(context, line, column)
        roles: Dict[str, List[str]] = {}
        for node in nodes:
            for dependency, role in _dependencies(node):
                name = context.nodes[dependency]["column"]
                if name == column:
                    continue
                bucket = roles.setdefault(name, [])
                if role not in bucket:
                    bucket.append(role)
        if use and use["key_column"] != column:
            others = [role for role in roles.get(use["key_column"], []) if role not in ("value", "key")]
            roles[use["key_column"]] = ["key", *others]
        operations = [node.get("operation") for node in nodes if node.get("operation") != "cast"]
        conditions = [condition.get("pretty") or condition.get("code") for node in nodes for condition in node.get("conditions", []) if condition.get("pretty") or condition.get("code")]
        formulas = [node.get("pretty") or node.get("expression") or "" for node in nodes if node.get("operation") != "cast" or roles]
        code = step.get("code") or context.code_span(step["line"], step["end_line"])
        constants = _constants(code)
        if use:
            for key, value in use["entries"]:
                for item in (key, value):
                    number = _number(item)
                    if number is not None:
                        constants.add(number)
                    elif isinstance(item, str):
                        constants |= set(_numbers_in(item, single=True))
        for node in nodes:
            for variable in node.get("uses", []):
                span = context.analysis.get("variables", {}).get(variable)
                if span:
                    constants |= _constants(context.code_span(span["line"], span["end_line"]))
        case = {
            "id": f"{rule['id']}.{len(rule['cases']) + 1}",
            "rule": rule["id"],
            "column": column,
            "line": step["line"],
            "end_line": step["end_line"],
            "code": code,
            "operation": operations[0] if operations else "conversion",
            "fills_missing": all(operation == "fill_missing" for operation in operations) if operations else False,
            "condition": " & ".join(dict.fromkeys(conditions)) or None,
            "formula": " ; ".join(dict.fromkeys(item for item in formulas if item)) or None,
            "inputs": [name for name, items in roles.items() if "value" in items],
            "selectors": [name for name, items in roles.items() if "value" not in items],
            "roles": roles,
            "lookup": {key: use[key] for key in ("name", "key_column", "entries", "line", "end_line")} if use else None,
            "constants": sorted(constants),
            "records": 0,
            "amount": None,
            "keys": [],
        }
        rule["cases"].append(case)
        context.cases[case["id"]] = case
        context.case_by_write[(line, column)] = case
        for name, items in roles.items():
            bucket = rule["roles"].setdefault(name, [])
            bucket.extend(role for role in items if role not in bucket)
        rule["lines"].append(line)
    if rule["cases"]:
        rule["inputs"] = [name for name, items in rule["roles"].items() if "value" in items]
        rule["selectors"] = [name for name, items in rule["roles"].items() if "value" not in items]
        context.rules.append(rule)
        context.rule_by_column[column] = rule


def _compose(job: _Job, context: _Context, field: str) -> None:
    job.stage("follow")
    known = {item["name"] for item in context.analysis["columns"]}
    if field not in known:
        raise HTTPException(400, f"The code has no field '{field}'.")
    structure = forward_structure(context.analysis, field)
    context.field = field
    context.structure = structure
    context.reached = [item["column"] for item in structure["fields"]]
    for column in context.reached:
        _add_rule(context, column)
    context.derived = {rule["column"] for rule in context.rules}
    others: List[str] = []
    for rule in context.rules:
        for name in rule["roles"]:
            if name != field and name not in context.reached and name not in others:
                others.append(name)
    context.context_fields = others
    context.scope = [field, *context.reached, *others]
    final = context.final
    context.identifier = _identifier(final, {field}) if final is not None else None
    cases = sum(len(rule["cases"]) for rule in context.rules)
    if structure["fields"]:
        job.log(
            "success", "scope",
            f"[{field}] flows into {len(structure['next'])} next field(s) and {len(context.reached)} downstream field(s) in total, {len(structure['ends'])} of them end field(s); {len(context.rules)} business rules with {cases} cases",
            field=field, next=len(structure["next"]), fields=len(context.reached), ends=len(structure["ends"]), rules=len(context.rules), cases=cases,
            names=structure["next"],
        )
    else:
        job.log("info", "scope_empty", f"The code analysis links [{field}] to no other field; the impact test checks every field.", field=field)
    if structure["selection"]:
        job.log("info", "selection", f"[{field}] also decides which records are kept (line(s) {', '.join(str(item['line']) for item in structure['selection'])})", field=field, lines=[str(item["line"]) for item in structure["selection"]])


def _require_downstream(context: _Context) -> None:
    if not context.reached:
        checked = " and the impact test" if context.impact.get("enabled") else ""
        raise HTTPException(400, f"No other field depends on '{context.field}': the code analysis{checked} found no field its content flows into.")


def _origin(context: _Context, column: str, label: Any, before: Optional[int] = None) -> Tuple[Optional[Dict[str, Any]], Optional[int]]:
    """The case that last wrote a cell before a statement (None: the final value), and its statement line."""
    for line in reversed(context.history.get(column, {}).get(label, [])):
        if before is not None and line >= before:
            continue
        case = context.case_by_write.get((line, column))
        if case is not None:
            return case, line
    return None, None


def _reach(context: _Context, column: str, label: Any, before: Optional[int], memo: Dict[Tuple[str, Any, Optional[int]], Any]) -> Optional[Tuple[str, List[str]]]:
    """Whether the value a record has in a column (before a statement; None: its final value) is built from the
    selected field: ('value', via) when the field's value is carried into it, ('decides', via) when the field only
    decides which case or parameter applies, None when the case that produced the value uses neither the field nor
    anything built from it (or the source delivered the value). via: what the producing case read that carries it."""
    if column == context.field:
        return ("start", [])
    key = (column, label, before)
    if key in memo:
        return memo[key]
    memo[key] = None  # a cycle between versions carries nothing
    case, line = _origin(context, column, label, before)
    result = None
    if case is not None:
        via: List[str] = []
        value = False
        for other, roles in case["roles"].items():
            if other == column:
                continue
            found = _reach(context, other, label, line, memo)
            if found is None:
                continue
            via.append(other)
            if "value" in roles and found[0] in ("start", "value"):
                value = True
        if via:
            result = ("value" if value else "decides", via)
    memo[key] = result
    return result


def _route_carried(context: _Context, route: List[str], label: Any) -> bool:
    """Whether a record's content travels along a route of fields: every field on it was produced by a case that read
    the field before it (in the version that existed at that point)."""
    before: Optional[int] = None
    for index in range(len(route) - 1, 0, -1):
        column, previous = route[index], route[index - 1]
        case, line = _origin(context, column, label, before)
        if case is None or previous not in case["roles"]:
            return False
        before = line
    return True


# Impact test ---------------------------------------------------------------------------------------

def _value_classes(series: pd.Series) -> Tuple[List[str], Dict[str, Any], Dict[Any, str]]:
    """(classes in a fixed order, representative record per class, class per record) of a field's values; a missing
    value is a class of its own."""
    class_of: Dict[Any, str] = {}
    representative: Dict[str, Any] = {}
    sort_key: Dict[str, Tuple[Any, ...]] = {}
    for label, value in series.items():
        key = "∅" if la._is_null(value) else _key_text(value)
        class_of[label] = key
        if key not in representative:
            representative[key] = label
            number = _number(value)
            sort_key[key] = (1, 0.0, "") if key == "∅" else (0, number, "") if number is not None else (0, math.inf, key)
    classes = sorted(representative, key=lambda key: sort_key[key])
    return classes, representative, class_of


def _alternatives(reference: pd.Series, rotations: int = MAX_ROTATIONS, partials: int = 2) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """The other values the impact test gives a field, and its value classes.

    rotate: every record takes the value of a record of another value class (all classes in turn, missing values
    included); partial: only the records of one value take another value (moving every record at once keeps records
    that share a value together, so a sum per group would not change); shift: every value becomes one that does not
    occur in the data (an unknown category, another amount)."""
    classes, representative, class_of = _value_classes(reference)
    alternatives: List[Dict[str, Any]] = []
    for offset in range(1, min(len(classes), rotations + 1)):
        alternatives.append({"id": f"rotate:{offset}", "kind": "rotate", "offset": offset})
    for index in range(min(len(classes) - 1, partials)):
        alternatives.append({"id": f"partial:{index}", "kind": "partial", "source": index, "target": index + 1})
    if any(not la._is_null(value) for value in reference.tolist()):
        alternatives.append({"id": "shift", "kind": "shift"})
    return alternatives, {"classes": classes, "representative": representative, "class_of": class_of}


def _alter(alternative: Dict[str, Any], values: Dict[str, Any], snapshots: Dict[int, pd.Series]) -> Any:
    """The function that gives a field its other values at a point of the run, from its unchanged values there: a
    record takes the value its counterpart had at the same point, so it behaves throughout as if its field were the
    counterpart's."""
    classes, representative, class_of = values["classes"], values["representative"], values["class_of"]
    if alternative["kind"] in ("rotate", "partial"):
        if alternative["kind"] == "rotate":
            offset = alternative["offset"]
            position = {key: index for index, key in enumerate(classes)}
            target = {label: representative[classes[(position[key] + offset) % len(classes)]] for label, key in class_of.items()}
        else:
            moved, into = classes[alternative["source"]], representative[classes[alternative["target"]]]
            target = {label: into for label, key in class_of.items() if key == moved}

        def alter(point: int, current: pd.Series) -> pd.Series:
            snapshot = snapshots[point]
            values_ = [snapshot.at[target[label]] if label in target and target[label] in snapshot.index else current.at[label] for label in current.index]
            return la._typed_series(values_, current)
        return alter

    def shift(point: int, current: pd.Series) -> pd.Series:
        snapshot = snapshots[point]
        values_ = [la._perturb(snapshot.at[label]) if label in snapshot.index else current.at[label] for label in current.index]
        return la._typed_series(values_, current)
    return shift


class _Replayer:
    """Re-runs the code statement by statement, optionally with other values of one field wherever it gets its value:
    at the start when the source delivers it and after every statement that writes or changes it."""

    def __init__(self, prepared: str, frame: pd.DataFrame, make_globals: Any, analysis: Dict[str, Any]):
        self.statements = ast.parse(prepared).body
        self.frame = frame
        self.make_globals = make_globals
        nodes = analysis["nodes"]
        self.write_lines: Dict[str, Set[int]] = {}
        for step in analysis["steps"]:
            for node in step["writes"]:
                self.write_lines.setdefault(nodes[node]["column"], set()).add(step["line"])
        # Columns the code drops before the end: their last state is compared too.
        self.intermediates = {item["name"] for item in analysis["columns"] if not item["in_output"]}

    def _writes(self, statement: ast.stmt, field: str) -> bool:
        end = statement.end_lineno or statement.lineno
        return any(statement.lineno <= line <= end for line in self.write_lines.get(field, ()))

    def baseline(self, fields: List[str]) -> Dict[str, Any]:
        """One unchanged run: per field the points where it gets its value and its values there, and the last state of
        every watched column."""
        namespace = self.make_globals(self.frame.copy(deep=True))
        watch = self.intermediates | set(fields)
        points: Dict[str, List[int]] = {field: [] for field in fields}
        snapshots: Dict[str, Dict[int, pd.Series]] = {field: {} for field in fields}
        last: Dict[str, pd.Series] = {}
        df = namespace["df"]
        previous: Dict[str, pd.Series] = {}
        for field in fields:
            if field in df.columns:
                points[field].append(-1)
                snapshots[field][-1] = df[field].copy()
                previous[field] = df[field].copy()
        self._remember(df, watch, last)
        for index, statement in enumerate(self.statements):
            try:
                exec(compile(ast.Module(body=[statement], type_ignores=[]), "<script>", "exec"), namespace)
            except Exception as exc:
                return {"error": {"line": statement.lineno, "message": f"{type(exc).__name__}: {exc}"[:300]}}
            df = namespace.get("df")
            if not isinstance(df, pd.DataFrame):
                return {"error": {"line": statement.lineno, "message": "df is no longer a DataFrame"}}
            for field in fields:
                if field not in df.columns:
                    previous.pop(field, None)
                    continue
                current = df[field]
                prior = previous.get(field)
                if prior is None:
                    moved = True
                else:
                    shared = current.index.intersection(prior.index)
                    moved = bool(la._changed_rows(prior.loc[shared], current.loc[shared]))
                if moved or self._writes(statement, field):
                    points[field].append(index)
                    snapshots[field][index] = current.copy()
                previous[field] = current.copy()
            self._remember(df, watch, last)
        return {"final": df, "points": points, "snapshots": snapshots, "last": last, "error": None}

    def run(self, field: str, points: List[int], alter: Any) -> Dict[str, Any]:
        namespace = self.make_globals(self.frame.copy(deep=True))
        watch = self.intermediates | {field}
        last: Dict[str, pd.Series] = {}
        wanted = set(points)

        def touch(point: int) -> None:
            df = namespace.get("df")
            if point in wanted and isinstance(df, pd.DataFrame) and field in df.columns:
                df[field] = alter(point, df[field])

        touch(-1)
        self._remember(namespace["df"], watch, last)
        for index, statement in enumerate(self.statements):
            try:
                exec(compile(ast.Module(body=[statement], type_ignores=[]), "<script>", "exec"), namespace)
            except Exception as exc:  # the code cannot handle the other values
                return {"error": {"line": statement.lineno, "message": f"{type(exc).__name__}: {exc}"[:300]}}
            df = namespace.get("df")
            if not isinstance(df, pd.DataFrame):
                return {"error": {"line": statement.lineno, "message": "df is no longer a DataFrame"}}
            touch(index)
            self._remember(df, watch, last)
        return {"final": namespace["df"], "last": last, "error": None}

    @staticmethod
    def _remember(df: Any, watch: Set[str], last: Dict[str, pd.Series]) -> None:
        if isinstance(df, pd.DataFrame):
            for name in watch:
                if name in df.columns:
                    last[name] = df[name].copy()


def _differences(base: Dict[str, Any], run: Dict[str, Any], field: str, intermediates: Set[str]) -> Tuple[Dict[str, List[Any]], int, int]:
    """(changed records per column, records the other values removed, records they added)."""
    base_final: pd.DataFrame = base["final"]
    final: pd.DataFrame = run["final"]
    shared = base_final.index.intersection(final.index)
    differences: Dict[str, List[Any]] = {}
    for column in base_final.columns:
        name = str(column)
        if name == field or column not in final.columns:
            continue
        labels = la._changed_rows(base_final.loc[shared, column], final.loc[shared, column])
        if labels:
            differences[name] = labels
    for name in intermediates - {field}:
        before, after = base["last"].get(name), run["last"].get(name)
        if before is None or after is None or name in base_final.columns:
            continue
        common = before.index.intersection(after.index)
        labels = la._changed_rows(before.loc[common], after.loc[common])
        if labels:
            differences[name] = labels
    return differences, int(len(base_final.index.difference(final.index))), int(len(final.index.difference(base_final.index)))


def _impact(job: _Job, context: _Context) -> None:
    """Re-run the code with other values of the selected field and compare every value of the result, record by record
    (see _alternatives and _Replayer). The test confirms the downstream fields the code analysis finds and catches
    any field the analysis cannot follow (for example values computed in a loop over the records)."""
    field = context.field
    frame = context.frame.iloc[:REPLAY_ROW_LIMIT] if len(context.frame) > REPLAY_ROW_LIMIT else context.frame
    started = time.perf_counter()
    report: Dict[str, Any] = {
        "enabled": False, "reason": None, "rows": len(frame), "partial": len(context.frame) > REPLAY_ROW_LIMIT,
        "classes": 0, "all_values": True, "alternatives": [], "planned": 0, "tested": 0, "budget_hit": False, "errors": [],
        "unexpected": [], "silent": [], "confirmed": 0, "selection": None, "ms": 0,
    }
    context.impact = report
    if not frame.index.is_unique:
        report["reason"] = "index"
        job.log("warn", "impact_skipped", "Impact test skipped: the records have no unique index.", reason="index")
        return
    replayer = _Replayer(context.prepared, frame, context.make_globals, context.analysis)
    base = replayer.baseline([field])
    if base.get("error"):
        report["reason"] = "error"
        job.log("warn", "impact_skipped", f"Impact test skipped: the code fails when it is re-run ({base['error']['message']})", reason="error", message=base["error"]["message"])
        return
    base_final: pd.DataFrame = base["final"]
    check = la.compare_frames(context.final, base_final) if context.final is not None else None
    if check is None or check["rows_removed"] or check["rows_added"] or any(check["changes"].values()):
        report["reason"] = "nondeterministic"
        job.log("warn", "impact_skipped", "Impact test skipped: the code gives a different result each time it runs.", reason="nondeterministic")
        return
    reference = base["last"].get(field)
    if reference is None:
        report["reason"] = "absent"
        return
    context.x_values = {label: value for label, value in reference.items()}
    alternatives, values = _alternatives(reference)
    report["classes"] = len(values["classes"])
    report["all_values"] = len(values["classes"]) - 1 <= MAX_ROTATIONS
    report["planned"] = len(alternatives)
    job.log(
        "info", "impact_start",
        f"Impact test: re-running the code {len(alternatives)} time(s) with other values of [{field}] ({len(values['classes'])} value class(es) in {len(frame)} records)",
        field=field, runs=len(alternatives), classes=len(values["classes"]), rows=len(frame),
    )
    first_records = set(context.record_labels)
    order = {label: position for position, label in enumerate(reference.index)}
    changed: Dict[str, Set[Any]] = {}
    for alternative in alternatives:
        job.check()
        if time.perf_counter() - started > IMPACT_BUDGET_SECONDS:
            report["budget_hit"] = True
            break
        moment = time.perf_counter()
        run = replayer.run(field, base["points"][field], _alter(alternative, values, base["snapshots"][field]))
        entry = {"id": alternative["id"], "kind": alternative["kind"], "changed": {}, "error": None, "rows_removed": 0, "rows_added": 0}
        report["tested"] += 1
        if run.get("error"):
            entry["error"] = run["error"]
            report["errors"].append({"alternative": alternative["id"], **run["error"]})
            report["alternatives"].append(entry)
            continue
        differences, removed, added = _differences(base, run, field, replayer.intermediates)
        entry["rows_removed"], entry["rows_added"] = removed, added
        final: pd.DataFrame = run["final"]
        x_after = run["last"].get(field)
        for column, labels in differences.items():
            entry["changed"][column] = len(labels)
            changed.setdefault(column, set()).update(labels)
            for label in sorted((label for label in labels if label in first_records), key=lambda item: order.get(item, 0)):
                if (column, label) in context.whatif:
                    continue
                if column in base_final.columns:
                    y_from, y_to = base_final.at[label, column], final.at[label, column]
                else:
                    y_from, y_to = base["last"][column].get(label), run["last"][column].get(label)
                context.whatif[(column, label)] = {
                    "alternative": alternative["id"],
                    "x_from": la.json_value(reference.get(label)),
                    "x_to": la.json_value(x_after.get(label)) if x_after is not None and label in x_after.index else None,
                    "y_from": la.json_value(y_from),
                    "y_to": la.json_value(y_to),
                }
        entry["ms"] = round((time.perf_counter() - moment) * 1000, 1)
        report["alternatives"].append(entry)
    context.changed = changed
    removed = max((item["rows_removed"] for item in report["alternatives"]), default=0)
    added = max((item["rows_added"] for item in report["alternatives"]), default=0)
    report["selection"] = {"removed": removed, "added": added} if removed or added else None
    report["enabled"] = report["tested"] > len(report["errors"])
    report["unexpected"] = sorted(column for column in changed if column not in context.reached and column != field)
    report["ms"] = round((time.perf_counter() - started) * 1000)
    if report["budget_hit"]:
        job.log("warn", "impact_budget", f"Impact test time budget reached after {report['tested']} of {report['planned']} runs", done=report["tested"], total=report["planned"])
    for error in report["errors"]:
        job.log("info", "impact_error", f"With other values of [{field}] the code fails at line {error['line']}: {error['message']}", field=field, line=error["line"], message=error["message"])


def _admit_unexpected(job: _Job, context: _Context) -> None:
    """Fields the impact test found to change although the code analysis does not link them to the selected field are
    added to the downstream fields, marked as found by the test, with their business rule."""
    known = {item["name"] for item in context.analysis["columns"]}
    unexpected = [column for column in context.impact.get("unexpected", []) if column in known and column not in context.reached]
    for column in unexpected:
        context.structure["fields"].append({"column": column, "level": max([item["level"] for item in context.structure["fields"]] + [0]) + 1, "distance": None, "direct": False, "end": True, "evidence": "impact"})
        context.reached.append(column)
        if column in context.context_fields:
            context.context_fields.remove(column)
        _add_rule(context, column)
    if unexpected:
        # Links among the downstream fields the code shows (a field found by the test may feed another), and a link
        # from the selected field, marked as found by the test, where the code shows none.
        structure = context.structure
        carriers = {context.field, *context.reached}
        known_edges = {(edge["from"], edge["to"]) for edge in structure["edges"]}
        for column in unexpected:
            for edge in forward_structure(context.analysis, column)["edges"]:
                if edge["from"] in carriers and edge["to"] in carriers and edge["to"] != context.field and (edge["from"], edge["to"]) not in known_edges:
                    structure["edges"].append({key: edge[key] for key in ("from", "to", "lines", "roles")})
                    known_edges.add((edge["from"], edge["to"]))
        for column in unexpected:
            if not any(target == column for _, target in known_edges):
                structure["edges"].append({"from": context.field, "to": column, "lines": [], "roles": ["impact"]})
                known_edges.add((context.field, column))
        level = {context.field: 0}
        for _ in range(len(context.reached) + 1):
            changed = False
            for source, target in known_edges:
                if source in level and level[source] + 1 > level.get(target, 0) and level[source] + 1 <= len(context.reached):
                    level[target] = level[source] + 1
                    changed = True
            if not changed:
                break
        successors = {source for source, target in known_edges if target != source}
        for item in structure["fields"]:
            item["level"] = level.get(item["column"], item["level"])
            item["end"] = item["column"] not in successors
        structure["ends"] = [item["column"] for item in structure["fields"] if item["end"]]
        structure["fields"].sort(key=lambda item: (item["level"], _first_line(context.analysis, context.dead, item["column"]), item["column"]))
        context.reached = [item["column"] for item in structure["fields"]]
        context.derived = {rule["column"] for rule in context.rules}
        context.context_fields = [
            name for name in dict.fromkeys([*context.context_fields, *[name for rule in context.rules for name in rule["roles"]]])
            if name != context.field and name not in context.reached
        ]
        context.scope = [context.field, *context.reached, *context.context_fields]
        job.log("warn", "impact_added", f"Found by the impact test only (the code analysis cannot follow how): {', '.join(unexpected)}", columns=unexpected)
    report = context.impact
    report["silent"] = [column for column in context.reached if not context.changed.get(column)]
    report["confirmed"] = sum(1 for column in context.reached if context.changed.get(column))
    if report.get("enabled"):
        job.log(
            "success" if not unexpected else "warn", "impact_done",
            f"Impact test: {report['confirmed']} of {len(context.reached)} downstream field(s) change when [{context.field}] changes",
            field=context.field, confirmed=report["confirmed"], fields=len(context.reached), ms=report["ms"],
        )


def _probe_catalogue(analysis: Dict[str, Any], prepared: str, frame: pd.DataFrame, make_globals: Any, fields: List[str]) -> Dict[str, Any]:
    """A quick impact test of every field for the setup preview (a few alternatives on the first records): the fields
    whose values change although the code analysis does not link them."""
    rows = frame.iloc[:PREVIEW_ROWS] if len(frame) > PREVIEW_ROWS else frame
    started = time.perf_counter()
    result: Dict[str, Any] = {"enabled": False, "reason": None, "rows": len(rows), "found": {}, "tested": [], "budget_hit": False}
    if not rows.index.is_unique:
        result["reason"] = "index"
        return result
    replayer = _Replayer(prepared, rows, make_globals, analysis)
    base = replayer.baseline(fields)
    if base.get("error"):
        result["reason"] = "error"
        return result
    again = replayer.baseline([])
    check = la.compare_frames(base["final"], again["final"]) if not again.get("error") else None
    if check is None or check["rows_removed"] or check["rows_added"] or any(check["changes"].values()):
        result["reason"] = "nondeterministic"
        return result
    result["enabled"] = True
    for field in fields:
        if time.perf_counter() - started > PREVIEW_BUDGET_SECONDS:
            result["budget_hit"] = True
            break
        reference = base["last"].get(field)
        if reference is None:
            continue
        alternatives, values = _alternatives(reference, rotations=1, partials=1)
        found: Set[str] = set()
        for alternative in alternatives:
            run = replayer.run(field, base["points"][field], _alter(alternative, values, base["snapshots"][field]))
            if run.get("error"):
                continue
            differences, _, _ = _differences(base, run, field, replayer.intermediates)
            found |= set(differences)
        result["tested"].append(field)
        result["found"][field] = sorted(found)
    result["ms"] = round((time.perf_counter() - started) * 1000)
    return result


def _categorical_columns(context: _Context, final: Any, numeric: Set[str]) -> List[str]:
    """Fields of the result that classify records (text with few values): the fields in scope first."""
    names = [str(column) for column in final.columns]
    ordered = [name for name in context.scope if name in names] + [name for name in names if name not in context.scope]
    columns = []
    for name in ordered:
        if name in numeric or name == context.identifier:
            continue
        distinct = final[name].dropna().astype(str).nunique()
        if 1 <= distinct <= MAX_CATEGORIES:
            columns.append(name)
    return columns[:8]


def _population(final: Any, columns: List[str], labels: List[Any]) -> Dict[str, List[str]]:
    """What a case's records are: the values they have in the classifying fields (balance-sheet type, exposure class,
    product type ...), so that each case is described for the records it actually concerns."""
    found: Dict[str, List[str]] = {}
    for name in columns:
        values = list(dict.fromkeys(_key_text(final.at[label, name]) for label in labels))
        if values:
            found[name] = values[:6]
    return found


def _compositions(context: _Context, final: Any, labels: List[Any], summable: Set[str], factors: Set[str]) -> Dict[str, Dict[str, Any]]:
    """Per downstream field: values delivered by the source, derived by each case (records, amounts, lookup keys) and
    missing."""
    numeric = _numeric_columns(final)
    classifying = _categorical_columns(context, final, numeric)
    compositions: Dict[str, Dict[str, Any]] = {}
    for rule in context.rules:
        column = rule["column"]
        present = column in final.columns
        delivered = {"records": 0, "amount": 0.0}
        missing = 0
        by_case: Dict[str, Dict[str, Any]] = {}
        for label in labels:
            value = final.at[label, column] if present else None
            amount = _number(value) if present else None
            if present and la._is_null(value):
                missing += 1
                continue
            case, line = _origin(context, column, label)
            if case is None:
                if present:
                    delivered["records"] += 1
                    delivered["amount"] += amount or 0.0
                continue
            entry = by_case.setdefault(case["id"], {"records": 0, "amount": 0.0, "keys": {}, "labels": []})
            entry["records"] += 1
            entry["amount"] += amount or 0.0
            if len(entry["labels"]) < 5000:
                entry["labels"].append(label)
            key = context.keys.get((line, column), {}).get(label, _MISSING)
            if key is not _MISSING:
                text = _key_text(key)
                slot = entry["keys"].setdefault(text, {"records": 0, "amount": 0.0, "labels": []})
                slot["records"] += 1
                slot["amount"] += amount or 0.0
                if len(slot["labels"]) < 5000:
                    slot["labels"].append(label)
        is_numeric = column in summable and present
        classifiers = [name for name in rule["selectors"] if name in final.columns and name not in numeric]

        def categories(labels_of_key: List[Any], key_column: str) -> Dict[str, List[str]]:
            found: Dict[str, List[str]] = {}
            for name in classifiers:
                if name == key_column:
                    continue
                values = list(dict.fromkeys(_key_text(final.at[label, name]) for label in labels_of_key))
                if values:
                    found[name] = values[:12]
            return found

        for case in rule["cases"]:
            entry = by_case.get(case["id"], {"records": 0, "amount": 0.0, "keys": {}, "labels": []})
            case["records"] = entry["records"]
            case["population"] = _population(final, classifying, entry["labels"])
            case["amount"] = _round(entry["amount"]) if is_numeric else None
            if case["lookup"]:
                keys = []
                key_column = case["lookup"]["key_column"]
                for key, value in case["lookup"]["entries"]:
                    slot = entry["keys"].get(_key_text(key), {"records": 0, "amount": 0.0, "labels": []})
                    keys.append({"key": _key_text(key), "value": la.json_value(value), "records": slot["records"], "amount": _round(slot["amount"]) if is_numeric else None, "mapped": True, "categories": categories(slot["labels"], key_column)})
                for text, slot in entry["keys"].items():
                    if not any(item["key"] == text for item in keys):
                        keys.append({"key": text, "value": None, "records": slot["records"], "amount": _round(slot["amount"]) if is_numeric else None, "mapped": False, "categories": categories(slot["labels"], key_column)})
                for text, count in context.unmapped.get((case["line"], column), {}).items():
                    keys.append({"key": text, "value": None, "records": 0, "amount": None, "mapped": False, "unmapped_rows": count})
                case["keys"] = keys
        compositions[column] = {
            "total_records": len(labels),
            "in_result": present,
            "numeric": is_numeric,
            "factor": column in factors,
            "delivered": {"records": delivered["records"], "amount": _round(delivered["amount"]) if is_numeric else None},
            "missing": missing,
            "derived": sum(entry["records"] for entry in by_case.values()),
        }
    return compositions


def _follow(job: _Job, context: _Context) -> Dict[str, Any]:
    """The content of the downstream fields, record by record: which records each field reaches and how, per case, per
    value of the selected field, the routes to the end fields and the forward trace of the first records."""
    final = context.final
    field = context.field
    numeric = _numeric_columns(final)
    factors = _factor_columns(final, numeric)
    summable = numeric - factors
    labels = list(final.index)
    final_labels = set(labels)
    compositions = _compositions(context, final, labels, summable, factors)
    if not context.x_values and field in final.columns:
        context.x_values = {label: final.at[label, field] for label in labels}
    x_values = context.x_values
    x_present = [x_values[label] for label in labels if label in x_values and not la._is_null(x_values[label])]
    distinct = list(dict.fromkeys(_key_text(value) for value in x_present))
    categorical = field not in summable and 0 < len(distinct) <= MAX_CATEGORIES
    profile: List[Dict[str, Any]] = []
    if categorical:
        counts: Dict[str, Dict[str, Any]] = {}
        for label in labels:
            x = x_values.get(label)
            key = "∅" if x is None or la._is_null(x) else _key_text(x)
            slot = counts.setdefault(key, {"key": key, "value": None if key == "∅" else la.json_value(x), "records": 0})
            slot["records"] += 1
        profile = sorted(counts.values(), key=lambda item: (-item["records"], item["key"] == "∅", item["key"]))

    memo: Dict[Tuple[str, Any, Optional[int]], Any] = {}
    reach: Dict[Tuple[str, Any], Any] = {}
    for column in context.reached:
        for label in labels:
            reach[(column, label)] = _reach(context, column, label, None, memo)

    def category(column: str, label: Any) -> Optional[str]:
        found = reach[(column, label)]
        if found is not None:
            return found[0]
        if label in context.changed.get(column, ()):
            return "tested"
        return None

    info = {item["column"]: item for item in context.structure["fields"]}
    incoming: Dict[str, List[str]] = {}
    for edge in context.structure["edges"]:
        incoming.setdefault(edge["to"], []).append(edge["from"])
    carriers = {field, *context.reached}
    record_ids = {label: _key_text(final.at[label, context.identifier]) for label in labels} if context.identifier else {}
    flows: Dict[str, Dict[str, Any]] = {}
    for column in context.reached:
        rule = context.rule_by_column.get(column)
        present = column in final.columns
        counts = {"total": len(labels), "reached": 0, "value": 0, "decides": 0, "tested": 0, "delivered": 0, "other": 0, "missing": 0, "missing_reached": 0}
        amount_total = 0.0
        amount_reached = 0.0
        x_amount = 0.0
        per_case: Dict[str, Dict[str, int]] = {}
        groups: Dict[str, Dict[str, Any]] = {}
        for label in labels:
            kind = category(column, label)
            value = final.at[label, column] if present else None
            missing = present and la._is_null(value)
            case, _ = _origin(context, column, label)
            if missing:
                counts["missing"] += 1
            if kind:
                counts[kind] += 1
                counts["reached"] += 1
                if missing:
                    counts["missing_reached"] += 1
            else:
                counts["delivered" if case is None else "other"] += 1
            if present and column in summable:
                amount = _number(value) or 0.0
                amount_total += amount
                if kind:
                    amount_reached += amount
            if kind == "value" and field in summable:
                x_amount += _number(x_values.get(label)) or 0.0
            if case is not None:
                slot = per_case.setdefault(case["id"], {"produced": 0, "reached": 0, "value": 0, "decides": 0, "tested": 0})
                slot["produced"] += 1
                if kind:
                    slot["reached"] += 1
                    slot[kind] += 1
            if categorical:
                x = x_values.get(label)
                key = "∅" if x is None or la._is_null(x) else _key_text(x)
                group = groups.setdefault(key, {"raw": None if key == "∅" else la.json_value(x), "records": 0, "reached": 0, "amount": 0.0, "outcomes": {}})
                group["records"] += 1
                if kind:
                    group["reached"] += 1
                if present:
                    if column in summable:
                        group["amount"] += _number(value) or 0.0
                    else:
                        outcome = "∅" if la._is_null(value) else _key_text(value)
                        slot = group["outcomes"].setdefault(outcome, {"value": None if outcome == "∅" else la.json_value(value), "records": 0})
                        slot["records"] += 1
        by_value = None
        if categorical:
            order = {item["key"]: position for position, item in enumerate(profile)}
            by_value = []
            for key, group in sorted(groups.items(), key=lambda item: order.get(item[0], 999)):
                outcomes = sorted(group["outcomes"].items(), key=lambda item: (-item[1]["records"], item[0]))
                by_value.append({
                    "key": key, "value": group["raw"], "records": group["records"], "reached": group["reached"],
                    "amount": _round(group["amount"]) if present and column in summable else None,
                    "outcomes": [slot for _, slot in outcomes[:MAX_OUTCOMES]] if present and column not in summable else None,
                    "more_outcomes": max(0, len(outcomes) - MAX_OUTCOMES),
                })
        example = None
        for label in context.record_labels:
            witness = context.whatif.get((column, label))
            if witness and label in final_labels:
                example = {"label": str(label), "id": record_ids.get(label), **witness}
                break
        cases = []
        for case in (rule or {}).get("cases", []):
            slot = per_case.get(case["id"], {"produced": 0, "reached": 0, "value": 0, "decides": 0, "tested": 0})
            cases.append({
                "id": case["id"],
                "reads": [{"column": name, "roles": roles} for name, roles in case["roles"].items() if name in carriers and name != column],
                **slot,
            })
        flows[column] = {
            "column": column,
            "level": info[column]["level"],
            "distance": info[column].get("distance"),
            "direct": bool(info[column]["direct"]),
            "end": bool(info[column]["end"]),
            "evidence": info[column].get("evidence", "code"),
            "via": [name for name in incoming.get(column, []) if name in carriers],
            "lines": sorted({line for edge in context.structure["edges"] if edge["to"] == column for line in edge["lines"]}),
            "in_result": present,
            "records": counts,
            "observed": len(context.changed.get(column, ())),
            "amount": {"summable": present and column in summable, "total": _round(amount_total) if present and column in summable else None, "reached": _round(amount_reached) if present and column in summable else None},
            "x_amount": _round(x_amount) if field in summable else None,
            "cases": cases,
            "by_value": by_value,
            "example": example,
        }

    # Routes: the chains of fields from the selected field to the end fields, and the records whose content travels
    # along each of them.
    successors: Dict[str, List[str]] = {}
    for edge in context.structure["edges"]:
        successors.setdefault(edge["from"], []).append(edge["to"])
    candidates: List[List[str]] = []

    def walk(path: List[str]) -> None:
        if len(candidates) >= ROUTE_CANDIDATES:
            return
        following = [name for name in successors.get(path[-1], []) if name not in path]
        if not following:
            if len(path) > 1:
                candidates.append(list(path))
            return
        for name in following:
            walk([*path, name])

    walk([field])
    routes = []
    for route in candidates:
        end = route[-1]
        carried = [label for label in labels if _route_carried(context, route, label)]
        summable_end = end in final.columns and end in summable
        routes.append({
            "fields": route,
            "records": len(carried),
            "amount": _round(sum(_number(final.at[label, end]) or 0.0 for label in carried)) if summable_end else None,
            "examples": [str(label) for label in carried[:3]],
        })
    routes.sort(key=lambda item: (-item["records"], len(item["fields"]), item["fields"]))
    # A field whose final values the content never reaches can still pass it on in an earlier state (its value is
    # overwritten after the next field has used it).
    for column, flow in flows.items():
        flow["passing"] = flow["records"]["reached"] == 0 and any(route["records"] and column in route["fields"][1:-1] for route in routes)
    other_routes = len(routes) - MAX_ROUTES if len(routes) > MAX_ROUTES else 0
    routes = routes[:MAX_ROUTES]

    # Forward trace of the first records.
    record_columns = [name for name in context.scope if name in final.columns]
    items = []
    for label in context.record_labels:
        if label not in final_labels:
            continue
        fields = {}
        for column in context.reached:
            found = reach[(column, label)]
            case, line = _origin(context, column, label)
            key = context.keys.get((line, column), {}).get(label, _MISSING) if case is not None else _MISSING
            fields[column] = {
                "case": case["id"] if case is not None else None,
                "key": None if key is _MISSING else _key_text(key),
                "reach": category(column, label),
                "via": found[1] if found else [],
                "whatif": context.whatif.get((column, label)),
            }
        items.append({
            "label": str(label),
            "id": record_ids.get(label),
            "x": la.json_value(x_values.get(label)),
            "values": {name: la.json_value(final.at[label, name]) for name in record_columns},
            "fields": fields,
        })
    records = {"columns": record_columns, "items": items, "total": len(labels), "identifier": context.identifier}
    nulled: Dict[str, Dict[str, Any]] = {}
    for statement in context.statements:
        for column, entry in (statement.get("nulled") or {}).items():
            if column not in context.scope:
                continue
            nodes = _writes_of(context, statement["line"], column)
            if nodes and all(node.get("operation") == "cast" for node in nodes):
                slot = nulled.setdefault(column, {"count": 0, "examples": [], "line": statement["line"]})
                slot["count"] += entry["count"]
                slot["examples"] = list(dict.fromkeys([*slot["examples"], *entry["examples"]]))[:5]
    selection = []
    for item in context.structure["selection"]:
        statement = next((entry for entry in context.statements if entry["line"] <= item["line"] <= entry["end_line"]), None)
        selection.append({**item, "statement_line": statement["line"] if statement else item["line"], "rows_removed": (statement or {}).get("rows_removed", 0)})
    # Field-to-field links with the role the source plays in the target's rule (a parameter-table key is a key).
    edges = []
    for edge in context.structure["edges"]:
        rule = context.rule_by_column.get(edge["to"])
        roles = list(rule["roles"].get(edge["from"], [])) if rule else []
        edges.append({"from": edge["from"], "to": edge["to"], "lines": edge["lines"], "roles": roles or edge["roles"]})
    reached_end = max((flows[name]["records"]["reached"] for name in context.structure["ends"] if name in flows), default=0)
    job.log(
        "success", "followed",
        f"Content followed in {len(labels)} records: {len(routes) + other_routes} route(s) to the end field(s); the end values of {reached_end} record(s) depend on [{field}]",
        records=len(labels), routes=len(routes) + other_routes, reached=reached_end, field=field,
    )
    for (line, column), missing in context.unmapped.items():
        if column in context.derived:
            job.log("warn", "unmapped", f"[{column}] line {line}: categories without a parameter: {', '.join(f'{key} ({count})' for key, count in missing.items())}", column=column, line=line, values=[f"{key} ({count})" for key, count in missing.items()])
    return {
        "compositions": compositions, "flows": flows, "routes": routes, "other_routes": other_routes, "records": records,
        "nulled": nulled, "numeric": sorted(numeric), "factors": sorted(factors), "categorical": categorical, "profile": profile,
        "selection": selection, "edges": edges,
    }


def _terms(context: _Context, content: Dict[str, Any]) -> List[Dict[str, Any]]:
    final = context.final
    numeric = set(content["numeric"])
    factors = set(content["factors"])
    info = {item["column"]: item for item in context.structure["fields"]}
    columns = {item["name"]: item for item in context.analysis["columns"]}
    inputs = [str(column) for column in context.frame.columns]
    selectors = {name for rule in context.rules for name in rule["selectors"]}
    used_by: Dict[str, List[str]] = {}
    for rule in context.rules:
        for name in rule["roles"]:
            used_by.setdefault(name, [])
            if rule["column"] not in used_by[name]:
                used_by[name].append(rule["column"])
    terms = []
    for name in context.scope:
        rule = context.rule_by_column.get(name)
        if name == context.field:
            role = "start"
        elif name in info:
            role = "next" if info[name]["direct"] else "downstream"
        else:
            role = "context"
        series = final[name] if final is not None and name in final.columns else context.frame[name] if name in context.frame.columns else None
        profile = None
        if series is not None:
            values = [value for value in series.tolist() if not la._is_null(value)]
            distinct = list(dict.fromkeys(_key_text(value) for value in values))
            profile = {
                "non_null": len(values), "total": len(series),
                "distinct": len(distinct), "values": distinct[:15] if name not in numeric or name in factors else [],
                "sum": _round(sum(_number(value) or 0.0 for value in values)) if name in numeric and name not in factors else None,
                "min": _round(min(_number(value) for value in values)) if name in numeric and values else None,
                "max": _round(max(_number(value) for value in values)) if name in numeric and values else None,
                "samples": [la.json_value(value) for value in values[:4]],
            }
        column = columns.get(name)
        terms.append({
            "column": name,
            "role": role,
            "level": info[name]["level"] if name in info else (0 if role == "start" else None),
            "end": bool(info[name]["end"]) if name in info else False,
            "evidence": info[name].get("evidence", "code") if name in info else None,
            "origin": "source" if name in inputs else "derived" if column and column["in_output"] else "intermediate",
            "in_result": final is not None and name in final.columns,
            "classifier": name in selectors and name not in numeric,
            "numeric": name in numeric,
            "factor": name in factors,
            "type": _value_type(series, name in numeric) if series is not None else "empty",
            "rule": rule["id"] if rule else None,
            "inputs": list(rule["roles"]) if rule else [],
            "used_by": used_by.get(name, []),
            "profile": profile,
            "composition": content["compositions"].get(name),
            "nulled": content["nulled"].get(name),
        })
    return terms


# --------------------------------------------------------------------------------------------------
# 3. Describe: business terms, rules and cases in business language
# --------------------------------------------------------------------------------------------------

def _language_rule(language: str) -> str:
    return f"Write every text in {LANGUAGE_NAMES.get(language, 'English')}."


def _case_facts(case: Dict[str, Any]) -> Dict[str, Any]:
    facts: Dict[str, Any] = {
        "case": case["id"],
        "statement_line": case["line"],
        "code": case["code"][:900],
        "applies_when": _plain(case["condition"]) or None,
        "computes": _plain(case["formula"]) or None,
        "only_fills_missing_values": case["fills_missing"],
        "value_columns": case["inputs"],
        "case_selector_columns": case["selectors"],
        "records_whose_final_value_comes_from_this_case": case["records"],
        "records_of_this_case_have": case.get("population") or {},
    }
    if case["lookup"]:
        facts["parameter_table"] = {
            "name": case["lookup"]["name"], "key_column": case["lookup"]["key_column"],
            "entries": [[la.json_value(key), la.json_value(value)] for key, value in case["lookup"]["entries"][:40]],
            "categories_without_parameter": [item["key"] for item in case["keys"] if not item["mapped"]],
            "source_categories_per_key": {item["key"]: item["categories"] for item in case["keys"] if item.get("categories")},
        }
    return facts


_POSITIONS = {
    "start": "the selected field (its content is followed downstream)",
    "next": "next field: its business rule reads the selected field directly",
    "downstream": "field further downstream: its rule reads a field that carries the selected field's content",
    "context": "other input of a downstream rule (not built from the selected field)",
}
_ORIGINS = {"source": "delivered by the source data", "derived": "created by the calculation", "intermediate": "helper column of the calculation (not in the result)"}


def _term_facts(term: Dict[str, Any]) -> Dict[str, Any]:
    profile = term["profile"] or {}
    facts: Dict[str, Any] = {
        "column": term["column"],
        "position": _POSITIONS[term["role"]],
        "end_field": term["end"],
        "origin": _ORIGINS[term["origin"]],
        "type": term["type"],
        "selects_cases": term["classifier"],
        "filled": f"{profile.get('non_null', 0)} of {profile.get('total', 0)} records",
    }
    if profile.get("values"):
        facts["values"] = profile["values"]
    elif profile.get("samples"):
        facts["samples"] = profile["samples"]
    return facts


def _describe_schema(columns: List[str], rules: List[Dict[str, Any]]) -> Dict[str, Any]:
    case_ids = [case["id"] for rule in rules for case in rule["cases"]]
    return {
        "type": "object", "additionalProperties": False, "required": ["domain", "terms", "rules"],
        "properties": {
            "domain": {"type": "string"},
            "terms": {"type": "array", "items": {
                "type": "object", "additionalProperties": False, "required": ["column", "name", "definition"],
                "properties": {"column": _enum(columns), "name": {"type": "string"}, "definition": {"type": "string"}},
            }},
            "rules": {"type": "array", "items": {
                "type": "object", "additionalProperties": False, "required": ["rule", "name", "statement", "cases"],
                "properties": {
                    "rule": _enum([rule["id"] for rule in rules]),
                    "name": {"type": "string"},
                    "statement": {"type": "string"},
                    "cases": {"type": "array", "items": {
                        "type": "object", "additionalProperties": False, "required": ["case", "label", "description"],
                        "properties": {"case": _enum(case_ids), "label": {"type": "string"}, "description": {"type": "string"}},
                    }},
                },
            }},
        },
    }


_DESCRIBE_SYSTEM = (
    "You are a business expert at a German bank documenting the CONTENT LINEAGE (business lineage) of a field of a "
    "calculation: what happens to its content downstream — the fields it flows into, the business rules that use it and the "
    "cases those rules distinguish — in the language of the business, not of the code.\n"
    "You receive the deterministic facts: the columns involved (the selected field, the fields downstream of it and the other "
    "inputs of their rules) and, per downstream field, one rule whose cases are the statements of the code that write it (with "
    "condition, formula, parameter tables and how many records each case produced).\n"
    "Return:\n"
    "- domain: the business domain in a few words (for example 'Credit risk – standardised approach') (only when asked);\n"
    "- terms: for every column listed, the business term a banking expert would use (name; include a common abbreviation in "
    "parentheses where one exists, e.g. 'Exposure value (EAD)') and a definition of 1–2 sentences of what the term MEANS, not how "
    "the code computes it; do not list its values (they are shown next to the definition);\n"
    "- rules: for every rule, a short business name, a statement of the rule in one or two sentences that a business analyst can "
    "verify, and every case (label of at most six words; description of when the case applies and what it yields).\n"
    "Rules for writing:\n"
    "- Refer to business terms by their business names, never by column names in brackets or code.\n"
    "- A case that only fills missing values means: the value delivered by the source is kept; the case supplies it only when "
    "it is missing. Say so.\n"
    "- A case of a parameter table: describe the parameter per category exactly as in the table (factors may be written as "
    "percentages, 0.2 = 20 %).\n"
    "- Use only numbers that appear in the facts of that rule (no record counts, no amounts, no article numbers).\n"
    "- Write category values of the data (product types, classes, codes, keys and values of parameter tables) exactly as they appear in the data, in single quotes (e.g. 'Limit', 'OffBalance'); never translate or rename them.\n"
    "- Do not mention rule or case ids (such as R1 or R1.3); name rules and cases in business language.\n"
    "- Do not cite regulations or articles.\n"
    "- Return every listed column, rule and case exactly once, with the exact ids."
)


def _allowed_numbers(rule: Dict[str, Any]) -> Set[float]:
    allowed = {0.0, 1.0, 100.0}
    for case in rule["cases"]:
        for value in case["constants"]:
            allowed.add(round(value, 9))
            allowed.add(round(value * 100, 9))
            allowed.add(round(value / 100, 9))
        allowed.add(float(case["line"]))
    return allowed


def _ungrounded_numbers(text: str, allowed: Set[float]) -> List[str]:
    bad = []
    text = re.sub(r"\bR\d+(?:\.\d+)*\b", " ", text or "")  # a rule or case id is not a number
    for readings in _numbers_in(text):
        if not any(round(value, 9) in allowed for value in readings):
            bad.append("/".join(_key_text(value) for value in sorted(readings)))
    return bad


def _describe_batch(job: _Job, llm: _LLM, context: _Context, terms: List[Dict[str, Any]], rules: List[Dict[str, Any]], first: bool) -> Tuple[Dict[str, Any], List[Dict[str, Any]], bool]:
    """(answer, remaining grounding issues, corrected) for one batch of terms and rules."""
    columns = [term["column"] for term in terms]
    rule_ids = [rule["id"] for rule in rules]
    facts = {
        "write_domain": first,
        "selected_field": context.field,
        "end_fields": context.structure.get("ends", []),
        "columns": [_term_facts(term) for term in terms],
        "rules": [
            {
                "rule": rule["id"],
                "derives_column": rule["column"],
                "column_also_delivered_by_source": rule["column"] in context.frame.columns,
                "cases": [_case_facts(case) for case in rule["cases"]],
            }
            for rule in rules
        ],
    }
    messages = [
        {"role": "system", "content": _DESCRIBE_SYSTEM + "\n" + _language_rule(context.language)},
        {"role": "user", "content": json.dumps(facts, ensure_ascii=False, default=str)},
    ]
    schema = _describe_schema(columns, rules)

    def issues_of(answer: Dict[str, Any]) -> List[Dict[str, Any]]:
        issues: List[Dict[str, Any]] = []
        named = {item.get("column") for item in answer.get("terms", []) if str(item.get("name") or "").strip()}
        for column in columns:
            if column not in named:
                issues.append({"code": "term_missing", "params": {"column": column}, "text": f"Column '{column}' has no business name."})
        by_rule = {item.get("rule"): item for item in answer.get("rules", [])}
        for rule in rules:
            item = by_rule.get(rule["id"])
            if not item or not str(item.get("statement") or "").strip():
                issues.append({"code": "rule_missing", "params": {"rule": rule["id"]}, "text": f"Rule {rule['id']} has no statement."})
                continue
            described = {case.get("case") for case in item.get("cases", []) if str(case.get("description") or "").strip()}
            for case in rule["cases"]:
                if case["id"] not in described:
                    issues.append({"code": "case_missing", "params": {"case": case["id"]}, "text": f"Case {case['id']} is not described."})
            for case in item.get("cases", []):
                if case.get("case") not in {entry["id"] for entry in rule["cases"]}:
                    issues.append({"code": "case_unknown", "params": {"case": case.get("case")}, "text": f"Case {case.get('case')} does not belong to rule {rule['id']}."})
            allowed = _allowed_numbers(rule)
            texts = [("statement", item.get("statement", ""))] + [
                (f"case {case.get('case')}", f"{case.get('label', '')} {case.get('description', '')}") for case in item.get("cases", [])
            ]
            for where, text in texts:
                bad = _ungrounded_numbers(str(text), allowed)
                if bad:
                    issues.append({
                        "code": "number_ungrounded", "params": {"rule": rule["id"], "where": where, "numbers": bad[:5]},
                        "text": f"Rule {rule['id']} ({where}) states number(s) {', '.join(bad[:5])} that are not parameters of this rule.",
                    })
        return issues

    answer = llm.json(messages, "content_model", schema)
    issues = issues_of(answer)
    corrected = False
    if issues:
        retry = messages + [
            {"role": "assistant", "content": json.dumps(answer, ensure_ascii=False)},
            {"role": "user", "content": "Correct these problems and return the complete answer again:\n- " + "\n- ".join(item["text"] for item in issues[:25])},
        ]
        second = llm.json(retry, "content_model", schema)
        second_issues = issues_of(second)
        if len(second_issues) <= len(issues):
            answer, issues, corrected = second, second_issues, True
    return answer, issues, corrected


def _describe(job: _Job, llm: _LLM, context: _Context, terms: List[Dict[str, Any]]) -> Dict[str, Any]:
    job.stage("describe")
    rules = context.rules
    batches = [rules[index:index + RULE_BATCH] for index in range(0, len(rules), RULE_BATCH)] or [[]]
    assigned: Set[str] = set()
    plan: List[Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]] = []
    by_column = {term["column"]: term for term in terms}
    for batch in batches:
        names: List[str] = []
        for rule in batch:
            for name in [rule["column"], *rule["inputs"], *rule["selectors"]]:
                if name in by_column and name not in assigned and name not in names:
                    names.append(name)
        assigned |= set(names)
        plan.append(([by_column[name] for name in names], batch))
    leftover = [term for term in terms if term["column"] not in assigned]
    if leftover:
        plan[-1] = (plan[-1][0] + leftover, plan[-1][1])
    job.progress(rules_total=len(rules), rules_described=0)
    job.log("info", "describing", f"Describing {len(terms)} business terms and {len(rules)} business rules in {len(plan)} batch(es)", terms=len(terms), rules=len(rules), batches=len(plan))
    texts: Dict[str, Any] = {"summary": "", "domain": "", "terms": {}, "rules": {}, "usage": {}}
    all_issues: List[Dict[str, Any]] = []
    corrected_any = False
    lock = threading.Lock()

    def run(index: int, batch_terms: List[Dict[str, Any]], batch_rules: List[Dict[str, Any]]) -> None:
        nonlocal corrected_any
        try:
            answer, issues, corrected = _describe_batch(job, llm, context, batch_terms, batch_rules, index == 0)
        except _Cancelled:
            raise
        except Exception as exc:
            job.log("warn", "llm_error", f"Describing batch {index + 1} failed: {str(exc)[:200]}", step="describe", message=str(exc)[:200])
            with lock:
                all_issues.extend({"code": "rule_missing", "params": {"rule": rule["id"]}, "text": f"Rule {rule['id']} could not be described."} for rule in batch_rules)
            return
        with lock:
            corrected_any = corrected_any or corrected
            if index == 0:
                texts["domain"] = str(answer.get("domain") or "").strip()
            allowed_columns = {term["column"] for term in batch_terms}
            for item in answer.get("terms", []):
                if item.get("column") in allowed_columns and str(item.get("name") or "").strip():
                    texts["terms"][item["column"]] = {"name": str(item["name"]).strip(), "definition": str(item.get("definition") or "").strip()}
            for rule in batch_rules:
                item = next((entry for entry in answer.get("rules", []) if entry.get("rule") == rule["id"]), None)
                if not item:
                    continue
                cases = {}
                for case in item.get("cases", []):
                    if case.get("case") in {entry["id"] for entry in rule["cases"]}:
                        cases[case["case"]] = {"label": str(case.get("label") or "").strip(), "description": str(case.get("description") or "").strip()}
                texts["rules"][rule["id"]] = {"name": str(item.get("name") or "").strip(), "statement": str(item.get("statement") or "").strip(), "cases": cases}
            all_issues.extend(issues)
        job.increment("rules_described", len(batch_rules))
        job.log("success", "described", f"Batch {index + 1}: {len(batch_terms)} terms and {len(batch_rules)} rules described" + (" after a correction round" if corrected else ""), batch=index + 1, terms=len(batch_terms), rules=len(batch_rules), corrected=corrected)

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        for future in [pool.submit(run, index, batch_terms, batch_rules) for index, (batch_terms, batch_rules) in enumerate(plan)]:
            future.result()
    for issue in all_issues:
        job.log("warn", "grounding_issue", issue["text"], code=issue["code"], **{key: value for key, value in issue["params"].items() if key != "code"})
    return {"texts": texts, "grounding": {"issues": all_issues, "corrected": corrected_any, "status": "verified" if not all_issues else "partial"}}


_USAGE_SYSTEM = (
    "You document the CONTENT LINEAGE of one field of a bank's calculation in the forward direction: what happens to the "
    "content of the selected field downstream — which fields it flows into next, which business rules use it and how, and "
    "which fields at the end of the calculation it finally reaches.\n"
    "You receive the selected field (business name, definition and, for a field with few values, its values) and every "
    "downstream field with its business name and definition, its position (next_field = its rule reads the selected field "
    "directly; otherwise 'via' names the fields through which the content arrives), its business rule with every case, and "
    "per case what it reads of the selected field or of the fields on the way, in which role: value = used to compute the "
    "result, key = key of a parameter table, condition = decides whether the case applies, group = values are aggregated per "
    "group, join = key that joins other data. A case that reads none of them produces its records' values without the selected "
    "field. 'found_by_impact_test' marks a field the code analysis cannot link: re-running the code with other values of the "
    "selected field changed it, but how is not visible in the code (for example values computed in a loop over the records).\n"
    "Return:\n"
    "- summary: 2–3 sentences in business terms on what the selected field determines downstream: what it decides or feeds "
    "directly, and which fields at the end of the calculation it finally reaches;\n"
    "- fields: for every downstream field, usage: one or two sentences on how the content of the selected field determines this "
    "field — through which business rule and cases, in which role, and with which parameter per category when a parameter table "
    "is keyed by it; for a field reached via other fields, name the field it arrives through. When cases that do not read it "
    "produce some records' values, say that for those records the value comes from that case. For a field found by the impact "
    "test, say only that its values change with the selected field, as the test shows, and that the code does not show how.\n"
    "Rules for writing:\n"
    "- Refer to fields by their business names, never by column names in brackets or code.\n"
    "- Write category values of the data (product types, classes, codes, keys and values of parameter tables) exactly as they appear in the data, in single quotes (e.g. 'Limit', 'OffBalance'); never translate or rename them.\n"
    "- Do not mention rule or case ids (such as R1 or R1.3); name rules and cases in business language.\n"
    "- Use only numbers that appear as parameters in the facts (factors may be written as percentages, 0.2 = 20 %); no record "
    "counts, no amounts, no article numbers; do not cite regulations.\n"
    "- Return every listed field exactly once, with its exact column."
)


def _usage_facts(context: _Context, content: Dict[str, Any], texts: Dict[str, Any]) -> Dict[str, Any]:
    def name(column: str) -> str:
        return texts["terms"].get(column, {}).get("name") or column

    field = context.field
    carriers = {field, *context.reached}
    selected: Dict[str, Any] = {"column": field, "business_term": name(field), "definition": texts["terms"].get(field, {}).get("definition", "")}
    if content["categorical"]:
        selected["values"] = [item["key"] for item in content["profile"][:15]]
    fields = []
    for column in context.reached:
        flow = content["flows"][column]
        rule = context.rule_by_column.get(column)
        entry: Dict[str, Any] = {
            "column": column,
            "business_term": name(column),
            "definition": texts["terms"].get(column, {}).get("definition", ""),
            "next_field": flow["direct"],
            "via": [name(other) for other in flow["via"] if other != field],
            "end_field": flow["end"],
            "found_by_impact_test": flow["evidence"] == "impact",
        }
        if rule:
            rule_text = texts["rules"].get(rule["id"], {})
            entry["rule"] = {"name": rule_text.get("name", ""), "statement": rule_text.get("statement", "")}
            entry["cases"] = []
            for case in rule["cases"]:
                case_text = rule_text.get("cases", {}).get(case["id"], {})
                item: Dict[str, Any] = {
                    "case": case["id"], "label": case_text.get("label", ""), "description": case_text.get("description", ""),
                    "reads": [{"field": name(other), "role": "/".join(roles)} for other, roles in case["roles"].items() if other in carriers and other != column],
                }
                if case["lookup"] and case["lookup"]["key_column"] in carriers:
                    item["parameter_table"] = {"key": name(case["lookup"]["key_column"]), "entries": [[la.json_value(key), la.json_value(value)] for key, value in case["lookup"]["entries"][:40]]}
                entry["cases"].append(item)
        fields.append(entry)
    return {"selected_field": selected, "downstream_fields": fields}


def _usage(job: _Job, llm: _LLM, context: _Context, content: Dict[str, Any], texts: Dict[str, Any]) -> Dict[str, Any]:
    """How the selected field reaches every downstream field, and a summary, in business language (closed set of fields,
    numbers checked against the parameters of the rules on the way, one correction round)."""
    columns = list(context.reached)
    schema = {
        "type": "object", "additionalProperties": False, "required": ["summary", "fields"],
        "properties": {
            "summary": {"type": "string"},
            "fields": {"type": "array", "items": {
                "type": "object", "additionalProperties": False, "required": ["column", "usage"],
                "properties": {"column": _enum(columns), "usage": {"type": "string"}},
            }},
        },
    }
    allowed: Set[float] = {0.0, 1.0, 100.0}
    for rule in context.rules:
        allowed |= _allowed_numbers(rule)
    messages = [
        {"role": "system", "content": _USAGE_SYSTEM + "\n" + _language_rule(context.language)},
        {"role": "user", "content": json.dumps(_usage_facts(context, content, texts), ensure_ascii=False, default=str)},
    ]

    def issues_of(answer: Dict[str, Any]) -> List[Dict[str, Any]]:
        issues: List[Dict[str, Any]] = []
        explained = {item.get("column"): str(item.get("usage") or "").strip() for item in answer.get("fields", [])}
        if not str(answer.get("summary") or "").strip():
            issues.append({"code": "usage_missing", "params": {"column": context.field}, "text": "The summary is missing."})
        for column in columns:
            if not explained.get(column):
                issues.append({"code": "usage_missing", "params": {"column": column}, "text": f"Field '{column}' has no usage."})
        for where, text in [("summary", str(answer.get("summary") or "")), *[(f"usage of {column}", text) for column, text in explained.items()]]:
            bad = _ungrounded_numbers(text, allowed)
            if bad:
                issues.append({
                    "code": "number_ungrounded", "params": {"where": where, "numbers": bad[:5]},
                    "text": f"The {where} states number(s) {', '.join(bad[:5])} that are not parameters of the rules.",
                })
        return issues

    job.log("info", "explaining", f"Explaining how [{context.field}] reaches {len(columns)} downstream field(s)", field=context.field, count=len(columns))
    try:
        answer = llm.json(messages, "forward_usage", schema)
        issues = issues_of(answer)
        corrected = False
        if issues:
            retry = messages + [
                {"role": "assistant", "content": json.dumps(answer, ensure_ascii=False)},
                {"role": "user", "content": "Correct these problems and return the complete answer again:\n- " + "\n- ".join(item["text"] for item in issues[:25])},
            ]
            second = llm.json(retry, "forward_usage", schema)
            second_issues = issues_of(second)
            if len(second_issues) <= len(issues):
                answer, issues, corrected = second, second_issues, True
    except _Cancelled:
        raise
    except Exception as exc:
        job.log("warn", "llm_error", f"Explaining the downstream fields failed: {str(exc)[:200]}", step="usage", message=str(exc)[:200])
        return {"summary": "", "usage": {}, "issues": [{"code": "usage_missing", "params": {"column": column}, "text": f"Field '{column}' has no usage."} for column in columns], "corrected": False}
    usage = {item["column"]: str(item.get("usage") or "").strip() for item in answer.get("fields", []) if item.get("column") in columns and str(item.get("usage") or "").strip()}
    for issue in issues:
        job.log("warn", "grounding_issue", issue["text"], code=issue["code"], **{key: value for key, value in issue["params"].items() if key != "code"})
    job.log("success", "explained", f"How [{context.field}] reaches each downstream field explained ({len(usage)} of {len(columns)})" + (" after a correction round" if corrected else ""), field=context.field, done=len(usage), total=len(columns), corrected=corrected)
    return {"summary": str(answer.get("summary") or "").strip(), "usage": usage, "issues": issues, "corrected": corrected}


# --------------------------------------------------------------------------------------------------
# 4. Translate
# --------------------------------------------------------------------------------------------------

def _text_items(texts: Dict[str, Any]) -> List[Tuple[str, str]]:
    items: List[Tuple[str, str]] = [("summary", texts["summary"]), ("domain", texts["domain"])]
    for column, term in texts["terms"].items():
        items += [(f"term␟{column}␟name", term["name"]), (f"term␟{column}␟definition", term["definition"])]
    for column, text in texts.get("usage", {}).items():
        items.append((f"usage␟{column}", text))
    for rule_id, rule in texts["rules"].items():
        items += [(f"rule␟{rule_id}␟name", rule["name"]), (f"rule␟{rule_id}␟statement", rule["statement"])]
        for case_id, case in rule["cases"].items():
            items += [(f"case␟{rule_id}␟{case_id}␟label", case["label"]), (f"case␟{rule_id}␟{case_id}␟description", case["description"])]
    return [(key, text) for key, text in items if str(text or "").strip()]


def _texts_from(items: Dict[str, str], original: Dict[str, Any]) -> Dict[str, Any]:
    """The structured texts of one language, built from translated items (missing items keep the original)."""
    texts = copy.deepcopy(original)
    for key, text in items.items():
        parts = key.split("␟")
        if parts[0] in ("summary", "domain"):
            texts[parts[0]] = text
        elif parts[0] == "term" and parts[1] in texts["terms"]:
            texts["terms"][parts[1]][parts[2]] = text
        elif parts[0] == "usage" and parts[1] in texts.get("usage", {}):
            texts["usage"][parts[1]] = text
        elif parts[0] == "rule" and parts[1] in texts["rules"]:
            texts["rules"][parts[1]][parts[2]] = text
        elif parts[0] == "case" and parts[2] in texts["rules"].get(parts[1], {}).get("cases", {}):
            texts["rules"][parts[1]]["cases"][parts[2]][parts[3]] = text
    return texts


def _data_values(context: _Context, terms: List[Dict[str, Any]]) -> List[str]:
    """Text values of the data that AI texts may name (parameter-table keys and values, categories of the fields in
    scope): translations must keep them exactly."""
    found: Set[str] = set()

    def add(value: Any) -> None:
        if isinstance(value, str) and len(value.strip()) >= 3 and re.search(r"[A-Za-z]", value):
            found.add(value.strip())

    for rule in context.rules:
        for case in rule["cases"]:
            for key, value in (case["lookup"] or {}).get("entries", []):
                add(key)
                add(value)
    for term in terms:
        if term["profile"] and not term["numeric"] and term["profile"]["distinct"] <= MAX_CATEGORIES:
            for value in term["profile"]["values"]:
                add(value)
    return sorted(found)


def _digits(text: str) -> List[str]:
    return sorted(re.sub(r"[^\d]", "", token) for token in re.findall(r"\d[\d.,' ]*\d|\d", text or ""))


def _translate(job: _Job, llm: _LLM, items: List[Tuple[str, str]], source: str, target: str, names: List[str], values: Optional[List[str]] = None) -> Tuple[Dict[str, str], int]:
    """Translated texts by key and the number of texts that kept their original wording.

    A translation must keep the numbers of its source text, the technical column names and the values of the data it
    contains (checked deterministically); a batch with rejected texts goes back once with the reasons."""
    protected = sorted({name for name in names if len(name) > 2}, key=len, reverse=True)
    data_values = [(value, re.compile(rf"(?<![\w]){re.escape(value)}(?![\w])")) for value in sorted(set(values or []), key=len, reverse=True)]

    def problems(original: str, translated: Any) -> Optional[str]:
        if not isinstance(translated, str) or bool(original.strip()) != bool(translated.strip()):
            return "empty or missing"
        if _digits(original) != _digits(translated):
            return "numbers changed"
        for name in protected:
            if f"[{name}]" in original and f"[{name}]" not in translated:
                return f"[{name}] dropped"
        for value, pattern in data_values:
            if pattern.search(original) and not pattern.search(translated):
                return f"data value '{value}' changed (keep it exactly)"
        return None

    batches = [items[index:index + TRANSLATE_BATCH] for index in range(0, len(items), TRANSLATE_BATCH)]
    result: Dict[str, str] = {}
    kept = 0
    lock = threading.Lock()

    def run(batch: List[Tuple[str, str]]) -> None:
        nonlocal kept
        job.check()
        keys = [key for key, _ in batch]
        original = dict(batch)
        schema = {
            "type": "object", "additionalProperties": False, "required": ["items"],
            "properties": {"items": {"type": "array", "items": {
                "type": "object", "additionalProperties": False, "required": ["key", "text"],
                "properties": {"key": _enum(keys), "text": {"type": "string"}},
            }}},
        }
        messages = [
            {"role": "system", "content": (
                f"You translate the content-lineage documentation of a bank from {LANGUAGE_NAMES[source]} into {LANGUAGE_NAMES[target]}: "
                "business terms, definitions, business rules and how a field reaches other fields. Translate every text faithfully "
                "in the professional register of German banking (use the established terms, e.g. 'Risikopositionswert' for "
                "'exposure value', 'Kreditumrechnungsfaktor' for 'credit conversion factor', 'risikogewichteter Positionsbetrag' for "
                "'risk-weighted exposure amount', and their English equivalents in the other direction). Keep abbreviations (EAD, RWA, "
                "CCF), every number exactly as written, values in quotes and "
                "anything in [brackets] unchanged. Values of the data (product types, classes, codes such as 'Loan' or 'OffBalance') "
                "are content: keep them exactly as written and never translate them. Return every key exactly once."
            )},
            {"role": "user", "content": json.dumps({"items": [{"key": key, "text": text} for key, text in batch]}, ensure_ascii=False)},
        ]
        try:
            answer = llm.json(messages, "translation", schema)
            texts = {item.get("key"): item.get("text") for item in answer.get("items", [])}
            rejected = {key: reason for key in keys if (reason := problems(original[key], texts.get(key)))}
            if rejected:
                retry = messages + [
                    {"role": "assistant", "content": json.dumps(answer, ensure_ascii=False)},
                    {"role": "user", "content": "Return the complete answer again and fix these items (keep numbers and [names] exactly): " + "; ".join(f"{key}: {reason}" for key, reason in list(rejected.items())[:30])},
                ]
                second = llm.json(retry, "translation", schema)
                second_texts = {item.get("key"): item.get("text") for item in second.get("items", [])}
                for key in list(rejected):
                    if not problems(original[key], second_texts.get(key)):
                        texts[key] = second_texts[key]
                        rejected.pop(key)
        except _Cancelled:
            raise
        except Exception as exc:
            job.log("warn", "llm_error", f"Translation into {LANGUAGE_NAMES[target]} failed: {str(exc)[:200]}", step="translate", message=str(exc)[:200])
            with lock:
                kept += len(batch)
            return
        with lock:
            for key in keys:
                if key in rejected:
                    kept += 1
                else:
                    result[key] = str(texts[key]).strip()

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        for future in [pool.submit(run, batch) for batch in batches]:
            future.result()
    return result, kept


# --------------------------------------------------------------------------------------------------
# Checks, findings, summary
# --------------------------------------------------------------------------------------------------

def _checks(context: _Context, content: Dict[str, Any], described: Dict[str, Any], translation: Dict[str, Any]) -> List[Dict[str, Any]]:
    checks: List[Dict[str, Any]] = []
    replay = context.replay
    impact = context.impact
    flows = content["flows"]

    def add(check_id: str, status: str, **params: Any) -> None:
        checks.append({"id": check_id, "status": status, "params": params})

    if replay.get("ok"):
        add("replay", "pass", cells=replay.get("cells_compared", 0))
    else:
        add("replay", "warn", mismatches=replay.get("mismatches", 0), cells=replay.get("cells_compared", 0), columns_match=replay.get("columns_match", True))
    # Completeness: every field whose values change with the selected field is among the downstream fields the code
    # analysis found; a field only the test finds is added, but the check shows that the code analysis missed it.
    code_fields = [column for column in context.reached if flows[column]["evidence"] == "code"]
    if impact.get("enabled"):
        problems = bool(impact["unexpected"]) or impact["budget_hit"] or bool(impact["errors"])
        add(
            "impact", "warn" if problems else "pass",
            runs=impact["tested"], rows=impact["rows"], classes=impact["classes"], fields=len(code_fields), unexpected=impact["unexpected"],
            budget_hit=impact["budget_hit"], errors=len(impact["errors"]), all_values=impact["all_values"],
        )
        # Every field the code links changes in the test, or passes the content on in an earlier state of its value.
        unobserved = [column for column in code_fields if not flows[column]["observed"] and not flows[column]["passing"]]
        add("observed", "pass" if not unobserved else "warn", fields=len(code_fields), observed=len(code_fields) - len(unobserved), unobserved=unobserved[:10])
    else:
        add("impact", "skipped", reason=impact.get("reason") or "disabled")
        add("observed", "skipped", reason=impact.get("reason") or "disabled")
    texts = described["texts"]
    unnamed = [column for column in context.scope if column not in texts["terms"]]
    add("terms", "pass" if not unnamed else "warn", total=len(context.scope), named=len(context.scope) - len(unnamed), missing=unnamed[:10])
    cases = [case["id"] for rule in context.rules for case in rule["cases"]]
    undescribed = [case_id for case_id in cases if not texts["rules"].get(context.cases[case_id]["rule"], {}).get("cases", {}).get(case_id, {}).get("description")]
    add("cases", "pass" if not undescribed else "warn", total=len(cases), described=len(cases) - len(undescribed), missing=undescribed[:10])
    unexplained = [column for column in context.reached if not texts.get("usage", {}).get(column)]
    add("usage", "pass" if not unexplained else "warn", total=len(context.reached), explained=len(context.reached) - len(unexplained), missing=unexplained[:10])
    numbers = [issue for issue in described["grounding"]["issues"] if issue["code"] == "number_ungrounded"]
    add("numbers", "pass" if not numbers else "warn", issues=len(numbers), corrected=described["grounding"]["corrected"])
    unmapped = sum(count for (line, column), values in context.unmapped.items() if column in context.derived for count in values.values())
    add("unmapped", "pass" if not unmapped else "warn", records=unmapped)
    missing = {column: flows[column]["records"]["missing"] for column in context.structure.get("ends", []) if column in flows and flows[column]["records"]["missing"]}
    add("complete", "pass" if not missing else "warn", missing=missing)
    if translation.get("languages"):
        add("translation", "pass" if not translation.get("kept") else "warn", languages=translation["languages"], kept=translation.get("kept", 0))
    return checks


def _findings(context: _Context, content: Dict[str, Any]) -> List[Dict[str, Any]]:
    findings: List[Dict[str, Any]] = []
    flows = content["flows"]

    def add(code: str, severity: str, **params: Any) -> None:
        findings.append({"id": "", "code": code, "severity": severity, "params": params})

    for column in context.structure.get("ends", []):
        records = flows.get(column, {}).get("records")
        if records and records["missing"]:
            add("end_missing", "critical", column=column, count=records["missing"], total=records["total"], reached=records["missing_reached"])
    for (line, column), values in context.unmapped.items():
        if column in context.derived:
            rule = context.rule_by_column[column]
            case = context.case_by_write.get((line, column))
            add("unmapped_category", "warning", rule=rule["id"], column=column, case=case["id"] if case else None, key_column=((case or {}).get("lookup") or {}).get("key_column"), values=[{"value": key, "rows": count} for key, count in values.items()])
    for column, entry in content["nulled"].items():
        add("source_not_numeric", "warning", column=column, count=entry["count"], examples=entry["examples"], line=entry["line"])
    for column in context.reached:
        if flows[column]["evidence"] == "impact":
            add("impact_only", "warning", column=column, records=flows[column]["records"]["reached"])
    for error in context.impact.get("errors", []):
        add("impact_error", "info", line=error["line"], message=error["message"][:200])
    for column in context.reached:
        records = flows[column]["records"]
        if records["delivered"] and records["reached"]:
            add("delivered_unaffected", "info", column=column, delivered=records["delivered"], reached=records["reached"])
    for item in content["selection"]:
        add("selection", "info", line=item["statement_line"], rows=item["rows_removed"], code=(item.get("code") or "")[:200])
    order = {"critical": 0, "warning": 1, "info": 2}
    findings.sort(key=lambda item: order[item["severity"]])
    for index, item in enumerate(findings, start=1):
        item["id"] = f"F{index:02d}"
    return findings


def _summary(context: _Context, content: Dict[str, Any], findings: List[Dict[str, Any]], checks: List[Dict[str, Any]]) -> Dict[str, Any]:
    flows = content["flows"]
    ends = context.structure.get("ends", [])
    impact = context.impact
    return {
        "field": context.field,
        "next": sum(1 for column in context.reached if flows[column]["direct"]),
        "fields": len(context.reached),
        "ends": len(ends),
        "end_names": ends,
        "terms": len(context.scope),
        "context": len(context.context_fields),
        "rules": len(context.rules),
        "cases": sum(len(rule["cases"]) for rule in context.rules),
        "records": len(context.final) if context.final is not None else 0,
        "reached": max((flows[column]["records"]["reached"] for column in ends if column in flows), default=0),
        "routes": len(content["routes"]) + content["other_routes"],
        "impact": {"enabled": bool(impact.get("enabled")), "confirmed": impact.get("confirmed", 0), "unexpected": len(impact.get("unexpected", []))},
        "findings": {severity: sum(1 for item in findings if item["severity"] == severity) for severity in ("critical", "warning", "info")},
        "checks": {status: sum(1 for item in checks if item["status"] == status) for status in ("pass", "warn", "fail", "skipped")},
        "replay_ok": bool(context.replay.get("ok")),
    }


# --------------------------------------------------------------------------------------------------
# Pipeline
# --------------------------------------------------------------------------------------------------

def _run_job(job: _Job) -> None:
    time.sleep(0.2)  # let the start request's response go out first
    context = _Context()
    context.language = job.state["language"]
    try:
        _load(job, context)
        job.check()
        _analyse(job, context)
        _trace(job, context)
        job.save(force=True)
        job.check()
        _compose(job, context, job.state["settings"]["field"])
        _impact(job, context)
        _admit_unexpected(job, context)
        _require_downstream(context)
        content = _follow(job, context)
        terms = _terms(context, content)
        with job.lock:
            job.state["settings"]["fields_reached"] = list(context.reached)
        job.save(force=True)
        job.check()

        llm = _LLM(job)
        job.set_model(llm.model)
        job.log("info", "model", f"Model: {llm.model} · temperature 0 · JSON-schema output", model=llm.model)
        described = _describe(job, llm, context, terms)
        job.check()
        usage = _usage(job, llm, context, content, described["texts"])
        described["texts"]["summary"] = usage["summary"]
        described["texts"]["usage"] = usage["usage"]
        grounding = described["grounding"]
        grounding["issues"] = [*grounding["issues"], *usage["issues"]]
        grounding["corrected"] = grounding["corrected"] or usage["corrected"]
        grounding["status"] = "verified" if not grounding["issues"] else "partial"
        job.save(force=True)
        job.check()

        # Every AI text in English and German.
        job.stage("translate")
        original = dict(described["texts"])
        items = _text_items(original)
        texts_by_language = {context.language: _texts_from({}, original)}
        translation = {"languages": [], "kept": 0}
        for target in [code for code in TEXT_LANGUAGES if code != context.language]:
            job.log("info", "translating", f"Translating {len(items)} AI texts into {LANGUAGE_NAMES[target]}", language=target.upper(), count=len(items))
            translated, kept = _translate(job, llm, items, context.language, target, list(context.scope), _data_values(context, terms))
            texts_by_language[target] = _texts_from(translated, original)
            translation["languages"].append(target)
            translation["kept"] += kept
            job.log("success", "translated", f"AI texts available in {LANGUAGE_NAMES[target]} ({len(translated)} texts; {kept} kept in the original wording)", language=target.upper(), count=len(translated), kept=kept)
        job.check()

        job.stage("assemble")
        checks = _checks(context, content, described, translation)
        findings = _findings(context, content)
        summary = _summary(context, content, findings, checks)
        origin = next((term["origin"] for term in terms if term["column"] == context.field), "source")
        result = {
            "mode": RESULT_MODE,
            "code": {"filename": job.state["source"]["code_filename"], "lines": context.lines, "line_count": len(context.lines)},
            "scope": {
                "field": context.field,
                "origin": origin,
                "next": [column for column in context.reached if content["flows"][column]["direct"]],
                "fields": list(context.reached),
                "ends": list(context.structure.get("ends", [])),
                "context": list(context.context_fields),
                "identifier": context.identifier,
                "records": summary["records"],
                "categorical": content["categorical"],
                "values": content["profile"],
                "edges": content["edges"],
                "selection": content["selection"],
            },
            "terms": terms,
            "rules": context.rules,
            "flows": content["flows"],
            "routes": content["routes"],
            "other_routes": content["other_routes"],
            "records": content["records"],
            "impact": context.impact,
            "filters": context.filters,
            "replay": context.replay,
            "ai": {
                "language": context.language,
                "texts": texts_by_language,
                "grounding": described["grounding"],
                "translation": translation,
            },
            "checks": checks,
            "findings": findings,
        }
        with job.lock:
            job.state["result"] = result
            job.state["summary"] = summary
            job.state["status"] = "completed"
        job.log(
            "success", "completed",
            f"Content lineage complete: [{context.field}] reaches {summary['fields']} downstream field(s), {summary['ends']} of them end field(s); {summary['rules']} rules",
            field=context.field, fields=summary["fields"], ends=summary["ends"], rules=summary["rules"],
        )
    except _Cancelled:
        with job.lock:
            job.state["status"] = "cancelled"
        job.log("warn", "cancelled", "Run cancelled.")
    except HTTPException as exc:
        with job.lock:
            job.state.update(status="failed", error=str(exc.detail))
        job.log("error", "failed", str(exc.detail), message=str(exc.detail))
    except Exception as exc:
        traceback.print_exc()
        with job.lock:
            job.state.update(status="failed", error=str(exc)[:500])
        job.log("error", "failed", f"Run failed: {str(exc)[:300]}", message=str(exc)[:300])
    finally:
        with job.lock:
            job.state["finished_at"] = _now()
            job.state["duration_ms"] = int((time.monotonic() - job.started) * 1000)
            job.state["stage"] = "finished"
        job.save(force=True)
        with _JOBS_LOCK:
            _JOBS.pop(job.state["id"], None)


# --------------------------------------------------------------------------------------------------
# API
# --------------------------------------------------------------------------------------------------

class RunRequest(BaseModel):
    execution_id: str
    field: str = ""
    language: str = "en"


@router.get("/sources")
def list_sources():
    executions = tla.list_sources()
    return {
        "clusters": executions["clusters"],
        "model": _configured_model(),
        "llm_configured": tla._llm_configured(),
    }


@router.get("/preview")
def preview(execution_id: str):
    """Every field of an execution and what its content flows into: the same computation as an analysis (code
    analysis on column versions), completed by a quick impact test of every field on the first records."""
    stored = tla._stored_execution(execution_id)
    code = tla._code_text(stored["code_id"])
    prepared = runtime._platform("prepare_user_code")(code)
    frame, _metadata, _table = runtime._platform("load_execution_input")(stored["dataset_id"], stored.get("table_id"))
    inputs = [str(column) for column in frame.columns]
    try:
        analysis = la.analyse(prepared, inputs)
    except la.LineageError as exc:
        raise HTTPException(400, str(exc)) from exc
    stored_rows = stored.get("data") or []
    result_frame = pd.DataFrame(stored_rows) if stored_rows else frame
    numeric = _numeric_columns(result_frame)
    try:
        probe = _probe_catalogue(analysis, prepared, frame, runtime._platform("execution_globals"), field_names(analysis, inputs))
    except Exception as exc:  # the preview still shows what the code analysis finds
        traceback.print_exc()
        probe = {"enabled": False, "reason": "error", "rows": 0, "found": {}, "tested": [], "budget_hit": False, "message": str(exc)[:200]}
    code_meta = tla._code_meta(stored["code_id"])
    lines = code.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    return {
        "execution_id": execution_id,
        "code": {"filename": code_meta.get("filename"), "line_count": len(lines)},
        "rows": len(stored_rows) if stored_rows else len(frame),
        "inputs": inputs,
        "fields": field_catalogue(analysis, inputs, numeric, probe),
        "probe": {key: probe.get(key) for key in ("enabled", "reason", "rows", "budget_hit", "ms")} | {"tested": len(probe.get("tested") or [])},
    }


@router.post("/runs")
def start_run(request: RunRequest):
    execution_id = (request.execution_id or "").strip()
    field = (request.field or "").strip()
    if not field:
        raise HTTPException(400, "Select the field whose content the analysis follows.")
    stored = tla._stored_execution(execution_id)
    if not tla._llm_configured():
        raise HTTPException(503, "OpenAI is not configured. Set OPENAI_API_KEY, or AZURE_OPENAI_API_KEY and AZURE_OPENAI_ENDPOINT.")
    with _JOBS_LOCK:
        if len(_JOBS) >= MAX_ACTIVE_RUNS:
            raise HTTPException(409, "Two content lineage runs are already in progress. Wait for one to finish or cancel it.")
        if any(job.state["source"]["execution_id"] == execution_id for job in _JOBS.values()):
            raise HTTPException(409, "This execution is already being analysed.")
    code = tla._code_text(stored["code_id"])
    code_meta = tla._code_meta(stored["code_id"])
    meta = tla._dataset_meta(stored["dataset_id"])
    cluster = None
    executed_date = None
    for candidate in db.get_all_clusters():
        for execution in db.get_cluster_executions(candidate["id"]):
            if execution.get("execution_id") == execution_id:
                cluster, executed_date = candidate, execution.get("executed_date")
                break
        if cluster:
            break
    language = request.language if request.language in LANGUAGE_NAMES else "en"
    summary = stored.get("summary") or {}
    run_id = str(uuid.uuid4())
    state = {
        "id": run_id,
        "status": "running",
        "stage": "queued",
        "created_at": _now(),
        "finished_at": None,
        "duration_ms": None,
        "language": language,
        "model": None,
        "usage": {"llm_calls": 0, "prompt_tokens": 0, "completion_tokens": 0, "tool_calls": 0},
        "source": {
            "execution_id": execution_id,
            "executed_date": executed_date,
            "cluster_id": (cluster or {}).get("id"),
            "cluster_name": (cluster or {}).get("name"),
            "reporting_date": (cluster or {}).get("reporting_date"),
            "dataset_id": stored["dataset_id"],
            "dataset_name": meta.get("user_name") or (cluster or {}).get("dataset_name") or meta.get("filename"),
            "dataset_file": meta.get("filename"),
            "table_id": stored.get("table_id"),
            "code_id": stored["code_id"],
            "code_filename": code_meta.get("filename") or (cluster or {}).get("code_filename") or "script.py",
            "code_lines": len(code.rstrip("\n").split("\n")),
            "rows": summary.get("rows_processed"),
        },
        "settings": {"mode": RESULT_MODE, "field": field, "fields_reached": [], "record_rows": RECORD_ROWS},
        "progress": {"rules_total": 0, "rules_described": 0},
        "summary": None,
        "result": None,
        "log": [],
        "error": None,
    }
    job = _Job(state)
    job.log("info", "run_started", f"Run started for [{field}] in execution {execution_id[:8]} of {state['source']['cluster_name'] or 'unknown cluster'}", field=field, execution=execution_id[:8], cluster=state["source"]["cluster_name"] or "")
    with _JOBS_LOCK:
        _JOBS[run_id] = job
    job.save(force=True)
    threading.Thread(target=_run_job, args=(job,), name=f"content-lineage-{run_id[:8]}", daemon=True).start()
    return job.snapshot(lite=True)


@router.get("/runs")
def list_runs():
    runs = []
    if RUNS_DIR.exists():
        for path in sorted(RUNS_DIR.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True)[:40]:
            try:
                state = _read_run(path.stem)
            except (HTTPException, OSError, json.JSONDecodeError):
                continue
            runs.append({key: state.get(key) for key in ("id", "status", "created_at", "finished_at", "duration_ms", "model", "language", "source", "settings", "summary")})
    runs.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"runs": runs}


@router.get("/runs/{run_id}")
def get_run(run_id: str, lite: bool = False):
    job = _JOBS.get(run_id)
    if job is not None:
        return job.snapshot(lite=lite)
    state = _read_run(run_id)
    if lite:
        state.pop("result", None)
    return state


@router.post("/runs/{run_id}/cancel")
def cancel_run(run_id: str):
    job = _JOBS.get(run_id)
    if job is None:
        raise HTTPException(404, "This run is not in progress.")
    job.cancel.set()
    job.log("warn", "cancel_requested", "Cancellation requested.")
    return {"id": run_id, "status": "cancelling"}


@router.delete("/runs/{run_id}")
def delete_run(run_id: str):
    if run_id in _JOBS:
        raise HTTPException(409, "Cancel the run before deleting it.")
    path = _run_path(run_id)
    if not path.exists():
        raise HTTPException(404, "Run not found")
    path.unlink()
    return {"id": run_id, "deleted": True}


# --------------------------------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------------------------------

_EXPORT: Dict[str, Dict[str, Any]] = {
    "en": {
        "sheets": ("Summary", "Downstream fields", "Routes", "Value flow", "Rules and cases", "Records", "Business glossary", "Checks"),
        "title": "Content lineage",
        "summary": {
            "field": "Field followed", "cluster": "Cluster", "execution": "Execution", "code": "Code", "dataset": "Dataset",
            "created": "Started (UTC)", "model": "Model", "language": "Language of the texts", "domain": "Domain", "summary": "Summary",
            "next": "Next fields", "fields": "Downstream fields", "ends": "End fields", "records": "Records",
            "reached": "Records whose end values depend on the field", "rules": "Business rules", "cases": "Cases",
            "impact": "Impact test",
        },
        "impact": "{runs} run(s) with other values of the field on {rows} records; {confirmed} of {fields} downstream fields change",
        "impact_off": "not run",
        "fields": (
            "Business term", "Column", "Step", "Position", "Reached via", "Evidence", "How the field reaches it", "Records", "Reached",
            "Value carried", "Case or parameter decided", "Shown by the impact test", "Not reached: delivered by the source",
            "Not reached: other case", "Missing", "Amount reached", "Amount total", "Business rule",
        ),
        "positions": {"next": "Next field", "downstream": "Further downstream"},
        "end": "end field",
        "evidence": {"code": "Code analysis and impact test", "impact": "Impact test only"},
        "routes": ("Route", "Records carried along the route", "Amount at the end field", "Example records"),
        "flow": ("Value of the followed field", "Downstream field", "Records", "Reached", "Values of the downstream field", "Amount"),
        "rules": ("Rule", "Business rule", "Derives", "Statement", "Case", "Case label", "Description", "Reads (followed content)", "Condition (code)", "Formula (code)", "Line", "Parameter", "Records produced", "Records reached", "Amount"),
        "records": ("Record", "Value of the followed field", "Downstream field", "Value", "Case", "How it is reached", "Read from", "What if: other value of the followed field", "What if: value of the downstream field"),
        "glossary": ("Business term", "Column", "Role", "Definition", "Rule", "Filled"),
        "checks": ("Check", "Status", "Details"),
        "roles": {"start": "Selected field", "next": "Next field", "downstream": "Further downstream", "context": "Other input"},
        "reach": {"value": "Value carried", "decides": "Case or parameter decided", "tested": "Shown by the impact test"},
        "uses": {"value": "value", "key": "table key", "condition": "condition", "group": "grouping", "join": "join key", "impact": "impact test"},
        "delivered": "delivered by the source",
        "status": {"pass": "Pass", "warn": "Warning", "fail": "Fail", "skipped": "Skipped"},
        "checks_names": {
            "replay": "Replay reproduces the stored execution", "impact": "Impact test finds no field the code analysis missed",
            "observed": "Every downstream field changes in the impact test", "terms": "Every term named and defined",
            "cases": "Every case described", "usage": "How the field reaches every downstream field explained",
            "numbers": "Numbers in AI texts are parameters of the code", "unmapped": "Every category has a parameter",
            "complete": "End fields complete", "translation": "Texts translated",
        },
    },
    "de": {
        "sheets": ("Zusammenfassung", "Nachgelagerte Felder", "Wege", "Wertefluss", "Regeln und Fälle", "Datensätze", "Fachliches Glossar", "Prüfungen"),
        "title": "Inhalts-Herkunft",
        "summary": {
            "field": "Verfolgtes Feld", "cluster": "Cluster", "execution": "Ausführung", "code": "Code", "dataset": "Datensatz",
            "created": "Gestartet (UTC)", "model": "Modell", "language": "Sprache der Texte", "domain": "Fachgebiet", "summary": "Zusammenfassung",
            "next": "Nächste Felder", "fields": "Nachgelagerte Felder", "ends": "Endfelder", "records": "Datensätze",
            "reached": "Datensätze, deren Endwerte vom Feld abhängen", "rules": "Fachliche Regeln", "cases": "Fälle",
            "impact": "Wirkungstest",
        },
        "impact": "{runs} Lauf/Läufe mit anderen Werten des Feldes auf {rows} Datensätzen; {confirmed} von {fields} nachgelagerten Feldern ändern sich",
        "impact_off": "nicht ausgeführt",
        "fields": (
            "Fachbegriff", "Spalte", "Schritt", "Position", "Erreicht über", "Nachweis", "Wie das Feld es erreicht", "Datensätze", "Erreicht",
            "Wert übernommen", "Fall oder Parameter entschieden", "Durch den Wirkungstest gezeigt", "Nicht erreicht: von der Quelle geliefert",
            "Nicht erreicht: anderer Fall", "Fehlend", "Betrag erreicht", "Betrag gesamt", "Fachliche Regel",
        ),
        "positions": {"next": "Nächstes Feld", "downstream": "Weiter nachgelagert"},
        "end": "Endfeld",
        "evidence": {"code": "Codeanalyse und Wirkungstest", "impact": "Nur Wirkungstest"},
        "routes": ("Weg", "Datensätze entlang des Weges", "Betrag im Endfeld", "Beispieldatensätze"),
        "flow": ("Wert des verfolgten Feldes", "Nachgelagertes Feld", "Datensätze", "Erreicht", "Werte des nachgelagerten Feldes", "Betrag"),
        "rules": ("Regel", "Fachliche Regel", "Leitet ab", "Aussage", "Fall", "Fallbezeichnung", "Beschreibung", "Liest (verfolgter Inhalt)", "Bedingung (Code)", "Formel (Code)", "Zeile", "Parameter", "Erzeugte Datensätze", "Erreichte Datensätze", "Betrag"),
        "records": ("Datensatz", "Wert des verfolgten Feldes", "Nachgelagertes Feld", "Wert", "Fall", "Wie es erreicht wird", "Gelesen aus", "Was wäre wenn: anderer Wert des verfolgten Feldes", "Was wäre wenn: Wert des nachgelagerten Feldes"),
        "glossary": ("Fachbegriff", "Spalte", "Rolle", "Definition", "Regel", "Befüllt"),
        "checks": ("Prüfung", "Status", "Details"),
        "roles": {"start": "Ausgewähltes Feld", "next": "Nächstes Feld", "downstream": "Weiter nachgelagert", "context": "Weiterer Eingang"},
        "reach": {"value": "Wert übernommen", "decides": "Fall oder Parameter entschieden", "tested": "Durch den Wirkungstest gezeigt"},
        "uses": {"value": "Wert", "key": "Tabellenschlüssel", "condition": "Bedingung", "group": "Gruppierung", "join": "Verknüpfungsschlüssel", "impact": "Wirkungstest"},
        "delivered": "von der Quelle geliefert",
        "status": {"pass": "Bestanden", "warn": "Warnung", "fail": "Nicht bestanden", "skipped": "Übersprungen"},
        "checks_names": {
            "replay": "Wiederholung reproduziert die gespeicherte Ausführung", "impact": "Wirkungstest findet kein Feld, das die Codeanalyse übersehen hat",
            "observed": "Jedes nachgelagerte Feld ändert sich im Wirkungstest", "terms": "Jeder Begriff benannt und definiert",
            "cases": "Jeder Fall beschrieben", "usage": "Wie das Feld jedes nachgelagerte Feld erreicht, ist erklärt",
            "numbers": "Zahlen in KI-Texten sind Parameter des Codes", "unmapped": "Jede Kategorie hat einen Parameter",
            "complete": "Endfelder vollständig", "translation": "Texte übersetzt",
        },
    },
}


def _labels(state: Dict[str, Any], language: str) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    result = state.get("result") or {}
    ai = result.get("ai") or {}
    texts = (ai.get("texts") or {}).get(language) or (ai.get("texts") or {}).get(ai.get("language") or "en") or {}
    return _EXPORT[language], texts


@router.get("/runs/{run_id}/export")
def export_run(run_id: str, language: str = "en"):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    language = language if language in _EXPORT else "en"
    state = _read_run(run_id)
    result = state.get("result")
    if state.get("status") != "completed" or not result:
        raise HTTPException(409, "Only completed runs can be exported.")
    if result.get("mode") != RESULT_MODE:
        raise HTTPException(409, "This analysis was made with an earlier version of the agent (backward from reported figures). Start a new analysis to export it.")
    words, texts = _labels(state, language)
    terms_text = texts.get("terms", {})
    rules_text = texts.get("rules", {})
    usage_text = texts.get("usage", {})
    scope = result["scope"]
    field = scope["field"]
    flows = result["flows"]
    rule_of = {rule["column"]: rule for rule in result["rules"]}
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="F5C400")
    wrap = Alignment(vertical="top", wrap_text=True)

    def sheet(title: str, headers: Iterable[str], widths: List[int], first: bool = False) -> Any:
        target = workbook.active if first else workbook.create_sheet()
        target.title = title[:31]
        target.append(list(headers))
        for index, cell in enumerate(target[1], start=1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = wrap
            target.column_dimensions[cell.column_letter].width = widths[index - 1] if index - 1 < len(widths) else 18
        target.freeze_panes = "A2"
        return target

    def finish(target: Any) -> None:
        for row in target.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap

    def term_name(column: str) -> str:
        return (terms_text.get(column) or {}).get("name") or column

    def case_label(case_id: str) -> str:
        rule_id = case_id.split(".")[0]
        return ((rules_text.get(rule_id) or {}).get("cases", {}).get(case_id) or {}).get("label") or case_id

    def value_text(value: Any) -> Any:
        return "∅" if value is None else value

    summary = state.get("summary") or {}
    source = state.get("source") or {}
    impact = result.get("impact") or {}
    overview = sheet(words["sheets"][0], [words["title"], ""], [40, 110], first=True)
    impact_line = words["impact"].format(runs=impact.get("tested", 0), rows=impact.get("rows", 0), confirmed=impact.get("confirmed", 0), fields=len(scope["fields"])) if impact.get("enabled") else words["impact_off"]
    for key, value in (
        ("field", f"{term_name(field)} ({field})"), ("cluster", source.get("cluster_name")), ("execution", source.get("execution_id")),
        ("code", source.get("code_filename")), ("dataset", source.get("dataset_name")), ("created", state.get("created_at")),
        ("model", state.get("model")), ("language", language.upper()), ("domain", texts.get("domain")), ("summary", texts.get("summary")),
        ("next", ", ".join(term_name(name) for name in scope["next"])), ("fields", len(scope["fields"])),
        ("ends", ", ".join(term_name(name) for name in scope["ends"])), ("records", summary.get("records")),
        ("reached", summary.get("reached")), ("rules", summary.get("rules")), ("cases", summary.get("cases")),
        ("impact", impact_line),
    ):
        overview.append([words["summary"][key], value])
    finish(overview)

    fields_sheet = sheet(words["sheets"][1], words["fields"], [28, 20, 7, 20, 26, 22, 60, 9, 9, 10, 12, 12, 14, 12, 9, 14, 14, 30])
    for column in scope["fields"]:
        flow = flows[column]
        records = flow["records"]
        rule = rule_of.get(column)
        position = words["positions"]["next" if flow["direct"] else "downstream"] + (f" · {words['end']}" if flow["end"] else "")
        fields_sheet.append([
            term_name(column), column, flow["level"], position, ", ".join(term_name(name) for name in flow["via"]),
            words["evidence"][flow["evidence"]], usage_text.get(column, ""), records["total"], records["reached"],
            records["value"], records["decides"], records["tested"], records["delivered"], records["other"], records["missing"],
            flow["amount"]["reached"], flow["amount"]["total"], f"{rule['id']} · {(rules_text.get(rule['id']) or {}).get('name', '')}" if rule else "",
        ])
    finish(fields_sheet)

    record_ids = {item["label"]: item.get("id") or item["label"] for item in (result.get("records") or {}).get("items") or []}
    routes_sheet = sheet(words["sheets"][2], words["routes"], [90, 14, 18, 30])
    for route in result.get("routes") or []:
        routes_sheet.append([" → ".join(term_name(name) for name in route["fields"]), route["records"], route["amount"], ", ".join(str(record_ids.get(label, label)) for label in route.get("examples") or [])])
    finish(routes_sheet)

    flow_sheet = sheet(words["sheets"][3], words["flow"], [26, 28, 10, 10, 50, 16])
    for column in scope["fields"]:
        for group in flows[column].get("by_value") or []:
            outcomes = "; ".join(f"{value_text(item['value'])} ({item['records']})" for item in group.get("outcomes") or [])
            flow_sheet.append([group["key"], term_name(column), group["records"], group["reached"], outcomes, group.get("amount")])
    finish(flow_sheet)

    carriers = {field, *scope["fields"]}
    rules_sheet = sheet(words["sheets"][4], words["rules"], [8, 30, 26, 60, 8, 26, 60, 34, 40, 40, 8, 30, 10, 10, 16])
    for rule in result["rules"]:
        rule_text = rules_text.get(rule["id"]) or {}
        reached = {case["id"]: case for case in flows.get(rule["column"], {}).get("cases", [])}
        for case in rule["cases"]:
            case_text = (rule_text.get("cases") or {}).get(case["id"]) or {}
            parameter = "; ".join(f"{item['key']} → {_key_text(item['value']) if item['value'] is not None else '∅'}" for item in case.get("keys") or [] if item.get("mapped"))
            reads = ", ".join(f"{term_name(name)} ({'/'.join(words['uses'].get(role, role) for role in roles)})" for name, roles in case["roles"].items() if name in carriers and name != rule["column"])
            rules_sheet.append([
                rule["id"], rule_text.get("name", ""), term_name(rule["column"]), rule_text.get("statement", ""), case["id"],
                case_text.get("label", ""), case_text.get("description", ""), reads, _plain(case.get("condition")), _plain(case.get("formula")),
                case["line"], parameter, case.get("records"), (reached.get(case["id"]) or {}).get("reached"), case.get("amount"),
            ])
    finish(rules_sheet)

    records = result.get("records") or {}
    record_sheet = sheet(words["sheets"][5], [records.get("identifier") or words["records"][0], *words["records"][1:]], [18, 18, 26, 16, 30, 26, 26, 22, 22])
    for item in records.get("items") or []:
        for column in scope["fields"]:
            entry = (item.get("fields") or {}).get(column) or {}
            case = entry.get("case")
            whatif = entry.get("whatif") or {}
            record_sheet.append([
                item.get("id") or item["label"], value_text(item.get("x")), term_name(column), value_text((item.get("values") or {}).get(column)),
                (case_label(case) + (f" [{entry['key']}]" if entry.get("key") else "")) if case else words["delivered"],
                words["reach"].get(entry.get("reach") or "", "–"), ", ".join(term_name(name) for name in entry.get("via") or []),
                value_text(whatif.get("x_to")) if whatif else "", value_text(whatif.get("y_to")) if whatif else "",
            ])
    finish(record_sheet)

    glossary = sheet(words["sheets"][6], words["glossary"], [30, 22, 22, 70, 10, 12])
    for term in result["terms"]:
        profile = term.get("profile") or {}
        glossary.append([
            term_name(term["column"]), term["column"], words["roles"].get(term["role"], term["role"]), (terms_text.get(term["column"]) or {}).get("definition", ""),
            term.get("rule") or "", f"{profile.get('non_null', 0)}/{profile.get('total', 0)}",
        ])
    finish(glossary)

    checks_sheet = sheet(words["sheets"][7], words["checks"], [52, 14, 90])
    for check in result.get("checks") or []:
        checks_sheet.append([words["checks_names"].get(check["id"], check["id"]), words["status"].get(check["status"], check["status"]), json.dumps(check.get("params") or {}, ensure_ascii=False, default=str)])
    finish(checks_sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    stem = re.sub(r"[^\w.-]+", "_", f"{Path(str(source.get('code_filename') or 'content')).stem}-{field}")[:60]
    filename = f"content-lineage-{stem}-{language}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
