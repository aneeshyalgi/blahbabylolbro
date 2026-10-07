"""Technical Lineage AI Agent: column-level lineage of the code behind an execution, verified at runtime and documented
by an LLM.

Every run analyses one execution (the code of its cluster applied to its dataset) in a fixed pipeline:

1. parse       - static data-flow analysis of the code (``lineage_analysis.analyse``): version nodes per column write,
                 data / condition / group-by / join dependencies, lookup tables, dataset-level filters and sorts;
2. replay      - the code is replayed statement by statement in the platform sandbox (the one ``/api/execute`` uses)
                 and compared cell by cell with the stored execution; every changed cell is attributed to a statement;
3. probe       - every input column is perturbed, emptied and filled in turn; the output columns that change are the
                 runtime dependencies, compared with the static ones per output column;
4. document    - the LLM explains every derived column; it may only cite the statements and columns of that column's
                 verified lineage (closed sets, checked deterministically, one correction round);
5. review      - a second LLM pass reviews each explanation against the code (four eyes); when static and runtime
                 lineage disagree, an investigator agent with tools (code, lineage, probes, cell traces) explains why.

Runs execute in a background thread and are polled by the browser; finished runs are stored as JSON and can be
exported as an Excel workbook or as an OpenLineage run event with the column-lineage facet.
"""
from __future__ import annotations

import copy
import json
import os
import re
import threading
import time
import traceback
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Set, Tuple

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

import agent_runtime as runtime
import database as db
import lineage_analysis as la

router = APIRouter(prefix="/api/lineage-agent", tags=["lineage-agent"])

RUNS_DIR = runtime.RESULTS_DIR / "lineage_agent"
DATASETS_DIR = Path(os.environ.get("DATASETS_DIR", str(runtime.APP_DATA_ROOT / "uploads" / "datasets")))
TRACE_ROWS = 200
REPLAY_ROW_LIMIT = 20000
PROBE_ROW_LIMIT = 5000
PROBE_BUDGET_SECONDS = 150
BATCH_SIZE = 8
WORKERS = 3
MAX_ACTIVE_RUNS = 2
MAX_INVESTIGATION_STEPS = 8
DEFAULT_OPENAI_MODEL = "gpt-4.1"  # LINEAGE_AGENT_MODEL overrides it
LANGUAGE_NAMES = {"en": "English", "de": "German", "es": "Spanish"}
_ID = re.compile(r"^[0-9a-f-]{8,64}$")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _code_text(code_id: str) -> str:
    """The code file with normalised line ends ("\\r\\r\\n" from earlier uploads would double every line number)."""
    path = runtime.CODE_DIR / f"{code_id}.py"
    if not path.exists():
        raise HTTPException(404, "The code file of this execution no longer exists.")
    with open(path, encoding="utf-8", errors="replace", newline="") as handle:
        text = handle.read()
    return text.replace("\r\r\n", "\n").replace("\r\n", "\n").replace("\r", "\n")


def _code_meta(code_id: str) -> Dict[str, Any]:
    path = runtime.CODE_DIR / f"{code_id}_meta.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def _stored_execution(execution_id: str) -> Dict[str, Any]:
    if not _ID.match(execution_id or ""):
        raise HTTPException(404, "Execution not found")
    path = runtime.RESULTS_DIR / f"{execution_id}.json"
    if not path.exists():
        raise HTTPException(404, "The result of this execution no longer exists.")
    try:
        stored = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise HTTPException(500, "The stored execution result cannot be read.") from exc
    if not stored.get("code_id") or not stored.get("dataset_id"):
        raise HTTPException(400, "The execution does not reference its code and dataset.")
    return stored


def _dataset_meta(dataset_id: str) -> Dict[str, Any]:
    path = DATASETS_DIR / f"{dataset_id}_meta.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def _llm_configured() -> bool:
    return bool(os.environ.get("OPENAI_API_KEY")) or bool(os.environ.get("AZURE_OPENAI_API_KEY") and os.environ.get("AZURE_OPENAI_ENDPOINT"))


def _configured_model() -> str:
    return os.environ.get("LINEAGE_AGENT_MODEL", "").strip() or (
        DEFAULT_OPENAI_MODEL if os.environ.get("OPENAI_API_KEY") else os.environ.get("AZURE_OPENAI_DEPLOYMENT_NAME", "")
    )


# --------------------------------------------------------------------------------------------------
# Job state
# --------------------------------------------------------------------------------------------------

class _Cancelled(Exception):
    pass


class _Job:
    def __init__(self, state: Dict[str, Any]):
        self.state = state
        self.lock = threading.Lock()
        self.cancel = threading.Event()
        self.started = time.monotonic()
        self.last_save = 0.0

    def log(self, level: str, code: str, message: str, **params: Any) -> None:
        with self.lock:
            self.state["log"].append({
                "t": round(time.monotonic() - self.started, 2), "level": level, "code": code, "params": params, "message": message,
            })

    def stage(self, stage: str) -> None:
        with self.lock:
            self.state["stage"] = stage

    def progress(self, **changes: Any) -> None:
        with self.lock:
            self.state["progress"].update(changes)

    def increment(self, key: str) -> None:
        with self.lock:
            self.state["progress"][key] = self.state["progress"].get(key, 0) + 1

    def count_call(self, prompt: int, completion: int, tools: int = 0) -> None:
        with self.lock:
            usage = self.state["usage"]
            usage["llm_calls"] += 1
            usage["prompt_tokens"] += prompt
            usage["completion_tokens"] += completion
            usage["tool_calls"] += tools

    def set_model(self, model: str) -> None:
        with self.lock:
            self.state["model"] = model

    def check(self) -> None:
        if self.cancel.is_set():
            raise _Cancelled()

    def snapshot(self, lite: bool = False) -> Dict[str, Any]:
        with self.lock:
            if lite:
                return copy.deepcopy({key: value for key, value in self.state.items() if key != "result"})
            return copy.deepcopy(self.state)

    def save(self, force: bool = False) -> None:
        now = time.monotonic()
        if not force and now - self.last_save < 2.0:
            return
        self.last_save = now
        _write_run(self.snapshot())


_JOBS: Dict[str, _Job] = {}
_JOBS_LOCK = threading.Lock()


def _run_path(run_id: str) -> Path:
    if not _ID.match(run_id or ""):
        raise HTTPException(404, "Run not found")
    return RUNS_DIR / f"{run_id}.json"


def _write_run(state: Dict[str, Any]) -> None:
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    path = _run_path(state["id"])
    temporary = path.with_name(f"{path.stem}.{uuid.uuid4().hex[:8]}.tmp")
    temporary.write_text(json.dumps(state, ensure_ascii=False, default=str), encoding="utf-8")
    for attempt in range(20):
        try:
            os.replace(temporary, path)
            return
        except PermissionError:
            if attempt == 19:
                temporary.unlink(missing_ok=True)
                raise
            time.sleep(0.05)


def _read_run(run_id: str) -> Dict[str, Any]:
    path = _run_path(run_id)
    if not path.exists():
        raise HTTPException(404, "Run not found")
    for attempt in range(10):
        try:
            state = json.loads(path.read_text(encoding="utf-8"))
            break
        except PermissionError:
            if attempt == 9:
                raise
            time.sleep(0.05)
    if state.get("status") == "running" and run_id not in _JOBS:
        state["status"] = "interrupted"  # the server restarted while it ran
    return state


# --------------------------------------------------------------------------------------------------
# LLM
# --------------------------------------------------------------------------------------------------

class _LLM:
    """JSON-schema chat calls (and a tool loop for the investigator) with retries, model fallback and usage counts."""

    def __init__(self, job: _Job):
        client, configured = runtime._llm_client()
        override = os.environ.get("LINEAGE_AGENT_MODEL", "").strip()
        uses_openai = bool(os.environ.get("OPENAI_API_KEY"))
        self.client = client
        self.model = override or (DEFAULT_OPENAI_MODEL if uses_openai else configured)
        self.fallback = configured if self.model != configured else None
        self.job = job
        self.strict = True
        self.seed: Optional[int] = 20241007
        self.lock = threading.Lock()

    def _create(self, **options: Any) -> Any:
        last: Optional[Exception] = None
        for attempt in range(4):
            self.job.check()
            extra: Dict[str, Any] = {"seed": self.seed} if self.seed is not None else {}
            try:
                return self.client.chat.completions.create(model=self.model, temperature=0, **extra, **options)
            except Exception as exc:  # provider errors: fall back, degrade or retry
                last = exc
                status = getattr(exc, "status_code", None)
                if status in (403, 404) and self.fallback:
                    with self.lock:
                        if self.fallback:
                            self.job.log("warn", "model_fallback", f"Model {self.model} is not available; continuing with {self.fallback}.", model=self.model, fallback=self.fallback)
                            self.model, self.fallback = self.fallback, None
                            self.job.set_model(self.model)
                    continue
                if status == 400 and self.seed is not None and "seed" in str(exc):
                    self.seed = None
                    continue
                if status == 400 and self.strict and ("response_format" in str(exc) or "json_schema" in str(exc)) and "response_format" in options:
                    self.strict = False
                    options["response_format"] = {"type": "json_object"}
                    continue
                if status in (408, 409, 429) or (status or 0) >= 500 or status is None:
                    time.sleep(1.5 * (attempt + 1))
                    continue
                raise
        raise RuntimeError(f"The language model did not answer: {last}")

    def json(self, messages: List[Dict[str, Any]], name: str, schema: Dict[str, Any]) -> Dict[str, Any]:
        last: Optional[Exception] = None
        for _ in range(3):
            response_format: Dict[str, Any] = (
                {"type": "json_schema", "json_schema": {"name": name, "strict": True, "schema": schema}} if self.strict else {"type": "json_object"}
            )
            response = self._create(messages=messages, response_format=response_format)
            usage = getattr(response, "usage", None)
            self.job.count_call(getattr(usage, "prompt_tokens", 0) or 0, getattr(usage, "completion_tokens", 0) or 0)
            content = response.choices[0].message.content or "{}"
            try:
                return json.loads(content)
            except json.JSONDecodeError as exc:
                last = exc
        raise RuntimeError(f"The language model did not return valid JSON: {last}")

    def tools(self, messages: List[Dict[str, Any]], tools: List[Dict[str, Any]], required: bool = False) -> Any:
        response = self._create(messages=messages, tools=tools, tool_choice="required" if required else "auto")
        usage = getattr(response, "usage", None)
        message = response.choices[0].message
        self.job.count_call(getattr(usage, "prompt_tokens", 0) or 0, getattr(usage, "completion_tokens", 0) or 0, len(message.tool_calls or []))
        return message


def _strings() -> Dict[str, Any]:
    return {"type": "array", "items": {"type": "string"}}


def _integers() -> Dict[str, Any]:
    return {"type": "array", "items": {"type": "integer"}}


def _overview_schema(outputs: List[str]) -> Dict[str, Any]:
    return {
        "type": "object", "additionalProperties": False, "required": ["purpose", "stages", "key_outputs"],
        "properties": {
            "purpose": {"type": "string"},
            "stages": {"type": "array", "items": {
                "type": "object", "additionalProperties": False, "required": ["title", "start_line", "end_line", "description"],
                "properties": {"title": {"type": "string"}, "start_line": {"type": "integer"}, "end_line": {"type": "integer"}, "description": {"type": "string"}},
            }},
            "key_outputs": {"type": "array", "items": {"type": "string", "enum": outputs}},
        },
    }


def _rule_schema(columns: List[str]) -> Dict[str, Any]:
    return {
        "type": "object", "additionalProperties": False, "required": ["text", "lines", "inputs"],
        "properties": {"text": {"type": "string"}, "lines": _integers(), "inputs": {"type": "array", "items": {"type": "string", "enum": columns}}},
    }


def _document_schema(batch: List[str], columns: List[str]) -> Dict[str, Any]:
    return {
        "type": "object", "additionalProperties": False, "required": ["columns"],
        "properties": {"columns": {"type": "array", "items": {
            "type": "object", "additionalProperties": False,
            "required": ["column", "meaning", "summary", "formula", "rules", "notes"],
            "properties": {
                "column": {"type": "string", "enum": batch},
                "meaning": {"type": "string"},
                "summary": {"type": "string"},
                "formula": {"type": "string"},
                "rules": {"type": "array", "items": _rule_schema(columns)},
                "notes": {"type": "array", "items": {
                    "type": "object", "additionalProperties": False, "required": ["text", "lines", "severity"],
                    "properties": {"text": {"type": "string"}, "lines": _integers(), "severity": {"type": "string", "enum": ["info", "warning"]}},
                }},
            },
        }}},
    }


def _review_schema(batch: List[str], columns: List[str]) -> Dict[str, Any]:
    return {
        "type": "object", "additionalProperties": False, "required": ["columns"],
        "properties": {"columns": {"type": "array", "items": {
            "type": "object", "additionalProperties": False,
            "required": ["column", "verdict", "issues", "summary", "formula", "rules"],
            "properties": {
                "column": {"type": "string", "enum": batch},
                "verdict": {"type": "string", "enum": ["confirmed", "corrected"]},
                "issues": _strings(),
                "summary": {"type": "string"},
                "formula": {"type": "string"},
                "rules": {"type": "array", "items": _rule_schema(columns)},
            },
        }}},
    }


INVESTIGATION_SCHEMA = {
    "type": "object", "additionalProperties": False, "required": ["findings"],
    "properties": {"findings": {"type": "array", "items": {
        "type": "object", "additionalProperties": False,
        "required": ["title", "detail", "severity", "resolution", "lines", "columns"],
        "properties": {
            "title": {"type": "string"},
            "detail": {"type": "string"},
            "severity": {"type": "string", "enum": ["info", "warning", "critical"]},
            "resolution": {"type": "string", "enum": ["static_gap", "not_exercised", "probe_limitation", "code_issue", "input_changed", "other"]},
            "lines": _integers(),
            "columns": _strings(),
        },
    }}},
}


# --------------------------------------------------------------------------------------------------
# Facts handed to the LLM
# --------------------------------------------------------------------------------------------------

def _numbered(lines: List[str], start: int = 1, end: Optional[int] = None) -> str:
    end = end or len(lines)
    return "\n".join(f"{number:>4} | {lines[number - 1]}" for number in range(max(1, start), min(len(lines), end) + 1))


def _profile(values: List[Any]) -> Dict[str, Any]:
    non_null = [value for value in values if not la._is_null(value)]
    samples: List[Any] = []
    for value in non_null:
        item = la.json_value(value)
        if item not in samples:
            samples.append(item)
        if len(samples) == 5:
            break
    return {"non_null": len(non_null), "total": len(values), "samples": samples}


class _Context:
    """Everything the run knows about the execution, its code and the analysis results."""

    def __init__(self) -> None:
        self.stored: Dict[str, Any] = {}
        self.code = ""
        self.lines: List[str] = []
        self.prepared = ""
        self.frame: Any = None
        self.make_globals: Callable[..., Any] = lambda frame: {}
        self.analysis: Dict[str, Any] = {}
        self.columns: Dict[str, Dict[str, Any]] = {}
        self.nodes: List[Dict[str, Any]] = []
        self.steps_by_line: Dict[int, Dict[str, Any]] = {}
        self.trace: Dict[str, Any] = {}
        self.final: Any = None
        self.statements: Dict[int, Dict[str, Any]] = {}
        self.dependencies: Dict[str, Dict[str, List[str]]] = {}
        self.probes: List[Dict[str, Any]] = []
        self.probe_cache: Dict[Tuple[str, str], Dict[str, Any]] = {}
        self.probe_frame: Any = None
        self.baseline: Any = None
        self.input_profiles: Dict[str, Dict[str, Any]] = {}
        self.output_profiles: Dict[str, Dict[str, Any]] = {}
        self.language = "en"

    def step_of_line(self, line: int) -> Optional[Dict[str, Any]]:
        for step in self.analysis.get("steps", []):
            if step["line"] <= line <= step["end_line"]:
                return step
        return None

    def allowed_lines(self, column: str) -> Set[int]:
        """Lines an explanation of ``column`` may cite: its writes, their statements, and the variables they use."""
        item = self.columns[column]
        lines: Set[int] = set()
        variables = self.analysis.get("variables", {})
        for node_id in item["chain"]:
            node = self.nodes[node_id]
            if node.get("line"):
                lines |= set(range(node["line"], (node.get("end_line") or node["line"]) + 1))
                step = self.step_of_line(node["line"])
                if step is not None:
                    lines |= set(range(step["line"], step["end_line"] + 1))
            for name in node.get("uses", []):
                span = variables.get(name)
                if span:
                    lines |= set(range(span["line"], span["end_line"] + 1))
            for condition in node.get("conditions", []):
                if condition.get("line"):
                    lines.add(condition["line"])
        for name in item.get("lookups", []):
            lookup = next((entry for entry in self.analysis.get("lookups", []) if entry["name"] == name), None)
            if lookup:
                lines |= set(range(lookup["line"], lookup["end_line"] + 1))
        for op in self.analysis.get("dataset_ops", []):
            if op.get("line"):
                lines.add(op["line"])
        return lines

    def allowed_columns(self, column: str) -> Set[str]:
        item = self.columns[column]
        allowed = {column}
        allowed |= {entry["column"] for entry in item["upstream"]}
        allowed |= {entry["column"] for entry in item["sources"]}
        for op in self.analysis.get("dataset_ops", []):
            allowed |= set(op["columns"])
        return allowed

    def facts(self, column: str) -> Dict[str, Any]:
        item = self.columns[column]
        writes = []
        for node_id in item["chain"]:
            node = self.nodes[node_id]
            runtime_line = None
            step = self.step_of_line(node["line"]) if node.get("line") else None
            if step is not None:
                statement = self.statements.get(step["line"])
                runtime_line = (statement or {}).get("changed", {}).get(column, 0)
            writes.append({
                "line": node.get("line"),
                "statement": node.get("statement", "")[:600],
                "operation": node.get("operation"),
                "loop_bindings": node.get("bindings") or {},
                "conditions": [condition.get("expanded") or condition.get("code") for condition in node.get("conditions", [])],
                "keeps_previous_value_for_other_rows": bool(node.get("retained")),
                "cells_changed_in_replay": runtime_line,
            })
        lookups = []
        for name in item.get("lookups", []):
            lookup = next((entry for entry in self.analysis.get("lookups", []) if entry["name"] == name), None)
            if lookup:
                lookups.append({"name": name, "line": lookup["line"], "entries": lookup["entries"][:40]})
        dependency = self.dependencies.get(column, {})
        return {
            "column": column,
            "role": item["role"],
            "exists_in_input": item["in_input"],
            "input_profile": self.input_profiles.get(column),
            "output_profile": self.output_profiles.get(column),
            "writes_in_order": writes,
            "direct_inputs": [
                {"column": entry["column"], "how": [f"{t['type']}/{t['subtype']}" for t in entry["transformations"]], "lines": entry["lines"]}
                for entry in item["upstream"]
            ],
            "source_input_columns": [f"{entry['column']} ({entry['mode']})" for entry in item["sources"]],
            "lookup_tables": lookups,
            "runtime_verification": {
                "confirmed_sources": dependency.get("confirmed", []),
                "static_only_sources": dependency.get("static_only", []),
                "runtime_only_sources": dependency.get("runtime_only", []),
            },
        }


# --------------------------------------------------------------------------------------------------
# LLM steps
# --------------------------------------------------------------------------------------------------

def _language_rule(language: str) -> str:
    return f"Write every text field in {LANGUAGE_NAMES.get(language, 'English')}. Keep column names, code and numbers exactly as written."


def _overview(job: _Job, llm: _LLM, context: _Context) -> Optional[Dict[str, Any]]:
    outputs = context.analysis["outputs"]
    steps = [
        {"lines": f"{step['line']}-{step['end_line']}", "kind": step["kind"], "writes": sorted({context.nodes[node]["column"] for node in step["writes"]})}
        for step in context.analysis["steps"]
    ]
    messages = [
        {"role": "system", "content": (
            "You are a senior data-lineage engineer at a bank. You document pandas scripts that the bank's reporting platform "
            "runs on a dataset (the DataFrame `df`). Describe what the script does for a reviewer who knows the business but "
            "not the code. Be precise and factual; describe only what the code does. " + _language_rule(context.language)
        )},
        {"role": "user", "content": (
            f"Script `{context.stored.get('code_filename') or 'script.py'}` with line numbers:\n```\n{_numbered(context.lines)}\n```\n\n"
            f"Input columns: {json.dumps(context.analysis['inputs'], ensure_ascii=False)}\n"
            f"Output columns: {json.dumps(outputs, ensure_ascii=False)}\n"
            f"Statements (static analysis): {json.dumps(steps, ensure_ascii=False)}\n\n"
            "Return:\n"
            "- purpose: 2-3 sentences on what the script computes and why (as far as the code shows it).\n"
            "- stages: the script's processing stages in order (3-8), each with a short title, its first and last line "
            "and one or two sentences on what happens there. Stages must not overlap and must cover the statements.\n"
            "- key_outputs: the output columns the script exists to produce (derived results, not pass-through columns)."
        )},
    ]
    try:
        answer = llm.json(messages, "script_overview", _overview_schema(outputs))
    except _Cancelled:
        raise
    except Exception as exc:
        job.log("warn", "llm_error", f"Script overview failed: {str(exc)[:200]}", step="overview", message=str(exc)[:200])
        return None
    line_count = len(context.lines)
    stages = []
    for stage in answer.get("stages", []):
        start, end = int(stage.get("start_line") or 0), int(stage.get("end_line") or 0)
        if 1 <= start <= end <= line_count:
            stages.append({"title": stage.get("title", ""), "start_line": start, "end_line": end, "description": stage.get("description", "")})
    return {
        "purpose": answer.get("purpose", ""),
        "stages": stages,
        "key_outputs": [column for column in answer.get("key_outputs", []) if column in outputs],
    }


def _document_messages(context: _Context, batch: List[str]) -> List[Dict[str, Any]]:
    facts = [context.facts(column) for column in batch]
    return [
        {"role": "system", "content": (
            "You are a senior data-lineage engineer at a bank. You write the technical lineage documentation of output columns "
            "of a pandas script used for regulatory reporting. The column lineage below was established by static analysis of "
            "the code and verified by replaying the code on the data; treat it as ground truth. Your job is to explain it.\n"
            "Rules:\n"
            "1. Explain only what the code does. Never invent logic, thresholds, regulations or data that the code and facts do not show.\n"
            "2. Every rule must cite the line numbers of the statements it describes (`lines`) and list the columns it reads "
            "(`inputs`). Cite only lines listed in the column's facts (its writes, the variables and lookup tables they use).\n"
            "3. Describe the writes in execution order. When a write only fills missing values, say that existing (reported) "
            "values are kept. When conditions select rows, state the condition in business terms and with the column names.\n"
            "4. `formula`: one compact line; put column names in square brackets, e.g. "
            "`[EAD] = [EAD] if reported, else [Assessment Base] × [CCF]`. Mention lookup tables by their variable name.\n"
            "5. `meaning`: one sentence on what the column represents, inferred from its name and its derivation. If the "
            "meaning is not evident, return an empty string. Do not claim regulatory definitions as facts.\n"
            "6. `notes`: edge cases visible in the code or the replay facts (values coerced to empty by a type conversion, "
            "categories a lookup table does not cover, values that stay empty, overwritten values). No speculation.\n"
            "7. `summary`: one or two sentences, plain and specific.\n"
            + _language_rule(context.language)
        )},
        {"role": "user", "content": (
            f"Script with line numbers:\n```\n{_numbered(context.lines)}\n```\n\n"
            f"Column facts:\n{json.dumps(facts, ensure_ascii=False, default=str)}\n\n"
            f"Document these columns: {json.dumps(batch, ensure_ascii=False)}"
        )},
    ]


def _grounding_issues(context: _Context, column: str, doc: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Deterministic checks: cited lines and inputs belong to the column's verified lineage; direct inputs are covered.

    Each issue carries a code and parameters (shown translated in the UI) and an English text (sent back to the model).
    """
    issues: List[Dict[str, Any]] = []

    def issue(code: str, text: str, **params: Any) -> None:
        if not any(item["text"] == text for item in issues):
            issues.append({"code": code, "params": params, "text": text})

    allowed_lines = context.allowed_lines(column)
    allowed_columns = context.allowed_columns(column)
    lookups = set(context.columns[column].get("lookups_all", []))
    variables = set(context.analysis.get("variables", {}))
    rules = doc.get("rules") or []
    if not rules:
        issue("no_rules", "No rules were given.")
    covered: Set[str] = set()
    for index, rule in enumerate(rules, start=1):
        lines = [int(line) for line in rule.get("lines", []) if isinstance(line, (int, float))]
        if not lines:
            issue("uncited_rule", f"Rule {index} cites no line.", rule=index)
        wrong_lines = [line for line in lines if line not in allowed_lines]
        if wrong_lines:
            issue("lines_outside", f"Rule {index} cites line(s) {wrong_lines}, which are not part of the lineage of [{column}]. Allowed lines: {sorted(allowed_lines)}.", rule=index, lines=", ".join(map(str, wrong_lines)))
        wrong_inputs = [name for name in rule.get("inputs", []) if name not in allowed_columns]
        if wrong_inputs:
            issue("inputs_outside", f"Rule {index} lists input(s) {wrong_inputs}, which [{column}] does not depend on.", rule=index, inputs=", ".join(wrong_inputs))
        covered |= set(rule.get("inputs", []))
    for name in re.findall(r"\[([^\[\]]+)\]", doc.get("formula") or ""):
        if name.strip() not in allowed_columns and name.strip() not in lookups and name.strip() not in variables:
            issue("formula_name", f"The formula names [{name}], which is not in the lineage of [{column}].", name=name)
    direct = [entry["column"] for entry in context.columns[column]["upstream"] if any(t["type"] == "DIRECT" for t in entry["transformations"])]
    missing = [name for name in direct if name not in covered]
    if missing:
        issue("missing_inputs", f"The rules do not mention the direct input(s) {missing}.", inputs=", ".join(missing))
    return issues


def _sanitise(context: _Context, column: str, doc: Dict[str, Any]) -> Dict[str, Any]:
    """Remove citations outside the verified lineage (used when the correction round did not fix them)."""
    allowed_lines = context.allowed_lines(column)
    allowed_columns = context.allowed_columns(column)
    clean = copy.deepcopy(doc)
    for rule in clean.get("rules", []):
        rule["lines"] = [line for line in rule.get("lines", []) if line in allowed_lines]
        rule["inputs"] = [name for name in rule.get("inputs", []) if name in allowed_columns]
    for note in clean.get("notes", []):
        note["lines"] = [line for line in note.get("lines", []) if line in allowed_lines]
    return clean


def _document_batch(job: _Job, llm: _LLM, context: _Context, batch: List[str], docs: Dict[str, Dict[str, Any]]) -> None:
    all_columns = sorted(set(context.analysis["inputs"]) | set(context.analysis["outputs"]) | {item["name"] for item in context.analysis["columns"]})
    messages = _document_messages(context, batch)
    try:
        answer = llm.json(messages, "column_documentation", _document_schema(batch, all_columns))
    except _Cancelled:
        raise
    except Exception as exc:
        job.log("warn", "llm_error", f"Documentation failed for {', '.join(batch)}: {str(exc)[:200]}", step="document", message=str(exc)[:200])
        for column in batch:
            docs[column] = {"status": "failed", "error": str(exc)[:300]}
        return
    by_column = {item["column"]: item for item in answer.get("columns", []) if item.get("column") in batch}
    # Grounding: violations go back to the model once, with the reasons.
    feedback = {}
    for column in batch:
        doc = by_column.get(column)
        if doc is None:
            feedback[column] = ["The column was not documented."]
            continue
        issues = _grounding_issues(context, column, doc)
        if issues:
            feedback[column] = [item["text"] for item in issues]
    corrected_columns: Set[str] = set()
    if feedback:
        for column, issues in feedback.items():
            job.log("info", "grounding_retry", f"[{column}] grounding check failed: {issues[0][:160]}", column=column, issues=len(issues))
        retry = messages + [
            {"role": "assistant", "content": json.dumps(answer, ensure_ascii=False)},
            {"role": "user", "content": (
                "Deterministic checks rejected parts of your answer. Fix exactly these problems and return the complete answer "
                "for all columns again:\n" + json.dumps(feedback, ensure_ascii=False, indent=1)
            )},
        ]
        try:
            second = llm.json(retry, "column_documentation", _document_schema(batch, all_columns))
            for item in second.get("columns", []):
                if item.get("column") in feedback:
                    by_column[item["column"]] = item
                    corrected_columns.add(item["column"])
        except _Cancelled:
            raise
        except Exception as exc:
            job.log("warn", "llm_error", f"Correction round failed: {str(exc)[:200]}", step="correction", message=str(exc)[:200])
    for column in batch:
        doc = by_column.get(column)
        if doc is None:
            docs[column] = {"status": "failed", "error": "not documented"}
            continue
        issues = _grounding_issues(context, column, doc)
        status = "verified" if not issues else "partial"
        if issues:
            doc = _sanitise(context, column, doc)
        docs[column] = {
            "meaning": doc.get("meaning", ""), "summary": doc.get("summary", ""), "formula": doc.get("formula", ""),
            "rules": doc.get("rules", []), "notes": doc.get("notes", []),
            "grounding": {"status": status, "issues": issues, "corrected": column in corrected_columns},
            "status": "documented",
        }
        job.increment("columns_documented")
        job.log("info", "documented", f"[{column}] documented" + (" (grounding partial)" if issues else ""), column=column, grounding=status)


def _review_batch(job: _Job, llm: _LLM, context: _Context, batch: List[str], docs: Dict[str, Dict[str, Any]]) -> None:
    documented = [column for column in batch if docs.get(column, {}).get("status") == "documented"]
    if not documented:
        return
    all_columns = sorted(set(context.analysis["inputs"]) | set(context.analysis["outputs"]) | {item["name"] for item in context.analysis["columns"]})
    payload = [
        {"column": column, "facts": context.facts(column), "documentation": {key: docs[column][key] for key in ("meaning", "summary", "formula", "rules", "notes")}}
        for column in documented
    ]
    messages = [
        {"role": "system", "content": (
            "You are the second reviewer (four-eyes principle) of technical lineage documentation at a bank. Check each column's "
            "documentation line by line against the code and the facts. Mark `confirmed` when every statement is correct and "
            "complete. Mark `corrected` when anything is wrong, missing or misleading (wrong order of writes, a missed "
            "condition, a fill described as an overwrite, a wrong input, a wrong formula); then list the issues and return the "
            "full corrected summary, formula and rules (same rules as the author: cite lines of the column's lineage, list inputs). "
            "For `confirmed`, return empty summary, formula and rules. Do not rewrite for style. " + _language_rule(context.language)
        )},
        {"role": "user", "content": (
            f"Script with line numbers:\n```\n{_numbered(context.lines)}\n```\n\n"
            f"Documentation to review:\n{json.dumps(payload, ensure_ascii=False, default=str)}"
        )},
    ]
    try:
        answer = llm.json(messages, "documentation_review", _review_schema(documented, all_columns))
    except _Cancelled:
        raise
    except Exception as exc:
        job.log("warn", "llm_error", f"Review failed: {str(exc)[:200]}", step="review", message=str(exc)[:200])
        for column in documented:
            docs[column]["review"] = {"verdict": "unavailable", "issues": []}
        return
    reviewed = {item["column"]: item for item in answer.get("columns", []) if item.get("column") in documented}
    for column in documented:
        item = reviewed.get(column)
        if item is None:
            docs[column]["review"] = {"verdict": "unavailable", "issues": []}
            continue
        verdict = item.get("verdict")
        review = {"verdict": verdict, "issues": item.get("issues", [])}
        if verdict == "corrected" and (item.get("summary") or item.get("rules")):
            proposal = {
                "summary": item.get("summary") or docs[column]["summary"],
                "formula": item.get("formula") or docs[column]["formula"],
                "rules": item.get("rules") or docs[column]["rules"],
            }
            issues = _grounding_issues(context, column, proposal)
            if not issues:
                review["original"] = {key: docs[column][key] for key in ("summary", "formula", "rules")}
                docs[column].update(proposal)
                docs[column]["grounding"] = {"status": "verified", "issues": [], "corrected": True}
                review["applied"] = True
            else:
                review["applied"] = False
                review["rejected_because"] = issues[:3]
        docs[column]["review"] = review
        job.log("info", "reviewed", f"[{column}] review: {verdict}" + (" (correction applied)" if review.get("applied") else ""), column=column, verdict=verdict, applied=bool(review.get("applied")))


TRANSLATE_BATCH = 4


def _translation_schema(columns: List[str]) -> Dict[str, Any]:
    column_item: Dict[str, Any] = {"type": "string", "enum": columns} if columns else {"type": "string"}
    text_pair = {
        "type": "object", "additionalProperties": False, "required": ["title", "description"],
        "properties": {"title": {"type": "string"}, "description": {"type": "string"}},
    }
    return {
        "type": "object", "additionalProperties": False, "required": ["overview", "columns", "investigation"],
        "properties": {
            "overview": {
                "type": "object", "additionalProperties": False, "required": ["purpose", "stages"],
                "properties": {"purpose": {"type": "string"}, "stages": {"type": "array", "items": text_pair}},
            },
            "columns": {"type": "array", "items": {
                "type": "object", "additionalProperties": False,
                "required": ["column", "meaning", "summary", "formula", "rules", "notes", "review_issues"],
                "properties": {
                    "column": column_item, "meaning": {"type": "string"}, "summary": {"type": "string"}, "formula": {"type": "string"},
                    "rules": _strings(), "notes": _strings(), "review_issues": _strings(),
                },
            }},
            "investigation": {"type": "array", "items": {
                "type": "object", "additionalProperties": False, "required": ["title", "detail"],
                "properties": {"title": {"type": "string"}, "detail": {"type": "string"}},
            }},
        },
    }


def _brackets(text: str) -> List[str]:
    """The distinct [bracketed] names of a formula: a translation must keep exactly these names."""
    return sorted({name.strip() for name in re.findall(r"\[([^\[\]]+)\]", text or "")})


_MARKED = re.compile(r"⟦([^⟦⟧]+)⟧")


class _Protector:
    """Wraps every known name (columns, variables, lookup tables) in ⟦ ⟧ so a translation cannot alter it."""

    def __init__(self, names: List[str]):
        names = sorted({name for name in names if name and len(name) > 1}, key=len, reverse=True)
        self.pattern = re.compile(r"(?<![\w⟦])(" + "|".join(re.escape(name) for name in names) + r")(?![\w⟧])") if names else None

    def protect(self, text: str) -> str:
        return self.pattern.sub(r"⟦\1⟧", text or "") if self.pattern else (text or "")

    @staticmethod
    def valid(original: str, translated: str) -> bool:
        """The translation keeps exactly the protected names and no stray markers."""
        if not isinstance(translated, str):
            return False
        if (translated.count("⟦") != translated.count("⟧")) or (bool(original.strip()) != bool(translated.strip())):
            return False
        return sorted(set(_MARKED.findall(original))) == sorted(set(_MARKED.findall(translated)))

    @staticmethod
    def restore(text: str) -> str:
        return (text or "").replace("⟦", "").replace("⟧", "")


def _translate(
    job: _Job, llm: _LLM, context: _Context, overview: Optional[Dict[str, Any]], docs: Dict[str, Dict[str, Any]],
    investigation: Optional[Dict[str, Any]], target: str,
) -> Dict[str, Any]:
    """Every AI text of the run in another language; structure, citations and names stay identical.

    Column, variable and lookup-table names are wrapped in ⟦ ⟧ before translation; each translated field must contain
    exactly the same names (checked deterministically). A batch with rejected fields is sent back once with the
    reasons; fields that still fail keep their original wording.
    """
    analysis = context.analysis
    protector = _Protector([
        *analysis.get("inputs", []), *analysis.get("outputs", []), *[item["name"] for item in analysis.get("columns", [])],
        *[item["name"] for item in analysis.get("lookups", [])], *analysis.get("variables", {}).keys(),
    ])
    documented = [column for column, doc in docs.items() if doc.get("status") == "documented"]
    batches = [documented[index:index + TRANSLATE_BATCH] for index in range(0, len(documented), TRANSLATE_BATCH)] or [[]]
    translation: Dict[str, Any] = {"overview": None, "columns": {}, "investigation": None, "rejected": 0}
    findings = (investigation or {}).get("findings", [])
    lock = threading.Lock()

    def payload_for(index: int, batch: List[str]) -> Dict[str, Any]:
        first = index == 0
        return {
            "overview": {
                "purpose": protector.protect(overview.get("purpose", "")),
                "stages": [{"title": protector.protect(stage["title"]), "description": protector.protect(stage["description"])} for stage in overview.get("stages", [])],
            } if overview and first else {"purpose": "", "stages": []},
            "columns": [
                {
                    "column": column, "meaning": protector.protect(docs[column].get("meaning", "")), "summary": protector.protect(docs[column].get("summary", "")),
                    "formula": protector.protect(docs[column].get("formula", "")),
                    "rules": [protector.protect(rule.get("text", "")) for rule in docs[column].get("rules", [])],
                    "notes": [protector.protect(note.get("text", "")) for note in docs[column].get("notes", [])],
                    "review_issues": [protector.protect(issue) for issue in (docs[column].get("review") or {}).get("issues", [])],
                }
                for column in batch
            ],
            "investigation": [{"title": protector.protect(item.get("title", "")), "detail": protector.protect(item.get("detail", ""))} for item in findings] if first else [],
        }

    def check(original: Dict[str, Any], answer: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
        """Accepted fields of an answer (restored) and the paths of rejected fields."""
        accepted: Dict[str, Any] = {"overview": None, "columns": {}, "investigation": None}
        rejected: List[str] = []
        source_overview, answer_overview = original["overview"], answer.get("overview") or {}
        if source_overview["purpose"] or source_overview["stages"]:
            stages = answer_overview.get("stages", [])
            if len(stages) == len(source_overview["stages"]):
                purpose = answer_overview.get("purpose", "")
                good_purpose = _Protector.valid(source_overview["purpose"], purpose)
                if not good_purpose:
                    rejected.append("overview.purpose")
                merged_stages = []
                for position, (source_stage, stage) in enumerate(zip(source_overview["stages"], stages)):
                    title_ok = _Protector.valid(source_stage["title"], stage.get("title", ""))
                    description_ok = _Protector.valid(source_stage["description"], stage.get("description", ""))
                    if not (title_ok and description_ok):
                        rejected.append(f"overview.stages[{position}]")
                    merged_stages.append({
                        "title": _Protector.restore(stage["title"] if title_ok else source_stage["title"]),
                        "description": _Protector.restore(stage["description"] if description_ok else source_stage["description"]),
                    })
                accepted["overview"] = {"purpose": _Protector.restore(purpose if good_purpose else source_overview["purpose"]), "stages": merged_stages}
            else:
                rejected.append("overview.stages (count)")
        if original["investigation"]:
            items = answer.get("investigation") or []
            if len(items) == len(original["investigation"]):
                merged = []
                for position, (source_item, item) in enumerate(zip(original["investigation"], items)):
                    ok = _Protector.valid(source_item["title"], item.get("title", "")) and _Protector.valid(source_item["detail"], item.get("detail", ""))
                    if not ok:
                        rejected.append(f"investigation[{position}]")
                    merged.append({key: _Protector.restore(item[key] if ok else source_item[key]) for key in ("title", "detail")})
                accepted["investigation"] = merged
            else:
                rejected.append("investigation (count)")
        by_column = {item.get("column"): item for item in answer.get("columns", [])}
        for source in original["columns"]:
            column = source["column"]
            item = by_column.get(column)
            if item is None:
                rejected.append(f"columns[{column}]")
                continue
            entry: Dict[str, Any] = {}
            for key in ("meaning", "summary", "formula"):
                ok = _Protector.valid(source[key], item.get(key, ""))
                if key == "formula" and ok:
                    ok = _brackets(_Protector.restore(item.get(key, ""))) == _brackets(_Protector.restore(source[key]))
                if not ok:
                    rejected.append(f"columns[{column}].{key}")
                entry[key] = _Protector.restore(item.get(key, "") if ok else source[key])
            for key in ("rules", "notes", "review_issues"):
                values = item.get(key, [])
                if len(values) != len(source[key]):
                    rejected.append(f"columns[{column}].{key} (count)")
                    entry[key] = [_Protector.restore(text) for text in source[key]]
                    continue
                merged = []
                for position, (source_text, text) in enumerate(zip(source[key], values)):
                    ok = _Protector.valid(source_text, text)
                    if not ok:
                        rejected.append(f"columns[{column}].{key}[{position}]")
                    merged.append(_Protector.restore(text if ok else source_text))
                entry[key] = merged
            accepted["columns"][column] = entry
        return accepted, rejected

    def translate_batch(index: int, batch: List[str]) -> None:
        job.check()
        original = payload_for(index, batch)
        messages = [
            {"role": "system", "content": (
                f"You translate the technical data-lineage documentation of a bank from {LANGUAGE_NAMES.get(context.language, 'English')} "
                f"into {LANGUAGE_NAMES[target]}. Translate every text field faithfully and precisely, in a professional banking register. "
                "Names wrapped in ⟦ ⟧ are column, variable or lookup-table names: copy each of them exactly, including the ⟦ ⟧ markers, "
                "and never translate, inflect or drop them. Also keep code, operators and values in quotes unchanged, and keep numbers in their "
                "exact notation (0.0 stays 0.0, 90000 stays 90000). Formulas "
                "are text too: translate their words (for example 'if reported, else') and keep their [brackets]. Keep every list with "
                "the same number of items in the same order. Empty strings stay empty."
            )},
            {"role": "user", "content": json.dumps(original, ensure_ascii=False)},
        ]
        schema = _translation_schema(batch)
        try:
            answer = llm.json(messages, "translation", schema)
            accepted, rejected = check(original, answer)
            if rejected:
                retry = messages + [
                    {"role": "assistant", "content": json.dumps(answer, ensure_ascii=False)},
                    {"role": "user", "content": (
                        "These fields changed, dropped or added a ⟦name⟧ (or changed a list length). Return the complete answer again "
                        "with every ⟦name⟧ copied exactly: " + ", ".join(rejected[:30])
                    )},
                ]
                second, second_rejected = check(original, llm.json(retry, "translation", schema))
                if len(second_rejected) <= len(rejected):
                    accepted, rejected = second, second_rejected
        except _Cancelled:
            raise
        except Exception as exc:
            job.log("warn", "llm_error", f"Translation into {LANGUAGE_NAMES[target]} failed: {str(exc)[:200]}", step="translate", message=str(exc)[:200])
            return
        with lock:
            if accepted["overview"] is not None:
                translation["overview"] = accepted["overview"]
            if accepted["investigation"] is not None:
                translation["investigation"] = accepted["investigation"]
            translation["columns"].update(accepted["columns"])
            translation["rejected"] += len(rejected)

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        for future in [pool.submit(translate_batch, index, batch) for index, batch in enumerate(batches)]:
            future.result()
    return translation


# --------------------------------------------------------------------------------------------------
# Investigator (tool-using agent for discrepancies between static and runtime lineage)
# --------------------------------------------------------------------------------------------------

INVESTIGATOR_TOOLS = [
    {"type": "function", "function": {
        "name": "read_code", "description": "Read lines of the script (with line numbers).",
        "parameters": {"type": "object", "properties": {"start_line": {"type": "integer"}, "end_line": {"type": "integer"}}, "required": ["start_line", "end_line"]},
    }},
    {"type": "function", "function": {
        "name": "column_lineage", "description": "Static lineage of a column: its writes in order, direct inputs, source columns and runtime verification.",
        "parameters": {"type": "object", "properties": {"column": {"type": "string"}}, "required": ["column"]},
    }},
    {"type": "function", "function": {
        "name": "run_probe",
        "description": "Re-run the script with one input column changed and report which output columns change. mode: perturb (non-empty values changed), null (all values emptied) or fill (empty values filled).",
        "parameters": {"type": "object", "properties": {"column": {"type": "string"}, "mode": {"type": "string", "enum": ["perturb", "null", "fill"]}}, "required": ["column", "mode"]},
    }},
    {"type": "function", "function": {
        "name": "statement_effects", "description": "What the statement starting at a line changed during the replay (cells per column, errors).",
        "parameters": {"type": "object", "properties": {"line": {"type": "integer"}}, "required": ["line"]},
    }},
    {"type": "function", "function": {
        "name": "trace_cell", "description": "History of one output cell during the replay: every statement that changed it, before/after values and the inputs read.",
        "parameters": {"type": "object", "properties": {"row": {"type": "integer"}, "column": {"type": "string"}}, "required": ["row", "column"]},
    }},
]


def _probe(context: _Context, column: str, mode: str) -> Dict[str, Any]:
    key = (column, mode)
    if key in context.probe_cache:
        return context.probe_cache[key]
    frame = context.probe_frame
    if frame is None or column not in frame.columns:
        return {"column": column, "mode": mode, "error": "unknown input column", "changed": {}}
    modified = la.probe_frame(frame, column, mode, la.fill_value_for(frame[column]))
    if modified is None:
        result = {"column": column, "mode": mode, "skipped": True, "changed": {}, "ms": 0}
    else:
        started = time.perf_counter()
        try:
            output = la.run_script(context.prepared, modified, context.make_globals)
            result = {"column": column, "mode": mode, "changed": la.changed_columns(context.baseline, output), "ms": round((time.perf_counter() - started) * 1000, 1)}
        except Exception as exc:
            result = {"column": column, "mode": mode, "error": f"{type(exc).__name__}: {exc}"[:300], "changed": {}, "ms": round((time.perf_counter() - started) * 1000, 1)}
    context.probe_cache[key] = result
    return result


def _tool(context: _Context, name: str, arguments: Dict[str, Any]) -> Any:
    if name == "read_code":
        start = max(1, int(arguments.get("start_line") or 1))
        end = min(len(context.lines), int(arguments.get("end_line") or start), start + 120)
        return _numbered(context.lines, start, end)
    if name == "column_lineage":
        column = str(arguments.get("column") or "")
        if column not in context.columns:
            return {"error": f"Unknown column. Known columns: {sorted(context.columns)}"}
        return context.facts(column)
    if name == "run_probe":
        if context.probe_frame is None:
            return {"error": "Probes are not available for this run."}
        return _probe(context, str(arguments.get("column") or ""), str(arguments.get("mode") or "perturb"))
    if name == "statement_effects":
        line = int(arguments.get("line") or 0)
        step = context.step_of_line(line)
        if step is None:
            return {"error": "No statement starts at or contains this line."}
        return {"statement": step["code"][:800], "static_writes": sorted({context.nodes[node]["column"] for node in step["writes"]}), "replay": context.statements.get(step["line"])}
    if name == "trace_cell":
        column = str(arguments.get("column") or "")
        row = str(arguments.get("row"))
        history = context.trace.get("cells", {}).get(column, {}).get(row)
        if history is None:
            return {"history": [], "note": "No statement changed this cell (or the row is beyond the traced rows)."}
        return {"history": history}
    return {"error": f"Unknown tool {name}"}


def _investigate(job: _Job, llm: _LLM, context: _Context, issues: List[Dict[str, Any]]) -> Dict[str, Any]:
    job.log("info", "investigation_start", f"Investigating {len(issues)} discrepancy(ies) between static and runtime lineage", issues=len(issues))
    messages: List[Dict[str, Any]] = [
        {"role": "system", "content": (
            "You are a lineage investigator. Static analysis of a pandas script and runtime probes (re-running the script with one "
            "input column changed) disagree in places, or parts of the script could not be analysed statically. Use the tools to "
            "find out why, then report one finding per issue. Typical causes: the static analysis could not resolve a construct "
            "(static_gap); the data never exercises a branch, so a dependency has no effect (not_exercised); a probe change "
            "cannot show the dependency, or the probe raised an error (probe_limitation); a defect in the code (code_issue); "
            "the input data changed since the execution (input_changed). Be concise and specific, cite lines and columns. "
            + _language_rule(context.language)
        )},
        {"role": "user", "content": (
            f"Script:\n```\n{_numbered(context.lines)}\n```\n\nIssues to investigate:\n{json.dumps(issues, ensure_ascii=False, default=str)}"
        )},
    ]
    steps: List[Dict[str, Any]] = []
    for step in range(MAX_INVESTIGATION_STEPS):
        job.check()
        try:
            # The first turn must gather evidence with a tool; afterwards the agent decides when it knows enough.
            message = llm.tools(messages, INVESTIGATOR_TOOLS, required=step == 0)
        except _Cancelled:
            raise
        except Exception as exc:
            job.log("warn", "llm_error", f"Investigation failed: {str(exc)[:200]}", step="investigate", message=str(exc)[:200])
            break
        calls = message.tool_calls or []
        if not calls:
            break
        messages.append({"role": "assistant", "content": message.content or "", "tool_calls": [
            {"id": call.id, "type": "function", "function": {"name": call.function.name, "arguments": call.function.arguments}} for call in calls
        ]})
        for call in calls:
            try:
                arguments = json.loads(call.function.arguments or "{}")
            except json.JSONDecodeError:
                arguments = {}
            output = _tool(context, call.function.name, arguments)
            steps.append({"tool": call.function.name, "arguments": arguments})
            job.log("info", "tool_call", f"{call.function.name}({json.dumps(arguments, ensure_ascii=False)})", tool=call.function.name, args=json.dumps(arguments, ensure_ascii=False)[:160])
            messages.append({"role": "tool", "tool_call_id": call.id, "content": json.dumps(output, ensure_ascii=False, default=str)[:12000]})
    messages.append({"role": "user", "content": "Report your findings now: one finding per issue."})
    try:
        answer = llm.json(messages, "investigation", INVESTIGATION_SCHEMA)
    except _Cancelled:
        raise
    except Exception as exc:
        job.log("warn", "llm_error", f"Investigation report failed: {str(exc)[:200]}", step="investigate", message=str(exc)[:200])
        return {"issues": issues, "steps": steps, "findings": []}
    findings = []
    line_count = len(context.lines)
    for finding in answer.get("findings", []):
        finding["lines"] = [line for line in finding.get("lines", []) if 1 <= int(line) <= line_count]
        finding["columns"] = [column for column in finding.get("columns", []) if column in context.columns]
        findings.append(finding)
    job.log("success", "investigation_done", f"Investigation finished with {len(findings)} finding(s)", findings=len(findings))
    return {"issues": issues, "steps": steps, "findings": findings}


# --------------------------------------------------------------------------------------------------
# Pipeline
# --------------------------------------------------------------------------------------------------

def _load(job: _Job, context: _Context) -> None:
    stored = _stored_execution(job.state["source"]["execution_id"])
    context.stored = stored
    context.code = _code_text(stored["code_id"])
    context.lines = context.code.split("\n")
    if context.lines and context.lines[-1] == "":
        context.lines.pop()
    context.prepared = runtime._platform("prepare_user_code")(context.code)
    frame, _metadata, _table = runtime._platform("load_execution_input")(stored["dataset_id"], stored.get("table_id"))
    context.frame = frame
    context.make_globals = runtime._platform("execution_globals")
    job.log("info", "code_loaded", f"Loaded {job.state['source']['code_filename']} ({len(context.lines)} lines)", file=job.state["source"]["code_filename"], lines=len(context.lines))
    job.log("info", "data_loaded", f"Loaded the input table: {len(frame)} rows × {len(frame.columns)} columns", rows=len(frame), columns=len(frame.columns))


def _parse(job: _Job, context: _Context) -> None:
    job.stage("parse")
    started = time.perf_counter()
    try:
        analysis = la.analyse(context.prepared, [str(column) for column in context.frame.columns])
    except la.LineageError as exc:
        raise HTTPException(400, str(exc)) from exc
    context.analysis = analysis
    context.nodes = analysis["nodes"]
    context.columns = {item["name"]: item for item in analysis["columns"]}
    writes = sum(1 for node in analysis["nodes"] if node["kind"] == "write")
    column_edges = [edge for edge in analysis["edges"] if edge["kind"] == "column"]
    job.log(
        "success", "static_done",
        f"Static data-flow analysis: {len(analysis['steps'])} statements, {writes} column writes, {len(column_edges)} column dependencies, {len(analysis['lookups'])} lookup table(s) ({round((time.perf_counter() - started) * 1000)} ms)",
        statements=len(analysis["steps"]), writes=writes, edges=len(column_edges), lookups=len(analysis["lookups"]),
    )
    for item in analysis["unresolved"]:
        job.log("warn", "unresolved", f"Line {item['line']}: not fully resolved statically ({item['reason']}: {item['detail']})", line=item["line"], reason=item["reason"], detail=item["detail"])
    stored_columns = [str(column) for column in (context.stored.get("summary") or {}).get("columns") or []]
    if stored_columns and stored_columns != analysis["outputs"]:
        job.log("warn", "outputs_mismatch", "The statically derived output columns differ from the stored execution's columns.")


def _replay(job: _Job, context: _Context) -> Dict[str, Any]:
    job.stage("replay")
    analysis = context.analysis
    reads: Dict[int, Dict[str, List[str]]] = {}
    for node in context.nodes:
        if node["kind"] != "write" or node.get("step") is None:
            continue
        step_line = analysis["steps"][node["step"]]["line"]
        used = reads.setdefault(step_line, {}).setdefault(node["column"], [])
        for dependency in [*[item["node"] for item in node["data"]], *node["control"], *node["group"], *node.get("join", [])]:
            name = context.nodes[dependency]["column"]
            if name not in used:
                used.append(name)
    checks = []
    for lookup in analysis["lookups"]:
        if lookup["kind"] != "mapping":
            continue
        for use in lookup.get("uses", []):
            if use.get("method") != "map" or not use.get("column"):
                continue
            step = context.step_of_line(use["line"])
            if step is None:
                continue
            checks.append({"name": lookup["name"], "line": use["line"], "statement_line": step["line"], "column": use["column"], "keys": [key for key, _ in lookup["entries"]]})
    frame = context.frame
    partial = len(frame) > REPLAY_ROW_LIMIT
    if partial:
        frame = frame.iloc[:REPLAY_ROW_LIMIT]
    started = time.perf_counter()
    trace = la.trace(context.prepared, frame, context.make_globals, reads, checks, TRACE_ROWS, cancelled=job.cancel.is_set)
    elapsed = round((time.perf_counter() - started) * 1000)
    context.trace = trace
    context.final = trace["final"]
    context.statements = {statement["line"]: statement for statement in trace["statements"]}
    if trace["error"]:
        job.log("error", "replay_error", f"Replay stopped at line {trace['error']['line']}: {trace['error']['message']}", line=trace["error"]["line"], message=trace["error"]["message"])
    stored_rows = context.stored.get("data") or []
    stored_columns = (context.stored.get("summary") or {}).get("columns") or []
    if partial:
        comparison = la.compare_with_execution(trace["final"], stored_rows[:REPLAY_ROW_LIMIT], stored_columns)
        comparison["rows_match"] = True
        comparison["ok"] = comparison["columns_match"] and comparison["mismatches"] == 0
    else:
        comparison = la.compare_with_execution(trace["final"], stored_rows, stored_columns)
    comparison.update({"ms": elapsed, "partial": partial, "rows": len(frame), "error": trace["error"]})
    changed_cells = sum(sum(statement["changed"].values()) for statement in trace["statements"])
    if comparison["ok"] and not trace["error"]:
        job.log("success", "replay_done", f"Replay reproduces the stored execution: {comparison['cells_compared']} cells identical; {changed_cells} cell changes attributed to statements ({elapsed} ms)", cells=comparison["cells_compared"], changes=changed_cells, ms=elapsed)
    elif not comparison["columns_match"]:
        job.log(
            "warn", "replay_columns",
            f"The code produces other columns than the stored execution (extra: {', '.join(comparison['extra_columns']) or '–'}; missing: {', '.join(comparison['missing_columns']) or '–'}); the code has probably changed since the execution",
            extra=comparison["extra_columns"], missing=comparison["missing_columns"], mismatches=comparison["mismatches"],
        )
    else:
        job.log("warn", "replay_mismatch", f"Replay differs from the stored execution in {comparison['mismatches']} of {comparison['cells_compared']} cells (input data may have changed since the execution)", mismatches=comparison["mismatches"], cells=comparison["cells_compared"])
    # Every runtime change must happen in a statement that statically writes that column.
    unexplained = []
    for step in analysis["steps"]:
        statement = context.statements.get(step["line"])
        if not statement:
            continue
        static_columns = {context.nodes[node]["column"] for node in step["writes"]}
        for column, count in statement["changed"].items():
            if column not in static_columns:
                unexplained.append({"line": step["line"], "column": column, "cells": count})
                job.log("warn", "write_unexplained", f"Line {step['line']} changed [{column}] at runtime, which the static analysis did not expect", line=step["line"], column=column)
    computed = [(str(cell.get("row")), str(cell.get("column"))) for cell in context.stored.get("computed_cells") or []]
    attributed = sum(1 for row, column in computed if trace["last_line"].get(column, {}).get(row) is not None)
    comparison["write_consistency"] = {"unexplained": unexplained}
    comparison["attribution"] = {"computed": len(computed), "attributed": attributed}
    return comparison


def _probes(job: _Job, context: _Context) -> Dict[str, Any]:
    job.stage("probe")
    frame = context.frame
    limited = len(frame) > PROBE_ROW_LIMIT
    probe_frame = frame.iloc[:PROBE_ROW_LIMIT].copy() if limited else frame
    context.probe_frame = probe_frame
    started = time.perf_counter()
    try:
        context.baseline = la.run_script(context.prepared, probe_frame, context.make_globals)
    except Exception as exc:
        job.log("warn", "probes_skipped", f"Probes skipped: the script fails on the probe input ({str(exc)[:160]})", reason="error", message=str(exc)[:160])
        context.probe_frame = None
        return {"enabled": False, "reason": "error", "message": str(exc)[:300], "probes": [], "rows": len(probe_frame)}
    baseline_ms = (time.perf_counter() - started) * 1000
    plan: List[Tuple[str, str]] = []
    for column in probe_frame.columns:
        for mode in ("perturb", "null", "fill"):
            plan.append((str(column), mode))
    job.progress(probes_total=len(plan), probes_done=0)
    job.log("info", "probes_start", f"Running {len(plan)} dependency probes on {len(probe_frame)} rows (one input column changed per run)", count=len(plan), rows=len(probe_frame))
    runtime_deps: Dict[str, Set[str]] = {}
    errors: Dict[str, List[str]] = {}
    budget_hit = False
    for index, (column, mode) in enumerate(plan):
        job.check()
        if time.perf_counter() - started > PROBE_BUDGET_SECONDS:
            budget_hit = True
            break
        result = _probe(context, column, mode)
        context.probes.append(result)
        if result.get("error"):
            errors.setdefault(column, []).append(mode)
        for output in result.get("changed", {}):
            runtime_deps.setdefault(output, set()).add(column)
        job.progress(probes_done=index + 1)
        if mode == "fill":
            affected = sorted({output for item in context.probes if item["column"] == column for output in item.get("changed", {})})
            job.log("info", "probe_column", f"[{column}] → {', '.join(affected) if affected else 'no output changes'}", column=column, affected=affected[:20], count=len(affected))
        job.save()
    if budget_hit:
        job.log("warn", "probes_budget", f"Probe time budget reached after {len(context.probes)} of {len(plan)} probes", done=len(context.probes), total=len(plan))
    for column, modes in errors.items():
        job.log("info", "probe_error", f"[{column}] probe(s) {', '.join(modes)} raised an error (the code depends on this column's type or values)", column=column, modes=", ".join(modes))
    confirmed_total = static_total = runtime_only_total = 0
    for column, item in context.columns.items():
        if not item["in_output"]:
            continue
        static = {entry["column"] for entry in item["sources"]}
        observed = runtime_deps.get(column, set())
        confirmed = sorted(static & observed)
        static_only = sorted(static - observed)
        runtime_only = sorted(observed - static)
        context.dependencies[column] = {"confirmed": confirmed, "static_only": static_only, "runtime_only": runtime_only}
        confirmed_total += len(confirmed)
        static_total += len(static)
        runtime_only_total += len(runtime_only)
    job.log(
        "success" if not runtime_only_total else "warn", "probes_done",
        f"Dependency probes: {confirmed_total} of {static_total} static source dependencies confirmed at runtime; {runtime_only_total} runtime-only",
        confirmed=confirmed_total, total=static_total, runtime_only=runtime_only_total,
    )
    return {
        "enabled": True, "rows": len(probe_frame), "limited": limited, "budget_hit": budget_hit, "baseline_ms": round(baseline_ms, 1),
        "probes": context.probes, "planned": len(plan),
    }


def _profiles(context: _Context) -> None:
    for column in context.frame.columns:
        context.input_profiles[str(column)] = _profile(context.frame[column].tolist())
    if context.final is not None:
        for column in context.final.columns:
            context.output_profiles[str(column)] = _profile(context.final[column].tolist())


def _findings(context: _Context, replay: Dict[str, Any], probes: Dict[str, Any]) -> List[Dict[str, Any]]:
    findings: List[Dict[str, Any]] = []
    analysis = context.analysis

    def add(code: str, severity: str, affected: List[str], lines: List[int], **params: Any) -> None:
        findings.append({"id": f"f{len(findings) + 1}", "code": code, "severity": severity, "columns": affected, "lines": [line for line in lines if line], "params": params})

    if not replay.get("columns_match", True):
        add("replay_columns", "critical", [*replay.get("extra_columns", []), *replay.get("missing_columns", [])][:10], [], extra=replay.get("extra_columns", []) or "–", missing=replay.get("missing_columns", []) or "–")
    elif not replay.get("ok"):
        add("replay_mismatch", "critical", sorted({item["column"] for item in replay.get("examples", [])})[:10], [], mismatches=replay.get("mismatches", 0), cells=replay.get("cells_compared", 0), examples=replay.get("examples", [])[:5])
    if replay.get("error"):
        add("replay_error", "critical", [], [replay["error"]["line"]], message=replay["error"]["message"])
    for item in replay.get("write_consistency", {}).get("unexplained", []):
        add("write_unexplained", "warning", [item["column"]], [item["line"]], cells=item["cells"])
    for column, dependency in context.dependencies.items():
        if dependency["runtime_only"]:
            add("runtime_only", "warning", [column], [], sources=dependency["runtime_only"])
    for item in analysis["unresolved"]:
        if item["reason"] == "import":
            add("import", "warning", [], [item["line"]], module=item["detail"])
        else:
            add("unresolved", "warning", [], [item["line"]], reason=item["reason"], detail=item["detail"])
    for entry in context.trace.get("unmapped", []):
        if entry["values"]:
            consumers = next((lookup["used_by"] for lookup in analysis["lookups"] if lookup["name"] == entry["name"]), [])
            add("unmapped_values", "warning", consumers, [entry["line"]], lookup=entry["name"], column=entry["column"], values=entry["values"], rows=sum(item["rows"] for item in entry["values"]))
    for statement in context.trace.get("statements", []):
        for column, entry in (statement.get("nulled") or {}).items():
            step = context.steps_by_line.get(statement["line"])
            operations = {context.nodes[node]["operation"] for node in (step or {}).get("writes", []) if context.nodes[node]["column"] == column}
            add("coerced_to_null" if "cast" in operations else "values_removed", "warning", [column], [statement["line"]], count=entry["count"], examples=entry["examples"])
    if context.final is not None:
        for column, item in context.columns.items():
            if not item["in_output"] or item["role"] in ("passthrough", "cast") or column not in context.final.columns:
                continue
            values = context.final[column].tolist()
            missing = [index for index, value in zip(context.final.index, values) if la._is_null(value)]
            if missing:
                add("remaining_nulls", "warning", [column], [], count=len(missing), total=len(values), rows=[str(label) for label in missing[:10]])
    for step in analysis["steps"]:
        statement = context.statements.get(step["line"])
        if statement is None or statement.get("error") or not step["writes"]:
            continue
        written = sorted({context.nodes[node]["column"] for node in step["writes"]})
        if not any(statement["changed"].get(column) for column in written) and not statement.get("added"):
            add("no_effect", "info", written[:12], [step["line"]], columns=written[:12])
    for item in analysis["dead_writes"]:
        add("dead_write", "info", [item["column"]], [item["line"]])
    for column, item in context.columns.items():
        for entry in item["upstream"]:
            if entry["historical"]:
                add("historical_read", "info", [column], entry["lines"], source=entry["column"])
    for node in context.nodes:
        if node["kind"] == "write" and node.get("via") in ("chained", "inplace"):
            add("chained_assignment", "warning", [node["column"]], [node["line"]])
    possible = sorted({column for column, item in context.columns.items() if item.get("possible")})
    if possible:
        add("possible_writes", "info", possible[:20], [], count=len(possible))
    for column, dependency in context.dependencies.items():
        item = context.columns[column]
        static_only = [name for name in dependency["static_only"] if name != column]
        if static_only and item["role"] != "passthrough":
            add("static_only", "info", [column], [], sources=static_only)
    errors: Dict[str, List[str]] = {}
    for probe in probes.get("probes", []):
        if probe.get("error"):
            errors.setdefault(probe["column"], []).append(f"{probe['mode']}: {probe['error'][:120]}")
    for column, messages in errors.items():
        add("probe_error", "info", [column], [], messages=messages[:3])
    order = {"critical": 0, "warning": 1, "info": 2}
    findings.sort(key=lambda item: order[item["severity"]])
    for index, item in enumerate(findings, start=1):
        item["id"] = f"F{index:02d}"
    return findings


def _checks(context: _Context, replay: Dict[str, Any], probes: Dict[str, Any], docs: Dict[str, Dict[str, Any]]) -> List[Dict[str, Any]]:
    analysis = context.analysis
    checks = []
    stored_columns = [str(column) for column in (context.stored.get("summary") or {}).get("columns") or []]
    checks.append({"id": "syntax", "status": "pass", "params": {"lines": len(context.lines)}})
    checks.append({"id": "outputs", "status": "pass" if stored_columns == analysis["outputs"] else "fail", "params": {"static": len(analysis["outputs"]), "stored": len(stored_columns)}})
    unresolved = len(analysis["unresolved"])
    checks.append({"id": "static_coverage", "status": "pass" if not unresolved else "warn", "params": {"statements": len(analysis["steps"]), "unresolved": unresolved}})
    checks.append({"id": "replay", "status": "pass" if replay.get("ok") and not replay.get("error") else "fail", "params": {"cells": replay.get("cells_compared", 0), "mismatches": replay.get("mismatches", 0), "partial": replay.get("partial", False)}})
    unexplained = len(replay.get("write_consistency", {}).get("unexplained", []))
    checks.append({"id": "write_consistency", "status": "pass" if not unexplained else "fail", "params": {"unexplained": unexplained}})
    attribution = replay.get("attribution", {"computed": 0, "attributed": 0})
    checks.append({"id": "cell_attribution", "status": "pass" if attribution["attributed"] == attribution["computed"] else "warn", "params": attribution})
    if probes.get("enabled"):
        confirmed = sum(len(item["confirmed"]) for item in context.dependencies.values())
        total = sum(len(item["confirmed"]) + len(item["static_only"]) for item in context.dependencies.values())
        runtime_only = sum(len(item["runtime_only"]) for item in context.dependencies.values())
        checks.append({"id": "dependency_agreement", "status": "pass" if not runtime_only else "warn", "params": {"confirmed": confirmed, "total": total, "runtime_only": runtime_only, "probes": len(probes.get("probes", []))}})
    else:
        checks.append({"id": "dependency_agreement", "status": "skipped", "params": {"reason": probes.get("reason", "disabled")}})
    documented = [doc for doc in docs.values() if doc.get("status") == "documented"]
    partial = [doc for doc in documented if doc["grounding"]["status"] != "verified"]
    failed = [doc for doc in docs.values() if doc.get("status") == "failed"]
    checks.append({"id": "ai_grounding", "status": "pass" if documented and not partial and not failed else ("warn" if documented else "skipped"), "params": {"verified": len(documented) - len(partial), "documented": len(documented), "failed": len(failed)}})
    reviewed = [doc for doc in documented if doc.get("review", {}).get("verdict") in ("confirmed", "corrected")]
    corrected = [doc for doc in reviewed if doc["review"]["verdict"] == "corrected"]
    checks.append({"id": "four_eyes", "status": "pass" if reviewed and len(reviewed) == len(documented) else ("warn" if documented else "skipped"), "params": {"reviewed": len(reviewed), "documented": len(documented), "corrected": len(corrected)}})
    return checks


def _cells(context: _Context) -> Dict[str, Any]:
    final = context.final
    if final is None:
        return {"columns": [], "rows": [], "input": {}, "changes": {}, "computed": []}
    rows = []
    for label, record in final.head(TRACE_ROWS).iterrows():
        rows.append({"label": str(label), "values": {str(column): la.json_value(value) for column, value in record.items()}})
    inputs = {}
    for label, record in context.frame.head(TRACE_ROWS).iterrows():
        inputs[str(label)] = {str(column): la.json_value(value) for column, value in record.items()}
    computed = sorted({f"{cell.get('row')}|{cell.get('column')}" for cell in context.stored.get("computed_cells") or []})
    return {
        "columns": [str(column) for column in final.columns], "rows": rows, "input": inputs, "input_columns": [str(column) for column in context.frame.columns],
        "changes": context.trace.get("cells", {}), "computed": computed, "total_rows": len(final),
    }


def _summary(context: _Context, replay: Dict[str, Any], probes: Dict[str, Any], findings: List[Dict[str, Any]], docs: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    analysis = context.analysis
    outputs = [item for item in analysis["columns"] if item["in_output"]]
    roles: Dict[str, int] = {}
    for item in analysis["columns"]:
        roles[item["role"]] = roles.get(item["role"], 0) + 1
    column_edges = [edge for edge in analysis["edges"] if edge["kind"] == "column"]
    direct = sum(1 for edge in column_edges if any(t["type"] == "DIRECT" for t in edge["transformations"]))
    indirect = sum(1 for edge in column_edges if any(t["type"] == "INDIRECT" for t in edge["transformations"]))
    confirmed = sum(len(item["confirmed"]) for item in context.dependencies.values())
    static_total = sum(len(item["confirmed"]) + len(item["static_only"]) for item in context.dependencies.values())
    runtime_only = sum(len(item["runtime_only"]) for item in context.dependencies.values())
    severities = {"critical": 0, "warning": 0, "info": 0}
    for finding in findings:
        severities[finding["severity"]] += 1
    documented = [doc for doc in docs.values() if doc.get("status") == "documented"]
    return {
        "outputs": len(outputs),
        "inputs": len(analysis["inputs"]),
        "derived": sum(1 for item in outputs if item["role"] not in ("passthrough",)),
        "roles": roles,
        "edges": len(column_edges),
        "direct": direct,
        "indirect": indirect,
        "lookups": len(analysis["lookups"]),
        "statements": len(analysis["steps"]),
        "writes": sum(1 for node in analysis["nodes"] if node["kind"] == "write"),
        "unresolved": len(analysis["unresolved"]),
        "replay_ok": bool(replay.get("ok")) and not replay.get("error"),
        "cells_compared": replay.get("cells_compared", 0),
        "cell_changes": sum(sum(statement["changed"].values()) for statement in context.trace.get("statements", [])),
        "computed": replay.get("attribution", {}).get("computed", 0),
        "attributed": replay.get("attribution", {}).get("attributed", 0),
        "probes": len(probes.get("probes", [])),
        "confirmed": confirmed,
        "static_dependencies": static_total,
        "runtime_only": runtime_only,
        "verified_pct": round(100 * confirmed / static_total) if static_total else None,
        "findings": severities,
        "documented": len(documented),
        "grounded": sum(1 for doc in documented if doc["grounding"]["status"] == "verified"),
        "reviewed": sum(1 for doc in documented if doc.get("review", {}).get("verdict") in ("confirmed", "corrected")),
        "corrected": sum(1 for doc in documented if doc.get("review", {}).get("verdict") == "corrected"),
    }


def _run_job(job: _Job) -> None:
    time.sleep(0.2)  # let the start request's response go out first
    context = _Context()
    context.language = job.state["language"]
    try:
        job.stage("load")
        _load(job, context)
        job.check()
        _parse(job, context)
        context.steps_by_line = {step["line"]: step for step in context.analysis["steps"]}
        job.save(force=True)
        job.check()
        replay = _replay(job, context)
        _profiles(context)
        job.save(force=True)
        job.check()
        if job.state["settings"]["probes"]:
            probes = _probes(job, context)
        else:
            probes = {"enabled": False, "reason": "disabled", "probes": []}
            job.log("info", "probes_skipped", "Dependency probes are switched off for this run.", reason="disabled")
        job.save(force=True)
        job.check()

        # AI documentation and review.
        job.stage("document")
        llm = _LLM(job)
        job.set_model(llm.model)
        job.log("info", "model", f"Model: {llm.model} · temperature 0 · JSON-schema output", model=llm.model)
        targets = [item["name"] for item in context.analysis["columns"] if item["in_output"] and item["role"] != "passthrough"]
        batches = [targets[index:index + BATCH_SIZE] for index in range(0, len(targets), BATCH_SIZE)]
        job.progress(columns_total=len(targets), columns_documented=0)
        job.log("info", "documenting", f"Documenting {len(targets)} derived column(s) in {len(batches)} batch(es)", columns=len(targets), batches=len(batches))
        docs: Dict[str, Dict[str, Any]] = {}
        with ThreadPoolExecutor(max_workers=WORKERS) as pool:
            overview_future = pool.submit(_overview, job, llm, context)
            futures = [pool.submit(_document_batch, job, llm, context, batch, docs) for batch in batches]
            for future in futures:
                future.result()
            overview = overview_future.result()
        job.save(force=True)
        job.check()
        job.stage("review")
        with ThreadPoolExecutor(max_workers=WORKERS) as pool:
            for future in [pool.submit(_review_batch, job, llm, context, batch, docs) for batch in batches]:
                future.result()
        job.check()

        # Discrepancies go to the investigator agent.
        issues: List[Dict[str, Any]] = []
        for column, dependency in context.dependencies.items():
            if dependency["runtime_only"]:
                issues.append({"kind": "runtime_only_dependency", "column": column, "sources": dependency["runtime_only"]})
        for item in context.analysis["unresolved"]:
            if item["reason"] != "import":
                issues.append({"kind": "unresolved_statement", **item})
        for item in replay.get("write_consistency", {}).get("unexplained", []):
            issues.append({"kind": "unexpected_runtime_write", **item})
        if not replay.get("columns_match", True):
            issues.append({"kind": "replayed_columns_differ_from_stored_execution", "extra_columns": replay.get("extra_columns"), "missing_columns": replay.get("missing_columns")})
        elif not replay.get("ok"):
            issues.append({"kind": "replay_mismatch", "mismatches": replay.get("mismatches"), "examples": replay.get("examples", [])[:5]})
        investigation = None
        if issues:
            job.stage("investigate")
            investigation = _investigate(job, llm, context, issues[:12])
        job.check()

        # English and German are always available; a Spanish run is translated into both.
        languages = [code for code in ("en", "de") if code != context.language]
        translations: Dict[str, Any] = {}
        if languages and any(doc.get("status") == "documented" for doc in docs.values()) or (languages and overview):
            job.stage("translate")
            job.log(
                "info", "translating", f"Translating the AI documentation into {', '.join(LANGUAGE_NAMES[code] for code in languages)}",
                languages=", ".join(code.upper() for code in languages),
            )
            with ThreadPoolExecutor(max_workers=len(languages)) as pool:
                futures = {code: pool.submit(_translate, job, llm, context, overview, docs, investigation, code) for code in languages}
                for code, future in futures.items():
                    translations[code] = future.result()
                    job.log(
                        "success", "translated", f"AI documentation available in {LANGUAGE_NAMES[code]} ({len(translations[code]['columns'])} columns; {translations[code]['rejected']} field(s) kept in the original wording)",
                        language=code.upper(), columns=len(translations[code]["columns"]), rejected=translations[code]["rejected"],
                    )

        job.stage("assemble")
        findings = _findings(context, replay, probes)
        checks = _checks(context, replay, probes, docs)
        for column, item in context.columns.items():
            if item["in_output"] and item["role"] == "passthrough":
                docs.setdefault(column, {"status": "passthrough"})
        result = {
            "code": {"filename": job.state["source"]["code_filename"], "lines": context.lines, "line_count": len(context.lines)},
            "analysis": context.analysis,
            "runtime": {
                "replay": replay,
                "statements": context.trace.get("statements", []),
                "unmapped": context.trace.get("unmapped", []),
                "probes": probes,
                "dependencies": context.dependencies,
            },
            "profiles": {"input": context.input_profiles, "output": context.output_profiles},
            "cells": _cells(context),
            "ai": {"overview": overview, "columns": docs, "investigation": investigation, "language": context.language, "translations": translations},
            "findings": findings,
            "checks": checks,
        }
        summary = _summary(context, replay, probes, findings, docs)
        summary["languages"] = [context.language, *translations]
        with job.lock:
            job.state["result"] = result
            job.state["summary"] = summary
            job.state["status"] = "completed"
        job.log(
            "success", "completed",
            f"Lineage complete: {summary['outputs']} output columns, {summary['edges']} column dependencies, "
            + (f"{summary['verified_pct']}% of static dependencies confirmed at runtime" if summary["verified_pct"] is not None else "runtime probes off"),
            outputs=summary["outputs"], edges=summary["edges"], verified=summary["verified_pct"],
        )
    except _Cancelled:
        with job.lock:
            job.state["status"] = "cancelled"
        job.log("warn", "cancelled", "Run cancelled.")
    except HTTPException as exc:
        with job.lock:
            job.state.update(status="failed", error=str(exc.detail))
        job.log("error", "failed", str(exc.detail), message=str(exc.detail))
    except Exception as exc:
        traceback.print_exc()
        with job.lock:
            job.state.update(status="failed", error=str(exc)[:500])
        job.log("error", "failed", f"Run failed: {str(exc)[:300]}", message=str(exc)[:300])
    finally:
        with job.lock:
            job.state["finished_at"] = _now()
            job.state["duration_ms"] = int((time.monotonic() - job.started) * 1000)
            job.state["stage"] = "finished"
        job.save(force=True)
        with _JOBS_LOCK:
            _JOBS.pop(job.state["id"], None)


# --------------------------------------------------------------------------------------------------
# API
# --------------------------------------------------------------------------------------------------

class RunRequest(BaseModel):
    execution_id: str
    language: str = "en"
    probes: bool = True


@router.get("/sources")
def list_sources():
    clusters = []
    for cluster in db.get_all_clusters():
        executions = []
        for execution in db.get_cluster_executions(cluster["id"]):
            execution_id = execution.get("execution_id")
            if not execution_id or not (runtime.RESULTS_DIR / f"{execution_id}.json").exists():
                continue
            summary = execution.get("summary") or {}
            executions.append({
                "execution_id": execution_id,
                "executed_date": execution.get("executed_date"),
                "rows": summary.get("rows_processed"),
                "columns": len(summary.get("columns") or []),
                "values_computed": summary.get("total_values_computed"),
                "code_filename": execution.get("code_filename"),
                "dataset_name": execution.get("dataset_name"),
            })
        clusters.append({
            "id": cluster["id"],
            "name": cluster["name"],
            "reporting_date": cluster.get("reporting_date"),
            "description": cluster.get("description") or "",
            "is_reference": bool(cluster.get("is_reference")),
            "dataset_id": cluster.get("dataset_id"),
            "dataset_name": cluster.get("dataset_name"),
            "code_id": cluster.get("code_id"),
            "code_filename": cluster.get("code_filename"),
            "executions": executions,
        })
    return {"clusters": clusters, "model": _configured_model(), "llm_configured": _llm_configured()}


@router.get("/preview")
def preview(execution_id: str):
    stored = _stored_execution(execution_id)
    code = _code_text(stored["code_id"])
    lines = code.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    meta = _dataset_meta(stored["dataset_id"])
    table = next((item for item in meta.get("tables") or [] if item.get("id") == stored.get("table_id")), None) or ((meta.get("tables") or [None])[0])
    inputs = [
        {"name": column.get("name"), "type": column.get("data_type"), "nulls": column.get("null_count"), "total": column.get("total_count")}
        for column in (table or {}).get("columns") or [] if column.get("name")
    ]
    summary = stored.get("summary") or {}
    code_meta = _code_meta(stored["code_id"])
    return {
        "execution_id": execution_id,
        "code": {"filename": code_meta.get("filename"), "lines": lines, "line_count": len(lines), "syntax_valid": code_meta.get("syntax_valid", True)},
        "inputs": inputs,
        "outputs": summary.get("columns") or [],
        "computed_by_column": summary.get("computed_by_column") or {},
        "rows": summary.get("rows_processed"),
        "values_computed": summary.get("total_values_computed"),
    }


@router.post("/runs")
def start_run(request: RunRequest):
    execution_id = (request.execution_id or "").strip()
    stored = _stored_execution(execution_id)
    if not _llm_configured():
        raise HTTPException(503, "OpenAI is not configured. Set OPENAI_API_KEY, or AZURE_OPENAI_API_KEY and AZURE_OPENAI_ENDPOINT.")
    with _JOBS_LOCK:
        if len(_JOBS) >= MAX_ACTIVE_RUNS:
            raise HTTPException(409, "Two lineage runs are already in progress. Wait for one to finish or cancel it.")
        if any(job.state["source"]["execution_id"] == execution_id for job in _JOBS.values()):
            raise HTTPException(409, "This execution is already being analysed.")
    code = _code_text(stored["code_id"])
    code_meta = _code_meta(stored["code_id"])
    meta = _dataset_meta(stored["dataset_id"])
    cluster = None
    executed_date = None
    for candidate in db.get_all_clusters():
        for execution in db.get_cluster_executions(candidate["id"]):
            if execution.get("execution_id") == execution_id:
                cluster, executed_date = candidate, execution.get("executed_date")
                break
        if cluster:
            break
    language = request.language if request.language in LANGUAGE_NAMES else "en"
    summary = stored.get("summary") or {}
    run_id = str(uuid.uuid4())
    state = {
        "id": run_id,
        "status": "running",
        "stage": "queued",
        "created_at": _now(),
        "finished_at": None,
        "duration_ms": None,
        "language": language,
        "model": None,
        "usage": {"llm_calls": 0, "prompt_tokens": 0, "completion_tokens": 0, "tool_calls": 0},
        "source": {
            "execution_id": execution_id,
            "executed_date": executed_date,
            "cluster_id": (cluster or {}).get("id"),
            "cluster_name": (cluster or {}).get("name"),
            "reporting_date": (cluster or {}).get("reporting_date"),
            "dataset_id": stored["dataset_id"],
            "dataset_name": meta.get("user_name") or (cluster or {}).get("dataset_name") or meta.get("filename"),
            "dataset_file": meta.get("filename"),
            "table_id": stored.get("table_id"),
            "code_id": stored["code_id"],
            "code_filename": code_meta.get("filename") or (cluster or {}).get("code_filename") or "script.py",
            "code_lines": len(code.rstrip("\n").split("\n")),
            "rows": summary.get("rows_processed"),
            "columns": len(summary.get("columns") or []),
            "values_computed": summary.get("total_values_computed"),
        },
        "settings": {"probes": bool(request.probes), "trace_rows": TRACE_ROWS, "probe_row_limit": PROBE_ROW_LIMIT, "batch_size": BATCH_SIZE},
        "progress": {"probes_total": 0, "probes_done": 0, "columns_total": 0, "columns_documented": 0},
        "summary": None,
        "result": None,
        "log": [],
        "error": None,
    }
    job = _Job(state)
    job.log("info", "run_started", f"Run started for execution {execution_id[:8]} of {state['source']['cluster_name'] or 'unknown cluster'}", execution=execution_id[:8], cluster=state["source"]["cluster_name"] or "")
    with _JOBS_LOCK:
        _JOBS[run_id] = job
    job.save(force=True)
    threading.Thread(target=_run_job, args=(job,), name=f"lineage-{run_id[:8]}", daemon=True).start()
    return job.snapshot(lite=True)


@router.get("/runs")
def list_runs():
    runs = []
    if RUNS_DIR.exists():
        for path in sorted(RUNS_DIR.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True)[:40]:
            try:
                state = _read_run(path.stem)
            except (HTTPException, OSError, json.JSONDecodeError):
                continue
            runs.append({key: state.get(key) for key in ("id", "status", "created_at", "finished_at", "duration_ms", "model", "language", "source", "summary")})
    runs.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"runs": runs}


@router.get("/runs/{run_id}")
def get_run(run_id: str, lite: bool = False):
    job = _JOBS.get(run_id)
    if job is not None:
        return job.snapshot(lite=lite)
    state = _read_run(run_id)
    if lite:
        state.pop("result", None)
    return state


@router.post("/runs/{run_id}/cancel")
def cancel_run(run_id: str):
    job = _JOBS.get(run_id)
    if job is None:
        raise HTTPException(404, "This run is not in progress.")
    job.cancel.set()
    job.log("warn", "cancel_requested", "Cancellation requested.")
    return {"id": run_id, "status": "cancelling"}


@router.delete("/runs/{run_id}")
def delete_run(run_id: str):
    if run_id in _JOBS:
        raise HTTPException(409, "Cancel the run before deleting it.")
    path = _run_path(run_id)
    if not path.exists():
        raise HTTPException(404, "Run not found")
    path.unlink()
    return {"id": run_id, "deleted": True}


# --------------------------------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------------------------------

_LABELS: Dict[str, Dict[str, Any]] = {
    "en": {
        "title": "Technical Lineage AI Agent",
        "sheets": {"summary": "Summary", "lineage": "Lineage", "columns": "Columns", "transformations": "Transformations", "cells": "Cell provenance", "findings": "Findings", "checks": "Checks", "probes": "Probes", "code": "Code"},
        "meta": {
            "cluster": "Cluster", "execution": "Execution", "executed": "Executed", "dataset": "Dataset", "code": "Code file", "run": "Analysis",
            "finished": "Analysis finished", "model": "Model", "outputs": "Output columns", "edges": "Column dependencies",
            "replay": "Replay reproduces the execution", "confirmed": "Static dependencies confirmed at runtime", "runtime_only": "Runtime-only dependencies",
            "findings": "Findings (critical / warning / info)", "purpose": "Purpose", "stages": "Processing stages", "language": "Language",
        },
        "yes": "yes", "no": "no",
        "headers": {
            "lineage": ["Output column", "Input", "Type", "Subtype", "Operation", "Lines", "Runtime verification", "Explanation"],
            "columns": ["Column", "Role", "Meaning", "Summary", "Formula", "Rules", "Source columns", "Lookup tables", "Confirmed at runtime", "Not observable", "Runtime only", "AI grounding", "Four-eyes review"],
            "transformations": ["Column", "Line", "Operation", "Statement", "Condition", "Loop binding", "Keeps other rows"],
            "cells": ["Row", "Column", "Line", "Before", "After", "Values read at that point"],
            "findings": ["ID", "Severity", "Finding", "Columns", "Lines", "Details"],
            "checks": ["Check", "Status", "Details"],
            "probes": ["Input column", "Probe", "Output columns changed (cells)", "Error", "ms"],
            "code": ["Line", "Code"],
        },
        "operations": {
            "copy": "Copy", "cast": "Type conversion", "fill_missing": "Fill missing values", "conditional": "Conditional rule", "lookup": "Lookup",
            "arithmetic": "Calculation", "aggregation": "Aggregation", "window": "Running / window value", "function": "Function", "string": "Text operation",
            "date": "Date operation", "rounding": "Rounding", "binning": "Binning", "constant": "Constant", "rename": "Rename", "transformation": "Transformation",
        },
        "roles": {"passthrough": "Pass-through", "cast": "Type-converted", "enriched": "Enriched", "overwritten": "Overwritten", "created": "Created", "renamed": "Renamed", "dropped": "Dropped", "intermediate": "Intermediate"},
        "verification": {"confirmed": "confirmed", "confirmed_via": "confirmed via its sources", "not_observable": "not observable"},
        "modes": {"direct": "direct", "indirect": "indirect"},
        "probe_modes": {"perturb": "values changed", "null": "emptied", "fill": "gaps filled"},
        "severity": {"critical": "critical", "warning": "warning", "info": "info"},
        "status": {"pass": "passed", "warn": "warning", "fail": "failed", "skipped": "skipped"},
        "grounding": {"verified": "verified", "partial": "partial"},
        "review": {"confirmed": "confirmed", "corrected": "corrected", "unavailable": "unavailable"},
        "passthrough": "Pass-through", "keeps_input": "Keeps the reported value", "lookup_table": "lookup table", "input_value": "input value",
        "line": "line", "lines": "lines", "no_input": "no input (constant)",
        "checks": {
            "syntax": "Code parsed", "outputs": "Output columns match", "static_coverage": "Static coverage", "replay": "Replay reproduces the execution",
            "write_consistency": "Runtime writes match", "cell_attribution": "Computed cells attributed", "dependency_agreement": "Static and runtime lineage agree",
            "ai_grounding": "AI explanations grounded", "four_eyes": "Four-eyes review",
        },
        "findings": {
            "replay_columns": "Code no longer matches the execution", "replay_mismatch": "Replay differs from the stored execution", "replay_error": "Replay stopped with an error",
            "write_unexplained": "Unexpected runtime write", "runtime_only": "Dependency found only at runtime", "unresolved": "Statement not fully resolved",
            "import": "Import outside the sandbox", "unmapped_values": "Values without a mapping", "coerced_to_null": "Values emptied by a type conversion",
            "values_removed": "Values emptied", "remaining_nulls": "Empty values remain", "no_effect": "No effect on this data", "dead_write": "Overwritten write",
            "historical_read": "Earlier state used", "chained_assignment": "Chained assignment", "possible_writes": "Columns selected at runtime",
            "static_only": "Dependency not observable", "probe_error": "Probe raised an error",
        },
    },
    "de": {
        "title": "Technische Herkunft – KI-Agent",
        "sheets": {"summary": "Überblick", "lineage": "Herkunft", "columns": "Spalten", "transformations": "Transformationen", "cells": "Zellherkunft", "findings": "Befunde", "checks": "Kontrollen", "probes": "Tests", "code": "Code"},
        "meta": {
            "cluster": "Cluster", "execution": "Ausführung", "executed": "Ausgeführt", "dataset": "Datensatz", "code": "Codedatei", "run": "Analyse",
            "finished": "Analyse beendet", "model": "Modell", "outputs": "Ausgabespalten", "edges": "Spaltenabhängigkeiten",
            "replay": "Wiederholung reproduziert die Ausführung", "confirmed": "Statische Abhängigkeiten zur Laufzeit bestätigt", "runtime_only": "Nur zur Laufzeit beobachtete Abhängigkeiten",
            "findings": "Befunde (kritisch / Warnung / Hinweis)", "purpose": "Zweck", "stages": "Verarbeitungsstufen", "language": "Sprache",
        },
        "yes": "ja", "no": "nein",
        "headers": {
            "lineage": ["Ausgabespalte", "Eingabe", "Typ", "Untertyp", "Operation", "Zeilen", "Laufzeitbestätigung", "Erläuterung"],
            "columns": ["Spalte", "Rolle", "Bedeutung", "Zusammenfassung", "Formel", "Regeln", "Quellspalten", "Mapping-Tabellen", "Zur Laufzeit bestätigt", "Nicht beobachtbar", "Nur zur Laufzeit", "KI-Belege", "Vier-Augen-Prüfung"],
            "transformations": ["Spalte", "Zeile", "Operation", "Anweisung", "Bedingung", "Schleifenbindung", "Andere Zeilen bleiben"],
            "cells": ["Zeile", "Spalte", "Codezeile", "Vorher", "Nachher", "An dieser Stelle gelesene Werte"],
            "findings": ["ID", "Schweregrad", "Befund", "Spalten", "Zeilen", "Details"],
            "checks": ["Kontrolle", "Status", "Details"],
            "probes": ["Eingabespalte", "Test", "Geänderte Ausgabespalten (Zellen)", "Fehler", "ms"],
            "code": ["Zeile", "Code"],
        },
        "operations": {
            "copy": "Kopie", "cast": "Typkonvertierung", "fill_missing": "Fehlende Werte füllen", "conditional": "Bedingte Regel", "lookup": "Nachschlagen (Mapping)",
            "arithmetic": "Berechnung", "aggregation": "Aggregation", "window": "Laufender Wert / Fenster", "function": "Funktion", "string": "Textoperation",
            "date": "Datumsoperation", "rounding": "Rundung", "binning": "Klassenbildung", "constant": "Konstante", "rename": "Umbenennung", "transformation": "Transformation",
        },
        "roles": {"passthrough": "Durchgereicht", "cast": "Typkonvertiert", "enriched": "Angereichert", "overwritten": "Überschrieben", "created": "Neu erzeugt", "renamed": "Umbenannt", "dropped": "Entfernt", "intermediate": "Zwischenspalte"},
        "verification": {"confirmed": "bestätigt", "confirmed_via": "über ihre Quellen bestätigt", "not_observable": "nicht beobachtbar"},
        "modes": {"direct": "direkt", "indirect": "indirekt"},
        "probe_modes": {"perturb": "Werte verändert", "null": "geleert", "fill": "Lücken gefüllt"},
        "severity": {"critical": "kritisch", "warning": "Warnung", "info": "Hinweis"},
        "status": {"pass": "bestanden", "warn": "Warnung", "fail": "nicht bestanden", "skipped": "übersprungen"},
        "grounding": {"verified": "bestätigt", "partial": "teilweise"},
        "review": {"confirmed": "bestätigt", "corrected": "korrigiert", "unavailable": "nicht verfügbar"},
        "passthrough": "Durchgereicht", "keeps_input": "Behält den gemeldeten Wert", "lookup_table": "Mapping-Tabelle", "input_value": "Eingabewert",
        "line": "Zeile", "lines": "Zeilen", "no_input": "keine Eingabe (Konstante)",
        "checks": {
            "syntax": "Code eingelesen", "outputs": "Ausgabespalten stimmen überein", "static_coverage": "Statische Abdeckung", "replay": "Wiederholung reproduziert die Ausführung",
            "write_consistency": "Laufzeit-Schreibzugriffe passen", "cell_attribution": "Berechnete Zellen zugeordnet", "dependency_agreement": "Statische und Laufzeit-Herkunft stimmen überein",
            "ai_grounding": "KI-Erläuterungen belegt", "four_eyes": "Vier-Augen-Prüfung",
        },
        "findings": {
            "replay_columns": "Code passt nicht mehr zur Ausführung", "replay_mismatch": "Wiederholung weicht von der gespeicherten Ausführung ab", "replay_error": "Wiederholung mit Fehler abgebrochen",
            "write_unexplained": "Unerwarteter Schreibzugriff zur Laufzeit", "runtime_only": "Abhängigkeit nur zur Laufzeit gefunden", "unresolved": "Anweisung nicht vollständig aufgelöst",
            "import": "Import außerhalb der Sandbox", "unmapped_values": "Werte ohne Zuordnung", "coerced_to_null": "Werte durch Typkonvertierung geleert",
            "values_removed": "Werte geleert", "remaining_nulls": "Leere Werte bleiben", "no_effect": "Keine Wirkung auf diese Daten", "dead_write": "Überschriebener Schreibzugriff",
            "historical_read": "Früherer Stand genutzt", "chained_assignment": "Verkettete Zuweisung", "possible_writes": "Zur Laufzeit gewählte Spalten",
            "static_only": "Abhängigkeit nicht beobachtbar", "probe_error": "Test löste einen Fehler aus",
        },
    },
    "es": {
        "title": "Linaje técnico – Agente IA",
        "sheets": {"summary": "Resumen", "lineage": "Linaje", "columns": "Columnas", "transformations": "Transformaciones", "cells": "Procedencia de celdas", "findings": "Hallazgos", "checks": "Controles", "probes": "Pruebas", "code": "Código"},
        "meta": {
            "cluster": "Clúster", "execution": "Ejecución", "executed": "Ejecutado", "dataset": "Conjunto de datos", "code": "Archivo de código", "run": "Análisis",
            "finished": "Análisis terminado", "model": "Modelo", "outputs": "Columnas de salida", "edges": "Dependencias de columnas",
            "replay": "La re-ejecución reproduce la ejecución", "confirmed": "Dependencias estáticas confirmadas en ejecución", "runtime_only": "Dependencias solo en ejecución",
            "findings": "Hallazgos (crítico / advertencia / nota)", "purpose": "Propósito", "stages": "Etapas de procesamiento", "language": "Idioma",
        },
        "yes": "sí", "no": "no",
        "headers": {
            "lineage": ["Columna de salida", "Entrada", "Tipo", "Subtipo", "Operación", "Líneas", "Verificación en ejecución", "Explicación"],
            "columns": ["Columna", "Rol", "Significado", "Resumen", "Fórmula", "Reglas", "Columnas de origen", "Tablas de mapeo", "Confirmadas en ejecución", "No observables", "Solo en ejecución", "Fundamentación IA", "Revisión a cuatro ojos"],
            "transformations": ["Columna", "Línea", "Operación", "Sentencia", "Condición", "Variable del bucle", "Conserva otras filas"],
            "cells": ["Fila", "Columna", "Línea", "Antes", "Después", "Valores leídos en ese punto"],
            "findings": ["ID", "Gravedad", "Hallazgo", "Columnas", "Líneas", "Detalles"],
            "checks": ["Control", "Estado", "Detalles"],
            "probes": ["Columna de entrada", "Prueba", "Columnas de salida modificadas (celdas)", "Error", "ms"],
            "code": ["Línea", "Código"],
        },
        "operations": {
            "copy": "Copia", "cast": "Conversión de tipo", "fill_missing": "Completar valores vacíos", "conditional": "Regla condicional", "lookup": "Búsqueda (mapeo)",
            "arithmetic": "Cálculo", "aggregation": "Agregación", "window": "Valor acumulado / ventana", "function": "Función", "string": "Operación de texto",
            "date": "Operación de fecha", "rounding": "Redondeo", "binning": "Agrupación en intervalos", "constant": "Constante", "rename": "Renombrado", "transformation": "Transformación",
        },
        "roles": {"passthrough": "Sin cambios", "cast": "Tipo convertido", "enriched": "Enriquecida", "overwritten": "Sobrescrita", "created": "Creada", "renamed": "Renombrada", "dropped": "Eliminada", "intermediate": "Intermedia"},
        "verification": {"confirmed": "confirmada", "confirmed_via": "confirmada vía sus orígenes", "not_observable": "no observable"},
        "modes": {"direct": "directa", "indirect": "indirecta"},
        "probe_modes": {"perturb": "valores modificados", "null": "vaciada", "fill": "huecos rellenados"},
        "severity": {"critical": "crítico", "warning": "advertencia", "info": "nota"},
        "status": {"pass": "superado", "warn": "advertencia", "fail": "no superado", "skipped": "omitido"},
        "grounding": {"verified": "verificada", "partial": "parcial"},
        "review": {"confirmed": "confirmada", "corrected": "corregida", "unavailable": "no disponible"},
        "passthrough": "Sin cambios", "keeps_input": "Conserva el valor informado", "lookup_table": "tabla de mapeo", "input_value": "valor de entrada",
        "line": "línea", "lines": "líneas", "no_input": "sin entrada (constante)",
        "checks": {
            "syntax": "Código analizado", "outputs": "Las columnas de salida coinciden", "static_coverage": "Cobertura estática", "replay": "La re-ejecución reproduce la ejecución",
            "write_consistency": "Las escrituras en ejecución coinciden", "cell_attribution": "Celdas calculadas atribuidas", "dependency_agreement": "Linaje estático y de ejecución coinciden",
            "ai_grounding": "Explicaciones de IA fundamentadas", "four_eyes": "Revisión a cuatro ojos",
        },
        "findings": {
            "replay_columns": "El código ya no coincide con la ejecución", "replay_mismatch": "La re-ejecución difiere de la ejecución guardada", "replay_error": "La re-ejecución se detuvo con un error",
            "write_unexplained": "Escritura inesperada en ejecución", "runtime_only": "Dependencia hallada solo en ejecución", "unresolved": "Sentencia no resuelta del todo",
            "import": "Importación fuera del entorno aislado", "unmapped_values": "Valores sin mapeo", "coerced_to_null": "Valores vaciados por una conversión de tipo",
            "values_removed": "Valores vaciados", "remaining_nulls": "Quedan valores vacíos", "no_effect": "Sin efecto sobre estos datos", "dead_write": "Escritura sobrescrita",
            "historical_read": "Estado anterior usado", "chained_assignment": "Asignación encadenada", "possible_writes": "Columnas elegidas en ejecución",
            "static_only": "Dependencia no observable", "probe_error": "La prueba provocó un error",
        },
    },
}


def _localized_ai(state: Dict[str, Any], language: str) -> Dict[str, Any]:
    """The AI texts of a run in one language: the run's own language, or its translation, field by field."""
    ai = copy.deepcopy((state.get("result") or {}).get("ai") or {})
    if language == state.get("language"):
        return ai
    translation = (ai.get("translations") or {}).get(language)
    if not translation:
        return ai
    if ai.get("overview") and translation.get("overview"):
        overview = ai["overview"]
        overview["purpose"] = translation["overview"].get("purpose") or overview.get("purpose", "")
        for stage, text in zip(overview.get("stages", []), translation["overview"].get("stages", [])):
            stage["title"] = text.get("title") or stage["title"]
            stage["description"] = text.get("description") or stage["description"]
    for column, text in (translation.get("columns") or {}).items():
        doc = (ai.get("columns") or {}).get(column)
        if not doc or doc.get("status") != "documented":
            continue
        for key in ("meaning", "summary", "formula"):
            if text.get(key):
                doc[key] = text[key]
        for rule, rule_text in zip(doc.get("rules", []), text.get("rules", [])):
            rule["text"] = rule_text or rule["text"]
        for note, note_text in zip(doc.get("notes", []), text.get("notes", [])):
            note["text"] = note_text or note["text"]
        if doc.get("review") and text.get("review_issues"):
            doc["review"]["issues"] = text["review_issues"]
    if ai.get("investigation") and translation.get("investigation"):
        for finding, text in zip(ai["investigation"].get("findings", []), translation["investigation"]):
            finding["title"] = text.get("title") or finding["title"]
            finding["detail"] = text.get("detail") or finding["detail"]
    return ai


def _raw_inputs(analysis: Dict[str, Any]) -> Set[str]:
    """Input columns whose values reach the output unchanged (or are dropped): the edges from them are input-to-output."""
    return {item["name"] for item in analysis["columns"] if item["in_input"] and item["role"] in ("passthrough", "dropped")}


def _operations_by_input(analysis: Dict[str, Any], column: Dict[str, Any]) -> Dict[str, List[str]]:
    """Operations of a column's writes per column they read (the column itself: writes keeping or adjusting its value)."""
    nodes = analysis["nodes"]
    result: Dict[str, List[str]] = {}
    for node_id in column["chain"]:
        node = nodes[node_id]
        operation = node.get("operation")
        if not operation:
            continue
        read = {nodes[item["node"]]["column"] for item in node["data"]}
        read |= {nodes[item]["column"] for item in [*node["control"], *node["group"], *node.get("join", [])]}
        if operation == "cast" or node.get("reads_self") or node.get("retained"):
            read.add(column["name"])
        for name in read:
            result.setdefault(name, [])
            if operation not in result[name]:
                result[name].append(operation)
    return result


def _lines_label(lines: List[int], labels: Dict[str, Any]) -> str:
    return f"{labels['line'] if len(lines) == 1 else labels['lines']} {', '.join(map(str, lines))}"


def _mapping_rows(state: Dict[str, Any], language: str = "en") -> List[Dict[str, Any]]:
    labels = _LABELS.get(language, _LABELS["en"])
    result = state["result"]
    analysis = result["analysis"]
    nodes = analysis["nodes"]
    dependencies = result["runtime"].get("dependencies", {})
    probes = bool(result["runtime"].get("probes", {}).get("enabled"))
    docs = _localized_ai(state, language).get("columns", {})
    inputs = set(analysis["inputs"])
    by_name = {item["name"]: item for item in analysis["columns"]}
    rows = []
    for column in analysis["columns"]:
        if not column["in_output"]:
            continue
        name = column["name"]
        confirmed = set(dependencies.get(name, {}).get("confirmed", []))
        by_input = _operations_by_input(analysis, column)

        def label(operations: List[str]) -> str:
            return ", ".join(labels["operations"].get(operation, operation) for operation in operations)

        def status(source: str) -> str:
            if not probes:
                return ""
            if source in confirmed:
                return labels["verification"]["confirmed"]
            if source not in inputs:
                sources = [item["column"] for item in by_name.get(source, {}).get("sources", [])]
                if sources and all(item in confirmed for item in sources):
                    return labels["verification"]["confirmed_via"]
            return labels["verification"]["not_observable"]

        operations = label(column["operations"]) or labels["passthrough"]
        summary = (docs.get(name) or {}).get("summary", "")
        if column["role"] == "passthrough":
            rows.append({"output": name, "input": name, "type": "DIRECT", "subtype": "IDENTITY", "operation": labels["passthrough"], "lines": "", "verification": status(name), "explanation": summary})
            continue
        if column["input_used"]:
            lines = sorted({nodes[node]["line"] for node in column["chain"] if nodes[node].get("line")})
            rows.append({"output": name, "input": name, "type": "DIRECT", "subtype": "IDENTITY", "operation": f"{labels['keeps_input']} · " + (label(by_input.get(name, [])) or operations), "lines": ", ".join(map(str, lines)), "verification": status(name), "explanation": summary})
        for entry in column["upstream"]:
            transformations = entry["transformations"]
            rows.append({
                "output": name, "input": entry["column"],
                "type": " + ".join(sorted({t["type"] for t in transformations})),
                "subtype": " + ".join(t["subtype"] for t in transformations),
                "operation": label(by_input.get(entry["column"], [])) or operations, "lines": ", ".join(map(str, entry["lines"])),
                "verification": status(entry["column"]), "explanation": summary,
            })
        for lookup in column["lookups"]:
            rows.append({"output": name, "input": f"{labels['lookup_table']} {lookup}", "type": "DIRECT", "subtype": "TRANSFORMATION", "operation": labels["operations"]["lookup"], "lines": "", "verification": "", "explanation": summary})
        if not column["upstream"] and not column["input_used"] and not column["lookups"]:
            rows.append({"output": name, "input": labels["no_input"], "type": "", "subtype": "", "operation": operations, "lines": ", ".join(str(nodes[node]["line"]) for node in column["chain"]), "verification": "", "explanation": summary})
    return rows


def _excel(state: Dict[str, Any], language: str = "en") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    labels = _LABELS.get(language, _LABELS["en"])
    result = state["result"]
    analysis = result["analysis"]
    nodes = analysis["nodes"]
    source = state["source"]
    summary = state.get("summary") or {}
    ai = _localized_ai(state, language)
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="F5C400")
    workbook = Workbook()
    yes, no = labels["yes"], labels["no"]

    def sheet(key: str, rows: List[List[Any]], widths: List[int]) -> None:
        worksheet = workbook.create_sheet(labels["sheets"][key])
        worksheet.append(labels["headers"][key])
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in rows:
            worksheet.append([value if not isinstance(value, (list, dict)) else json.dumps(value, ensure_ascii=False) for value in row])
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        worksheet.freeze_panes = "A2"

    overview = workbook.active
    overview.title = labels["sheets"]["summary"]
    overview.append([labels["title"]])
    overview["A1"].font = Font(bold=True, size=14)
    findings = summary.get("findings") or {}
    meta = labels["meta"]
    rows = [
        (meta["cluster"], source.get("cluster_name")), (meta["execution"], source.get("execution_id")), (meta["executed"], source.get("executed_date")),
        (meta["dataset"], source.get("dataset_name")), (meta["code"], source.get("code_filename")), (meta["run"], state["id"]),
        (meta["finished"], state.get("finished_at")), (meta["model"], state.get("model")), (meta["language"], language.upper()),
        (meta["outputs"], summary.get("outputs")), (meta["edges"], summary.get("edges")),
        (meta["replay"], yes if summary.get("replay_ok") else no),
        (meta["confirmed"], f"{summary.get('confirmed')}/{summary.get('static_dependencies')}"),
        (meta["runtime_only"], summary.get("runtime_only")),
        (meta["findings"], f"{findings.get('critical', 0)} / {findings.get('warning', 0)} / {findings.get('info', 0)}"),
    ]
    for key, value in rows:
        overview.append([key, value])
    if ai.get("overview"):
        overview.append([])
        overview.append([meta["purpose"], ai["overview"].get("purpose", "")])
        for index, stage in enumerate(ai["overview"].get("stages", []), start=1):
            overview.append([f"{meta['stages']} {index}: {stage['title']} ({stage['start_line']}–{stage['end_line']})", stage["description"]])
    overview.column_dimensions["A"].width = 46
    overview.column_dimensions["B"].width = 110
    for row in overview.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet("lineage", [[row["output"], row["input"], row["type"], row["subtype"], row["operation"], row["lines"], row["verification"], row["explanation"]] for row in _mapping_rows(state, language)],
          [24, 26, 12, 26, 34, 10, 22, 80])
    docs = ai.get("columns", {})
    column_rows = []
    for column in analysis["columns"]:
        doc = docs.get(column["name"]) or {}
        dependency = result["runtime"].get("dependencies", {}).get(column["name"], {})
        grounding = (doc.get("grounding") or {}).get("status", "")
        verdict = (doc.get("review") or {}).get("verdict", "")
        column_rows.append([
            column["name"], labels["roles"].get(column["role"], column["role"]), doc.get("meaning", ""), doc.get("summary", ""), doc.get("formula", ""),
            "\n".join(f"{rule.get('text')} ({_lines_label(rule.get('lines', []), labels)})" for rule in doc.get("rules", [])),
            ", ".join(f"{entry['column']} ({labels['modes'][entry['mode']]})" for entry in column["sources"]), ", ".join(column["lookups_all"]),
            ", ".join(dependency.get("confirmed", [])), ", ".join(name for name in dependency.get("static_only", []) if name != column["name"]), ", ".join(dependency.get("runtime_only", [])),
            labels["grounding"].get(grounding, grounding), labels["review"].get(verdict, verdict),
        ])
    sheet("columns", column_rows, [24, 14, 40, 60, 50, 80, 50, 30, 40, 30, 30, 14, 16])
    transformation_rows = []
    for column in analysis["columns"]:
        for node_id in column["chain"]:
            node = nodes[node_id]
            transformation_rows.append([
                column["name"], node.get("line"), labels["operations"].get(node.get("operation"), node.get("operation")),
                node.get("statement", ""), "; ".join(condition.get("expanded") or condition.get("code") or "" for condition in node.get("conditions", [])),
                ", ".join(f"{key} = {value}" for key, value in (node.get("bindings") or {}).items()), yes if node.get("retained") else "",
            ])
    sheet("transformations", transformation_rows, [24, 8, 22, 80, 60, 24, 14])
    cell_rows = []
    for column, by_row in result["cells"].get("changes", {}).items():
        for row, history in by_row.items():
            for change in history:
                cell_rows.append([row, column, change["line"], change["before"], change["after"], ", ".join(f"{key} = {value}" for key, value in change.get("inputs", {}).items())])
    cell_rows.sort(key=lambda item: (int(item[0]) if str(item[0]).isdigit() else 0, item[1], item[2]))
    sheet("cells", cell_rows, [8, 24, 10, 18, 18, 80])
    sheet("findings", [
        [item["id"], labels["severity"][item["severity"]], labels["findings"].get(item["code"], item["code"]), ", ".join(item["columns"]), ", ".join(map(str, item["lines"])), json.dumps(item["params"], ensure_ascii=False, default=str)]
        for item in result["findings"]
    ], [8, 12, 34, 40, 12, 100])
    sheet("checks", [[labels["checks"].get(item["id"], item["id"]), labels["status"].get(item["status"], item["status"]), json.dumps(item["params"], ensure_ascii=False)] for item in result["checks"]], [40, 16, 80])
    probe_rows = [
        [probe["column"], labels["probe_modes"].get(probe["mode"], probe["mode"]), ", ".join(f"{key} ({value})" for key, value in probe.get("changed", {}).items()), probe.get("error", ""), probe.get("ms")]
        for probe in result["runtime"]["probes"].get("probes", [])
    ]
    sheet("probes", probe_rows, [26, 18, 90, 50, 8])
    sheet("code", [[index, line] for index, line in enumerate(result["code"]["lines"], start=1)], [8, 140])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _openlineage(state: Dict[str, Any], language: str = "en") -> Dict[str, Any]:
    labels = _LABELS.get(language, _LABELS["en"])
    result = state["result"]
    analysis = result["analysis"]
    source = state["source"]
    namespace = "dataflow-platform"
    input_name = source.get("dataset_name") or source.get("dataset_id")
    output_name = f"{input_name} · execution {source.get('execution_id', '')[:8]}"
    docs = _localized_ai(state, language).get("columns", {})
    raw = _raw_inputs(analysis)
    fields: Dict[str, Any] = {}
    for column in analysis["columns"]:
        if not column["in_output"]:
            continue
        description = (docs.get(column["name"]) or {}).get("summary") or ""
        input_fields = []
        if column["role"] == "passthrough" or column["input_used"]:
            input_fields.append({"namespace": namespace, "name": input_name, "field": column["name"], "transformations": [{"type": "DIRECT", "subtype": "IDENTITY", "description": labels["input_value"], "masking": False}]})
        for entry in column["upstream"]:
            input_fields.append({
                "namespace": namespace,
                "name": input_name if entry["column"] in raw else output_name,
                "field": entry["column"],
                "transformations": [{"type": t["type"], "subtype": t["subtype"], "description": _lines_label(entry["lines"], labels), "masking": False} for t in entry["transformations"]],
            })
        fields[column["name"]] = {"inputFields": input_fields, "transformationDescription": description}

    def schema(names: List[str]) -> Dict[str, Any]:
        return {"_producer": "https://github.com/dataflow-platform/technical-lineage-agent", "_schemaURL": "https://openlineage.io/spec/facets/1-1-1/SchemaDatasetFacet.json", "fields": [{"name": name} for name in names]}

    return {
        "eventType": "COMPLETE",
        "eventTime": state.get("finished_at") or _now(),
        "producer": "https://github.com/dataflow-platform/technical-lineage-agent",
        "schemaURL": "https://openlineage.io/spec/2-0-2/OpenLineage.json#/definitions/RunEvent",
        "run": {"runId": state["id"]},
        "job": {"namespace": namespace, "name": source.get("code_filename") or "script.py"},
        "inputs": [{"namespace": namespace, "name": input_name, "facets": {"schema": schema(analysis["inputs"])}}],
        "outputs": [{
            "namespace": namespace, "name": output_name,
            "facets": {
                "schema": schema(analysis["outputs"]),
                "columnLineage": {
                    "_producer": "https://github.com/dataflow-platform/technical-lineage-agent",
                    "_schemaURL": "https://openlineage.io/spec/facets/1-2-0/ColumnLineageDatasetFacet.json",
                    "fields": fields,
                    "dataset": [
                        {"namespace": namespace, "name": input_name, "field": column, "transformations": [{"type": "INDIRECT", "subtype": op["kind"], "description": _lines_label([op["line"]], labels) if op.get("line") else "", "masking": False}]}
                        for op in analysis.get("dataset_ops", []) for column in op["columns"]
                    ],
                },
            },
        }],
    }


@router.get("/runs/{run_id}/export")
def export_run(run_id: str, format: str = "xlsx", language: str = "en"):
    job = _JOBS.get(run_id)
    state = job.snapshot() if job is not None else _read_run(run_id)
    if state.get("status") != "completed" or not state.get("result"):
        raise HTTPException(409, "Only completed runs can be exported.")
    language = language if language in _LABELS else "en"
    stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", f"technical_lineage_{state['source'].get('cluster_name') or 'run'}_{run_id[:8]}_{language}")
    if format == "openlineage":
        content = json.dumps(_openlineage(state, language), ensure_ascii=False, indent=2).encode("utf-8")
        return Response(content, media_type="application/json", headers={"Content-Disposition": f'attachment; filename="{stem}.openlineage.json"'})
    return Response(
        _excel(state, language),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{stem}.xlsx"'},
    )
