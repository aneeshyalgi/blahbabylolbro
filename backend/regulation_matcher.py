"""Regulation–Release Note Matcher: links every release note of the selected files to the regulation provisions it
affects and returns a structured, verified traceability matrix.

Each release note runs through a fixed pipeline:

1. interpret  – the LLM states the change in regulatory terms and writes search queries in regulation language;
2. retrieve   – hybrid search (BM25F with glossary expansion plus embeddings) over the selected regulations, plus the
                provisions whose title names a concept of the note; candidates are whole paragraphs (or passages of
                very long paragraphs such as the definitions in Article 4(1));
3. adjudicate – the LLM links the note only to retrieved candidates (by id, so provisions cannot be invented), with a
                verbatim quote from both the provision and the release note;
4. verify     – deterministic checks: quotes verbatim, provision scope fits the note, no empowerment (mandate)
                clause, shared regulatory concept. Failing links go back to the LLM once with the reasons; links that
                still fail are rejected and kept in the audit trail.

Runs execute in a background thread; the browser polls their progress, so a run survives navigation and reloads.
Finished runs are stored as JSON and can be exported as an Excel audit workbook.
"""
from __future__ import annotations

import copy
import json
import os
import re
import threading
import time
import traceback
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

import agent_runtime as runtime
import regulation_text as rt

router = APIRouter(prefix="/api/regulation-matcher", tags=["regulation-matcher"])

RUNS_DIR = runtime.RESULTS_DIR / "regulation_matcher"
MAX_NOTES = 250
WORKERS = 4
CANDIDATE_LIMIT = 16
REFERENCE_SLOTS = 4
CONCEPT_SLOTS = 3
CANDIDATES_PER_UNIT = 3
PARAGRAPH_MAX_CHARS = 3000
QUERY_DEPTH = 60
RANK_K = 10
MAX_LINKS = 5
NOTE_TEXT_LIMIT = 4000
QUOTE_REPAIR_SIMILARITY = runtime.QUOTE_REPAIR_SIMILARITY
DEFAULT_OPENAI_MODEL = "gpt-4.1"  # strongest generally available model; REGULATION_MATCHER_MODEL overrides it
LANGUAGE_NAMES = {"en": "English", "de": "German", "es": "Spanish"}
# Every run's texts are available in these languages: the AI-written texts through a translation step, the
# deterministic log lines and verification messages through the templates below.
TEXT_LANGUAGES = ("en", "de")
_NO_CANDIDATES = {
    "en": "No provision of the selected regulations matched this release note.",
    "de": "Keine Vorschrift der ausgewählten Regulierungen passte zu dieser Release Note.",
}

_ID = re.compile(r"^[0-9a-f-]{8,64}$")
_JIRA = re.compile(r"\b(?:[A-Z][A-Z0-9]+-\d+|\d{5,7})\b")
_CORRELATION = re.compile(r"correlation table|entsprechungstabelle", re.IGNORECASE)
_KNOWN_ACTS = [
    ("575/2013", "CRR"), ("2013/36", "CRD"), ("2019/876", "CRR II"), ("2024/1623", "CRR III"), ("2024/1619", "CRD VI"),
    ("2015/61", "LCR DR"), ("2021/451", "ITS Reporting"), ("2016/867", "AnaCredit"), ("2017/2402", "Securitisation Reg."),
]


# --------------------------------------------------------------------------------------------------
# Texts in English and German
# --------------------------------------------------------------------------------------------------

# German names of the glossary concepts and exposure classes that the deterministic checks report.
CONCEPT_NAMES_DE = {
    "exposure value": "Risikopositionswert",
    "risk weight": "Risikogewicht",
    "risk-weighted exposure amount": "risikogewichteter Positionsbetrag",
    "credit conversion factor": "Umrechnungsfaktor",
    "guarantee": "Garantie",
    "undrawn commitment": "nicht in Anspruch genommene Zusage",
    "institution": "Institut",
    "corporate": "Unternehmen",
    "retail": "Mengengeschäft",
    "central government": "Zentralstaat",
    "regional government": "regionale Gebietskörperschaft",
    "public sector entity": "öffentliche Stelle",
    "immovable property": "Immobilieneigentum",
    "exposures in default": "ausgefallene Positionen",
    "specific credit risk adjustment": "spezifische Kreditrisikoanpassung",
    "general credit risk adjustment": "allgemeine Kreditrisikoanpassung",
    "accounting value": "Buchwert",
    "credit risk mitigation": "Kreditrisikominderung",
    "credit quality step": "Bonitätsstufe",
    "equity exposure": "Beteiligungsposition",
    "covered bond": "gedeckte Schuldverschreibung",
    "residual maturity": "Restlaufzeit",
    "own funds requirement": "Eigenmittelanforderung",
    "standardised approach": "Standardansatz",
    "internal ratings based approach": "IRB-Ansatz",
    "probability of default": "Ausfallwahrscheinlichkeit",
    "loss given default": "Verlustquote bei Ausfall",
    "exposure class": "Risikopositionsklasse",
    "loan": "Darlehen",
    "small and medium-sized enterprise": "KMU",
    "accrued interest": "aufgelaufene Zinsen",
    "nominal amount": "Nominalbetrag",
    "derivative": "Derivat",
    "leverage ratio": "Verschuldungsquote",
    "large exposures": "Großkredite",
    "liquidity coverage ratio": "Liquiditätsdeckungsquote",
    "net stable funding ratio": "strukturelle Liquiditätsquote",
}
_SCOPE_LABELS_DE = {
    "credit_obligation": "Darlehen (Kreditverpflichtungen)", "non_credit_obligation": "sonstige Aktiva ohne Kreditverpflichtung",
    "central_government": "Zentralstaaten", "regional_government": "regionale Gebietskörperschaften",
    "public_sector_entity": "öffentliche Stellen", "multilateral_development_bank": "multilaterale Entwicklungsbanken",
    "international_organisation": "internationale Organisationen", "institution": "Institute", "corporate": "Unternehmen",
    "retail": "Mengengeschäft", "immovable_property": "Immobilien", "default": "ausgefallene Positionen", "equity": "Beteiligungen",
    "covered_bond": "gedeckte Schuldverschreibungen", "ciu": "OGA", "securitisation": "Verbriefungen",
}


def _concept_label(name: str, language: str) -> str:
    return CONCEPT_NAMES_DE.get(name, name) if language == "de" else name


def _scope_list(classes: List[str], language: str) -> str:
    labels = _SCOPE_LABELS_DE if language == "de" else runtime._SCOPE_LABELS
    return ", ".join(sorted(labels.get(name, name) for name in classes))


# Verification problems: the English text goes back to the model in the correction round; both languages are
# shown on the page (corrections, reasons of rejected links) and in the export.
_PROBLEMS = {
    "unknown_candidate": {
        "en": "Link to '{candidate}' names no candidate id; use only the listed ids [C1]…",
        "de": "Der Bezug auf „{candidate}“ nennt keine Kandidaten-ID; zulässig sind nur die aufgeführten IDs [C1]…",
    },
    "quote_other_unit": {
        "en": "the regulation quote is from {located}, not from {reference}. Cite the candidate that holds it, or quote {reference} itself.",
        "de": "Das Regulierungszitat stammt aus {located}, nicht aus {reference}. Den Kandidaten zitieren, der es enthält, oder {reference} selbst zitieren.",
    },
    "quote_not_verbatim": {
        "en": "the regulation quote is not verbatim in the candidate text. Closest text: “{closest}”. Copy the sentence character for character.",
        "de": "Das Regulierungszitat steht nicht wörtlich im Kandidatentext. Nächstliegender Text: „{closest}“. Den Satz zeichengenau übernehmen.",
    },
    "note_quote_not_verbatim": {
        "en": "the release-note quote is not verbatim in the release note. Copy the words exactly as written, typos included.",
        "de": "Das Release-Note-Zitat steht nicht wörtlich in der Release Note. Die Wörter exakt wie geschrieben übernehmen, einschließlich Tippfehlern.",
    },
    "scope_conflict": {
        "en": "scope conflict — the release note concerns {note}, but '{title}' applies to {provision}. Link the provision that governs these exposures instead.",
        "de": "Konflikt im Anwendungsbereich – die Release Note betrifft {note}, „{title}“ gilt aber für {provision}. Stattdessen die Vorschrift verknüpfen, die diese Positionen regelt.",
    },
    "mandate": {
        "en": "this is an empowerment clause (technical standards / delegated act), not a substantive requirement.",
        "de": "Dies ist eine Ermächtigungsklausel (technische Standards / delegierter Rechtsakt), keine materielle Anforderung.",
    },
    "concept_missing": {
        "en": (
            "the provision text does not deal with {concepts}, which is what the release note changes. Drop the link unless "
            "the change really affects what this provision governs, and then link the provision that deals with the changed "
            "item instead."
        ),
        "de": (
            "Der Vorschriftentext behandelt nicht {concepts} – genau das ändert die Release Note. Den Bezug streichen, sofern "
            "die Änderung nicht tatsächlich berührt, was diese Vorschrift regelt, und stattdessen die Vorschrift verknüpfen, "
            "die den geänderten Posten behandelt."
        ),
    },
    "reference_gap": {
        "en": (
            "{source} ({source_reference}) applies its rule in accordance with {reference}, which is candidate [{target}] "
            "{target_reference}. Link [{target}] (direct when it sets the treatment of the changed item) or list it in rejected "
            "with the reason it does not apply."
        ),
        "de": (
            "{source} ({source_reference}) wendet seine Regel gemäß {reference} an – das ist Kandidat [{target}] {target_reference}. "
            "[{target}] verknüpfen (direkt, wenn er die Behandlung des geänderten Postens festlegt) oder unter den verworfenen "
            "Kandidaten mit Begründung aufführen."
        ),
    },
    "no_link_challenge": {
        "en": (
            "You found no link, but the release note concerns {concepts}, and these candidates deal with it: {listed}. "
            "Re-examine whether the change affects the quantity these provisions govern. Keep 'no link' only if none of them "
            "governs the changed quantity, and say why in no_link_reason. Return the complete JSON."
        ),
        "de": (
            "Es wurde kein Bezug gefunden, obwohl die Release Note {concepts} betrifft und diese Kandidaten dies behandeln: "
            "{listed}. Erneut prüfen, ob die Änderung die von diesen Vorschriften geregelte Größe berührt. „Kein Bezug“ nur "
            "beibehalten, wenn keine davon die geänderte Größe regelt, und dies in no_link_reason begründen. Das vollständige "
            "JSON erneut zurückgeben."
        ),
    },
}


def _problem(code: str, label: Optional[str] = None, **params: Any) -> Dict[str, Any]:
    """A verification problem in both languages: `texts` with the candidate label (sent to the model), `bodies`
    without it (the reason of a rejected link)."""
    texts: Dict[str, str] = {}
    bodies: Dict[str, str] = {}
    for language in TEXT_LANGUAGES:
        values = dict(params)
        if "concepts" in params:
            values["concepts"] = ", ".join(_concept_label(name, language) for name in params["concepts"])
        if language != "en" and "note_classes" in params:
            values["note"] = _scope_list(params["note_classes"], language)
            values["provision"] = _scope_list(params["provision_classes"], language)
        body = _PROBLEMS[code][language].format(**values)
        texts[language] = f"{label}: {body}" if label else body
        bodies[language] = body[:1].upper() + body[1:]
    return {"code": code, "text": texts["en"], "texts": texts, "bodies": bodies}


def _plural(count: int, one: str, other: str) -> str:
    return one if count == 1 else other


def _log_texts(code: str, p: Dict[str, Any]) -> Tuple[str, str]:
    """English and German text of a coded log line."""
    if code == "run_started":
        return (
            f"Run started: {p['notes']} release notes from {p['files']} file(s) against {p['regulations']} regulation(s).",
            f"Lauf gestartet: {p['notes']} {_plural(p['notes'], 'Release Note', 'Release Notes')} aus {p['files']} "
            f"{_plural(p['files'], 'Datei', 'Dateien')} gegen {p['regulations']} {_plural(p['regulations'], 'Regulierung', 'Regulierungen')}.",
        )
    if code == "load_indexes":
        return "Loading regulation indexes and search structures", "Regulierungsindizes und Suchstrukturen werden geladen"
    if code == "indexes_loaded":
        sources, semantic = p["sources"], p["semantic"]
        return (
            f"Loaded {len(sources)} regulation index(es): "
            + ", ".join(f"{item['short']} ({item['articles']} articles, {item['passages']} passages)" for item in sources)
            + (f"; semantic index: {', '.join(semantic)}" if semantic else "; keyword search only"),
            f"{len(sources)} {_plural(len(sources), 'Regulierungsindex', 'Regulierungsindizes')} geladen: "
            + ", ".join(f"{item['short']} ({item['articles']} Artikel, {item['passages']} Passagen)" for item in sources)
            + (f"; semantischer Index: {', '.join(semantic)}" if semantic else "; nur Stichwortsuche"),
        )
    if code == "model":
        return (
            f"Model: {p['model']} · temperature 0 · JSON-schema output · {p['workers']} notes in parallel",
            f"Modell: {p['model']} · Temperatur 0 · JSON-Schema-Ausgabe · {p['workers']} Notes parallel",
        )
    if code == "model_fallback":
        return (
            f"Model {p['model']} is not available; continuing with {p['fallback']}.",
            f"Modell {p['model']} ist nicht verfügbar; weiter mit {p['fallback']}.",
        )
    if code == "interpreted":
        terms = ", ".join(p["terms"])
        return f"Interpreted: {terms or 'no regulatory concept named'}", f"Interpretiert: {terms or 'kein regulatorischer Begriff genannt'}"
    if code == "retrieved":
        return (
            f"Retrieved {p['candidates']} candidate provisions from a pool of {p['pool']} ({p['queries']} queries"
            f"{', semantic + keyword' if p['semantic'] else ', keyword'})",
            f"{p['candidates']} Kandidatenvorschriften aus einem Pool von {p['pool']} gefunden ({p['queries']} Suchanfragen, "
            f"{'semantisch + Stichwort' if p['semantic'] else 'Stichwort'})",
        )
    if code == "correction_round":
        en, de = p["texts"]["en"], p["texts"]["de"]
        return (
            f"Correction round: {len(en)} issue(s) sent back — " + " | ".join(item[:160] for item in en),
            f"Korrekturrunde: {len(de)} {_plural(len(de), 'Problem', 'Probleme')} an das Modell zurückgemeldet — "
            + " | ".join(item[:160] for item in de),
        )
    if code == "verification_rejected":
        return (
            f"{p['count']} link(s) rejected after verification",
            f"{p['count']} {_plural(p['count'], 'Bezug', 'Bezüge')} nach der Prüfung verworfen",
        )
    if code == "review":
        if p["escalated"]:
            return (
                "Four-eyes review: reviewer disagrees with every link – escalated for human review",
                "Vier-Augen-Prüfung: Der Prüfer widerspricht jedem Bezug – zur manuellen Prüfung eskaliert",
            )
        en_parts = [part for part in (
            f"rejected {', '.join(p['rejected'])}" if p["rejected"] else "",
            f"downgraded {', '.join(p['downgraded'])} to indirect" if p["downgraded"] else "",
        ) if part]
        de_parts = [part for part in (
            f"{', '.join(p['rejected'])} verworfen" if p["rejected"] else "",
            f"{', '.join(p['downgraded'])} auf indirekt herabgestuft" if p["downgraded"] else "",
        ) if part]
        confirmed = p["confirmed"]
        return (
            "Four-eyes review: " + (", ".join(en_parts) or f"all {confirmed} link(s) confirmed"),
            "Vier-Augen-Prüfung: " + ("; ".join(de_parts) or ("Bezug bestätigt" if confirmed == 1 else f"alle {confirmed} Bezüge bestätigt")),
        )
    if code == "result":
        links = p["links"]
        return (
            "Result: " + (", ".join(f"{item['reference']} ({item['link_type']})" for item in links) or "no link"),
            "Ergebnis: " + (", ".join(
                f"{item['reference']} ({'direkt' if item['link_type'] == 'direct' else 'indirekt'})" for item in links
            ) or "kein Bezug"),
        )
    if code == "translation_failed":
        return (
            f"Translation failed: {p['error']} – the texts of this note stay in their original language",
            f"Übersetzung fehlgeschlagen: {p['error']} – die Texte dieser Note bleiben in der Originalsprache",
        )
    if code == "translated":
        kept_en = f", {p['kept']} kept in the original language after failed checks" if p["kept"] else ""
        kept_de = f", {p['kept']} nach nicht bestandener Prüfung in der Originalsprache belassen" if p["kept"] else ""
        return (
            f"Texts translated for {p['notes']} release notes ({p['texts']} texts{kept_en}); every result is available in English and German.",
            f"Texte für {p['notes']} {_plural(p['notes'], 'Release Note', 'Release Notes')} übersetzt ({p['texts']} Texte{kept_de}); "
            "jedes Ergebnis liegt auf Deutsch und Englisch vor.",
        )
    if code == "note_failed":
        return f"Failed: {p['error']}", f"Fehlgeschlagen: {p['error']}"
    if code in ("run_completed", "run_cancelled"):
        s = p["summary"]
        cancelled = code == "run_cancelled"
        return (
            ("Run cancelled. " if cancelled else "Run completed. ")
            + f"{s['links']} links ({s['direct']} direct, {s['indirect']} indirect) to {s['provisions']} provisions; "
            f"{s['quotes_verified']}/{s['quotes_total']} quotes verified; coverage {s['coverage']}%.",
            ("Lauf abgebrochen. " if cancelled else "Lauf abgeschlossen. ")
            + f"{s['links']} {_plural(s['links'], 'Bezug', 'Bezüge')} ({s['direct']} direkt, {s['indirect']} indirekt) zu "
            f"{s['provisions']} {_plural(s['provisions'], 'Vorschrift', 'Vorschriften')}; {s['quotes_verified']}/{s['quotes_total']} "
            f"Zitate verifiziert; Abdeckung {s['coverage']} %.",
        )
    if code == "cancel_requested":
        return (
            "Cancellation requested; notes in progress finish their current step.",
            "Abbruch angefordert; laufende Notes beenden ihren aktuellen Schritt.",
        )
    if code == "run_failed":
        return f"Run failed: {p['error']}", f"Lauf fehlgeschlagen: {p['error']}"
    return str(p.get("error") or code), str(p.get("error") or code)


# --------------------------------------------------------------------------------------------------
# Sources
# --------------------------------------------------------------------------------------------------

def _short_name(document: Dict[str, Any]) -> str:
    """'CRR' for Regulation (EU) No 575/2013, else a short form of the title or file name."""
    title = str(document.get("document_title") or "")
    haystack = f"{title} {document.get('filename') or ''}"
    for number, short in _KNOWN_ACTS:
        if number in haystack or number.replace("/", "_") in haystack:
            return short
    match = re.search(r"\b(Regulation|Directive|Verordnung|Richtlinie)\s*\((EU|EG|EC)\)\s*(?:No\s*|Nr\.\s*)?(\d{2,4}/\d{1,4})", title, re.IGNORECASE)
    if match:
        return f"{match.group(1).title()} {match.group(3)}"
    return Path(str(document.get("filename") or "Regulation")).stem[:28]


def _text(value: Any) -> str:
    if runtime._is_missing(value):
        return ""
    return re.sub(r"[ \t]+", " ", str(value)).strip()


def _note_from_fields(context: Dict[str, Any], sheet: str, ordinal: int, record: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    fields = [{"name": str(key), "value": _text(value)} for key, value in record.items() if _text(value)]
    if not fields:
        return None
    jira = runtime._record_field(record, "jira")
    if not jira:
        found = _JIRA.search(" ".join(field["value"] for field in fields))
        jira = found.group(0) if found else ""
    problem = runtime._record_field(record, "problem")
    solution = runtime._record_field(record, "solution")
    described = [value for value in (problem, solution) if value]
    if not described:
        described = [field["value"] for field in fields if field["name"].lower() not in {"page", "jira id"}]
    text = "\n".join(described)
    if len(re.sub(r"\W", "", text)) < 12:
        return None
    return {
        "key": f"{context.get('id')}:{sheet}:{ordinal}",
        "jira_id": str(jira),
        "file_id": str(context.get("id")),
        "file": str(context.get("filename") or ""),
        "kind": context.get("kind") or "excel",
        "sheet": sheet,
        "ordinal": ordinal,
        "fields": fields,
        "problem": problem,
        "solution": solution,
        "text": text[:NOTE_TEXT_LIMIT],
        "full_text": " ".join(field["value"] for field in fields),
    }


def _notes_from_context(context: Dict[str, Any]) -> List[Dict[str, Any]]:
    notes: List[Dict[str, Any]] = []
    if (context.get("kind") or "excel") == "pdf":
        # PDF passages: a passage naming a Jira ID starts a note; following passages without one continue it.
        current: Optional[Dict[str, Any]] = None
        ordinal = 0
        for sheet in context.get("sheets", []):
            for record in sheet.get("records", []):
                text = _text(record.get("Text"))
                if not text:
                    continue
                jira = _text(record.get("Jira ID"))
                if current is not None and (not jira or jira == current["jira"]) and len(current["text"]) + len(text) < NOTE_TEXT_LIMIT:
                    current["text"] += "\n" + text
                    continue
                ordinal += 1
                current = {"jira": jira, "page": record.get("Page"), "sheet": sheet.get("name"), "text": text, "ordinal": ordinal}
                notes.append(current)
        built = []
        identified = any(item["jira"] for item in notes)
        for item in notes:
            if identified and not item["jira"] and len(item["text"]) < 160:
                continue  # a heading such as 'Release 4.2 – IReF Kalkulator' in a file whose notes carry Jira IDs
            record = {"Page": item["page"], **({"Jira ID": item["jira"]} if item["jira"] else {}), "Text": item["text"]}
            note = _note_from_fields(context, str(item["sheet"]), item["ordinal"], record)
            if note:
                built.append(note)
        return built
    for sheet in context.get("sheets", []):
        for ordinal, record in enumerate(sheet.get("records", []), start=1):
            note = _note_from_fields(context, str(sheet.get("name")), ordinal, record)
            if note:
                notes.append(note)
    return notes


def _release_note_files(file_ids: Optional[List[str]] = None) -> List[Tuple[Dict[str, Any], List[Dict[str, Any]]]]:
    contexts = runtime._release_note_contexts(set(file_ids) if file_ids else None)
    if file_ids:
        order = {file_id: position for position, file_id in enumerate(file_ids)}
        contexts.sort(key=lambda context: order.get(str(context.get("id")), 0))
    return [(context, _notes_from_context(context)) for context in contexts]


def _ready_regulations(regulation_ids: List[str]) -> List[Dict[str, Any]]:
    """The selected regulation documents, or an HTTP error when one is missing or not indexed yet."""
    by_id = {str(item.get("id")): item for item in runtime._platform("list_regulation_documents")()}
    documents = []
    for regulation_id in regulation_ids:
        document = by_id.get(regulation_id)
        if not document:
            raise HTTPException(404, "A selected regulation no longer exists. Refresh the page and select again.")
        status = document.get("index_status")
        if status != "ready" or document.get("index_version") != rt.INDEX_VERSION:
            reason = "failed to index" if status == "failed" else "is still being indexed"
            raise HTTPException(409, f"'{document.get('filename')}' {reason}. Wait until it is ready in the Regulations tab.")
        documents.append(document)
    return documents


def _regulation_sources(regulation_ids: List[str]) -> List[Dict[str, Any]]:
    sources = []
    for document in _ready_regulations(regulation_ids):
        regulation_id = str(document.get("id"))
        try:
            index = runtime._regulation_index_for(document)
        except runtime.ToolError as exc:
            raise HTTPException(409, str(exc)) from exc
        by_unit: Dict[int, List[int]] = {}
        for position, passage in enumerate(index["passages"]):
            if passage["kind"] != "footnote":
                by_unit.setdefault(passage["unit"], []).append(position)
        sources.append({
            "id": regulation_id,
            "document": document,
            "index": index,
            "search": runtime._regulation_search_index(document, index),
            "embeddings": runtime._regulation_embeddings_for(document),
            "short": _short_name(document),
            "by_unit": by_unit,
        })
    return sources


# --------------------------------------------------------------------------------------------------
# LLM access
# --------------------------------------------------------------------------------------------------

class _LLM:
    """JSON-schema chat calls with retries, a model fallback and usage accounting."""

    def __init__(self, job: "_Job"):
        client, configured = runtime._llm_client()
        override = os.environ.get("REGULATION_MATCHER_MODEL", "").strip()
        uses_openai = bool(os.environ.get("OPENAI_API_KEY"))
        self.client = client
        self.model = override or (DEFAULT_OPENAI_MODEL if uses_openai else configured)
        self.fallback = configured if self.model != configured else None
        self.job = job
        self.strict = True
        self.seed: Optional[int] = 20240611
        self.lock = threading.Lock()

    def json(self, messages: List[Dict[str, str]], name: str, schema: Dict[str, Any]) -> Dict[str, Any]:
        last: Optional[Exception] = None
        for attempt in range(4):
            if self.job.cancel.is_set():
                raise _Cancelled()
            response_format: Dict[str, Any] = (
                {"type": "json_schema", "json_schema": {"name": name, "strict": True, "schema": schema}}
                if self.strict else {"type": "json_object"}
            )
            options: Dict[str, Any] = {"seed": self.seed} if self.seed is not None else {}
            try:
                response = self.client.chat.completions.create(
                    model=self.model, messages=messages, temperature=0, response_format=response_format, **options,
                )
            except Exception as exc:  # provider errors: fall back, degrade or retry
                last = exc
                status = getattr(exc, "status_code", None)
                if status in (403, 404) and self.fallback:
                    with self.lock:
                        if self.fallback:
                            self.job.say("warn", "model_fallback", model=self.model, fallback=self.fallback)
                            self.model, self.fallback = self.fallback, None
                            self.job.set_model(self.model)
                    continue
                if status == 400 and self.seed is not None and "seed" in str(exc):
                    self.seed = None
                    continue
                if status == 400 and self.strict and ("response_format" in str(exc) or "json_schema" in str(exc)):
                    self.strict = False
                    continue
                if status in (408, 409, 429) or (status or 0) >= 500 or status is None:
                    time.sleep(1.5 * (attempt + 1))
                    continue
                raise
            usage = getattr(response, "usage", None)
            self.job.count_call(getattr(usage, "prompt_tokens", 0) or 0, getattr(usage, "completion_tokens", 0) or 0)
            content = response.choices[0].message.content or "{}"
            try:
                return json.loads(content)
            except json.JSONDecodeError as exc:
                last = exc
                continue
        raise RuntimeError(f"The language model did not return a valid answer: {last}")


class _Cancelled(Exception):
    pass


def _string_array() -> Dict[str, Any]:
    return {"type": "array", "items": {"type": "string"}}


INTERPRET_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["summary", "summary_en", "change_kind", "affected_items", "concepts", "queries"],
    "properties": {
        "summary": {"type": "string"},
        "summary_en": {"type": "string"},
        "change_kind": {"type": "string", "enum": ["calculation", "input_data", "classification", "reporting", "methodology", "other"]},
        "affected_items": _string_array(),
        "concepts": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["term", "source_phrase", "effect"],
                "properties": {"term": {"type": "string"}, "source_phrase": {"type": "string"}, "effect": {"type": "string"}},
            },
        },
        "queries": _string_array(),
    },
}

DECISION_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["assessment", "links", "rejected", "no_link_reason"],
    "properties": {
        "assessment": {"type": "string"},
        "links": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": [
                    "candidate", "link_type", "confidence", "approach", "affected_element", "rationale",
                    "regulation_quote", "release_note_quote",
                ],
                "properties": {
                    "candidate": {"type": "string"},
                    "link_type": {"type": "string", "enum": ["direct", "indirect"]},
                    "confidence": {"type": "integer"},
                    "approach": {"type": "string", "enum": ["standardised", "irb", "both", "not_specific"]},
                    "affected_element": {"type": "string"},
                    "rationale": {"type": "string"},
                    "regulation_quote": {"type": "string"},
                    "release_note_quote": {"type": "string"},
                },
            },
        },
        "rejected": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["candidate", "reason"],
                "properties": {"candidate": {"type": "string"}, "reason": {"type": "string"}},
            },
        },
        "no_link_reason": {"type": "string"},
    },
}

REVIEW_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["reviews"],
    "properties": {
        "reviews": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["link", "verdict", "reason"],
                "properties": {
                    "link": {"type": "string"},
                    "verdict": {"type": "string", "enum": ["confirm", "downgrade", "reject"]},
                    "reason": {"type": "string"},
                },
            },
        },
    },
}


def _review_messages(note: Dict[str, Any], interpretation: Dict[str, Any], links: List[Dict[str, Any]], language: str) -> List[Dict[str, str]]:
    name = LANGUAGE_NAMES.get(language, "English")
    system = (
        "You are the independent second reviewer (four-eyes principle) of a regulatory traceability analysis. A first "
        "analyst linked a bank release note to provisions of EU banking regulation. Review every proposed link on its "
        "own merits, from the provision text:\n"
        "- confirm: the change in the release note affects what the provision governs, for the release note's portfolio "
        "and purpose, and the grade is right;\n"
        "- downgrade: the link is graded direct, but the provision is affected only indirectly (it uses a value the change "
        "alters, or it applies only under an approach or use the release note does not mention);\n"
        "- reject: the provision is not affected – it uses the same words for another purpose (e.g. reference data for "
        "IRB conversion-factor estimates, securitisation, market risk, liquidity or disclosure) or its scope does not cover "
        "the release note's exposures.\n"
        "Direct means the provision defines, sets or prescribes the quantity, input or treatment the release note changes; "
        "indirect means it uses or consumes that changed value. Do not reject a link merely because it is indirect, and do "
        f"not reward shared vocabulary. Give one review per link with a one-sentence reason in {name}. {_DOMAIN_TERMS}"
    )
    parts = [_note_block(note), "", f"Analyst interpretation: {interpretation.get('summary_en') or ''}", "", "Proposed links:"]
    for index, link in enumerate(links, start=1):
        parts.append(
            f"[L{index}] {link['regulation']} {link['reference']} — {link['title']} ({' › '.join(link['path'][-3:])})\n"
            f"Graded: {link['link_type']} · approach: {link['approach']} · affected element: {link['affected_element']}\n"
            f"Analyst rationale: {link['rationale']}\n"
            f"Provision text: {link['provision_text'][:1600]}"
        )
        parts.append("")
    return [{"role": "system", "content": system}, {"role": "user", "content": "\n".join(parts)}]


def _review(
    llm: "_LLM", note: Dict[str, Any], interpretation: Dict[str, Any], links: List[Dict[str, Any]], language: str,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], bool]:
    """Apply the four-eyes review: confirmed and downgraded links stay, rejected ones move to the rejected list. When
    the reviewer rejects every link, the links stay but the note is escalated for human review (returns True)."""
    result = llm.json(_review_messages(note, interpretation, links, language), "link_review", REVIEW_SCHEMA)
    verdicts: Dict[int, Dict[str, str]] = {}
    for item in result.get("reviews") or []:
        match = re.search(r"\d+", str(item.get("link") or ""))
        if match:
            verdicts[int(match.group(0)) - 1] = {"verdict": str(item.get("verdict") or "confirm"), "reason": str(item.get("reason") or "").strip()}
    rejections = [position for position, verdict in verdicts.items() if verdict["verdict"] == "reject" and position < len(links)]
    escalate = bool(links) and len(rejections) == len(links)
    kept: List[Dict[str, Any]] = []
    dismissed: List[Dict[str, Any]] = []
    for position, link in enumerate(links):
        verdict = verdicts.get(position, {"verdict": "confirm", "reason": ""})
        if verdict["verdict"] == "reject" and not escalate:
            dismissed.append({
                "candidate": link["candidate"], "reference": link["reference"], "title": link["title"],
                "regulation": link["regulation"], "page_label": link["page_label"], "source": "review", "reason": verdict["reason"],
            })
            continue
        if verdict["verdict"] == "downgrade" and link["link_type"] == "direct":
            link["link_type"] = "indirect"
            link["confidence"] = max(0, link["confidence"] - 8)
            link["checks"].append({"id": "review", "status": "corrected", "params": {"reason": verdict["reason"]}})
        elif verdict["verdict"] == "reject":
            link["confidence"] = min(link["confidence"], 55)
            link["checks"].append({"id": "review", "status": "warn", "params": {"reason": verdict["reason"]}})
        else:
            link["checks"].append({"id": "review", "status": "pass", "params": {"reason": verdict["reason"]}})
        link["band"] = "high" if link["confidence"] >= 80 else "medium" if link["confidence"] >= 60 else "low"
        kept.append(link)
    kept.sort(key=lambda link: (link["link_type"] != "direct", -link["confidence"]))
    return kept, dismissed, escalate


_DOMAIN_TERMS = (
    "Release notes are often German and use banking business terms; translate them into the vocabulary of EU banking "
    "regulation. Examples: 'Anteilige Zinsen' = accrued interest; 'Carrying Amount' / 'Buchwert' = carrying amount, the "
    "accounting value of an asset; 'Einzelwertberichtigungen (EWB)' = specific credit risk adjustments; "
    "'Pauschalwertberichtigungen (PWB)' = general credit risk adjustments; 'Darlehen an Private Haushalte' = loans to "
    "households (retail credit obligations); 'Betriebsmitteldarlehen an Unternehmen' = working capital loans to "
    "corporates; 'Datenlieferung' / 'angeliefert' = input data delivery; 'Sicherheiten' = collateral (credit risk "
    "mitigation); 'Zusage' / 'Kreditlinie' = commitments (off-balance-sheet items)."
)


def _interpret_messages(note: Dict[str, Any], language: str) -> List[Dict[str, str]]:
    name = LANGUAGE_NAMES.get(language, "English")
    system = (
        "You are a senior regulatory reporting analyst at a European bank (CRR/CRD, IReF, AnaCredit, COREP/FINREP). "
        "You read one release note of the bank's regulatory calculation engine and state precisely, in regulatory terms, "
        f"what changed. {_DOMAIN_TERMS}\n\n"
        "Return JSON with:\n"
        f"- summary: one or two sentences in {name}: what changed and which regulatory quantity or input it affects.\n"
        "- summary_en: the same in English, in regulation vocabulary.\n"
        "- change_kind: calculation | input_data | classification | reporting | methodology | other.\n"
        "- affected_items: the instruments, counterparties or portfolios concerned, in English regulation vocabulary "
        "(e.g. 'loans to households (retail exposures)').\n"
        "- concepts: each regulatory concept the change touches: term (English regulation term), source_phrase (the exact "
        f"words of the release note, verbatim, in its original language), effect (what the change does to it, in {name}).\n"
        "- queries: 4 to 6 search queries in English regulation language that find the provisions governing the affected "
        "quantity or input, phrased the way the regulation states the rule (e.g. 'exposure value of an asset item is its "
        "accounting value remaining after specific credit risk adjustments'), not as questions. Cover (a) the provision "
        "that defines the affected quantity, (b) provisions in which the changed value feeds another calculation (e.g. "
        "risk weights that depend on the amount of specific credit risk adjustments, expected-loss comparisons, own "
        "funds items) and (c) the treatment prescribed for the changed input itself."
    )
    return [{"role": "system", "content": system}, {"role": "user", "content": _note_block(note)}]


def _note_block(note: Dict[str, Any]) -> str:
    lines = [f"Release note (file '{note['file']}', {note['sheet']}):"]
    for field in note["fields"]:
        lines.append(f"- {field['name']}: {field['value'][:1500]}")
    return "\n".join(lines)


def _decision_messages(note: Dict[str, Any], interpretation: Dict[str, Any], candidates: List[Dict[str, Any]], language: str) -> List[Dict[str, str]]:
    name = LANGUAGE_NAMES.get(language, "English")
    system = (
        "You are a senior regulatory reporting specialist performing a traceability review: which provisions of the "
        "selected regulations does this release note affect? You decide only among the numbered candidate provisions, "
        "which were retrieved from the regulation text.\n\n"
        "Rules:\n"
        "1. Link a candidate only when the change in the release note affects what the provision governs. Judge from the "
        "provision text, not from shared keywords. Follow the reference chain: when a provision you link prescribes the "
        "treatment by reference to another candidate ('shall treat general credit risk adjustments in accordance with "
        "Article 62(c)'), link that candidate too, because it holds the substantive rule; that a linked provision refers to it "
        "is a reason to link it, never a reason to reject it.\n"
        "2. link_type: direct = the provision defines, sets or prescribes the quantity, input or treatment the release note "
        "changes (e.g. the note adds accrued interest to the carrying amount of loans, and the provision defines exposure "
        "value as the accounting value of an asset item); indirect = the note changes an upstream input of a quantity the "
        "provision uses, or the provision prescribes a treatment that consumes the changed value.\n"
        "3. Prefer the provision that defines the affected quantity itself over provisions that only use the same words in "
        "another context; a provision that mentions the same item for another purpose is not affected (e.g. a note on "
        "the accrued interest in the carrying amount does not affect the reference data for IRB conversion-factor "
        "estimates, securitisation positions or liquidity ratios). Never link an empowerment clause ('EBA shall develop "
        "draft regulatory technical standards…', "
        "'power is delegated to the Commission…'), a review clause or a correlation table.\n"
        "4. Check scope: the provision's title and position (part, title, chapter) must apply to what the release note "
        "affects. Loans to households are retail credit obligations, so an article on 'Other non credit-obligation assets' "
        "does not govern them; an article on exposures to institutions does not govern loans to corporates; general credit "
        "risk adjustments are not specific credit risk adjustments. A provision that expressly leaves the changed item "
        "out (e.g. 'measured without taking into account any credit risk adjustments' for a change to credit risk "
        "adjustments) is not affected by the change: do not link it.\n"
        "5. When the standardised approach and the IRB approach treat the quantity differently and both are candidates, link "
        "both and set approach accordingly; otherwise approach = not_specific unless the provision is approach-specific.\n"
        f"6. At most {MAX_LINKS} links, most important first. Do not force a link: if no candidate governs the change, "
        "return no links and explain why in no_link_reason.\n"
        "7. regulation_quote: copy VERBATIM, character for character, the shortest complete sentence or list item of the "
        "candidate text that establishes the link (at most about 400 characters). Do not correct, translate or shorten "
        "words; only use '…' to skip text inside the same candidate.\n"
        "8. release_note_quote: copy VERBATIM the words of the release note that show the change, in the original "
        "language, including typos and double spaces.\n"
        f"9. rationale: 2 to 4 sentences in {name}: what changes in the bank's data or calculation, which regulatory "
        "quantity it feeds, and how the provision governs that quantity. Explain German business terms.\n"
        f"10. affected_element: a short label in {name} for the regulatory element affected (e.g. 'Exposure value – "
        "accounting value after specific credit risk adjustments').\n"
        "11. confidence: 0 to 100, how certain the link is given the provision text.\n"
        f"12. rejected: for the strongest candidates you did not link (up to 6), one sentence in {name} on why not.\n"
        f"13. assessment: one or two sentences in {name} on which quantity or input the change affects."
    )
    concept_lines = "; ".join(
        f"{item.get('term')} ('{item.get('source_phrase')}')" for item in interpretation.get("concepts", [])[:8]
    )
    parts = [
        _note_block(note),
        "",
        f"Analyst interpretation: {interpretation.get('summary_en') or ''}",
        f"Affected items: {', '.join(interpretation.get('affected_items', [])[:6]) or '—'}",
        f"Regulatory concepts: {concept_lines or '—'}",
        "",
        "Candidate provisions:",
    ]
    for candidate in candidates:
        location = " › ".join(candidate["path"][-3:]) if candidate["path"] else ""
        parts.append(
            f"[{candidate['id']}] {candidate['regulation']} · {candidate['reference']}"
            f"{' — ' + candidate['title'] if candidate['title'] else ''}\n"
            f"Location: {location or '—'} · page {candidate['page_label']}\n"
            f"Text: {candidate['text']}"
        )
        parts.append("")
    return [{"role": "system", "content": system}, {"role": "user", "content": "\n".join(parts)}]


# --------------------------------------------------------------------------------------------------
# Retrieval
# --------------------------------------------------------------------------------------------------

def _embed_queries(queries: List[str]) -> Dict[str, Any]:
    """Unit vectors for the queries (one embedding request), or {} when semantic search is unavailable."""
    import numpy as np

    client, model = rt.embedding_client()
    if client is None or not queries:
        return {}
    try:
        vectors = rt._embed(client, model, queries)
    except Exception:
        return {}
    result = {}
    for query, vector in zip(queries, vectors):
        array = np.asarray(vector[: rt.EMBEDDING_DIMENSIONS], dtype=np.float32)
        norm = float(np.linalg.norm(array)) or 1.0
        result[query] = array / norm
    return result


def _candidate_for(source: Dict[str, Any], passage_position: int) -> Optional[Dict[str, Any]]:
    index = source["index"]
    passage = index["passages"][passage_position]
    unit_position = passage["unit"]
    unit = index["units"][unit_position]
    if passage["kind"] == "footnote" or unit["kind"] not in ("article", "annex") or _CORRELATION.search(unit.get("title") or ""):
        return None
    members = source["by_unit"].get(unit_position, [])
    paragraph = str(passage.get("paragraph") or "")
    chosen: List[int] = [passage_position]
    reference = rt.passage_reference(unit, passage)
    key: Tuple[Any, ...] = (source["id"], unit_position, "i", passage_position)
    if unit["kind"] == "article":
        if paragraph:
            group = [i for i in members if str(index["passages"][i].get("paragraph") or "") == paragraph]
            if sum(len(index["passages"][i]["text"]) for i in group) <= PARAGRAPH_MAX_CHARS:
                chosen, key = group, (source["id"], unit_position, "p", paragraph)
        elif sum(len(index["passages"][i]["text"]) for i in members) <= PARAGRAPH_MAX_CHARS:
            chosen, key, reference = members, (source["id"], unit_position, "u", ""), unit["label"]
    points: List[str] = []
    for i in chosen:
        for point in index["passages"][i].get("points") or []:
            if point not in points:
                points.append(point)
    if key[2] == "i" and paragraph and points:
        # A slice of a long paragraph: name the points it covers, e.g. 'Article 4(1)(94)–(96)'.
        reference += f"({points[0]})" if len(points) == 1 else f"({points[0]})–({points[-1]})"
    pages = sorted({page for i in chosen for page in range(index["passages"][i]["page"], index["passages"][i].get("page_end", index["passages"][i]["page"]) + 1)})
    text = " ".join(index["passages"][i]["text"] for i in chosen)
    return {
        "key": key,
        "regulation_id": source["id"],
        "regulation": source["short"],
        "regulation_file": source["document"].get("filename"),
        "unit": unit_position,
        "kind": unit["kind"],
        "label": unit["label"],
        "number": unit.get("number"),
        "title": unit.get("title") or "",
        "path": list(unit.get("path") or []),
        "reference": reference,
        "paragraph": paragraph or None,
        "points": points,
        "passages": chosen,
        "pages": pages,
        "page": pages[0] if pages else None,
        "page_label": f"{pages[0]}–{pages[-1]}" if len(pages) > 1 else str(pages[0] if pages else ""),
        "text": text,
        "formula": any(index["passages"][i].get("formula") for i in chosen),
    }


def _flags(candidate: Dict[str, Any], note_text: str) -> Dict[str, bool]:
    return {
        "mandate": bool(runtime._MANDATE_CLAUSE.search(candidate["text"])) and len(candidate["text"]) < 1500,
        "side": bool(runtime._SIDE_PROVISION.match(candidate["title"])),
        "scope": bool(runtime._scope_conflict(note_text, [{"title": candidate["title"]}])),
    }


def _demotion(flags: Dict[str, bool]) -> float:
    factor = 1.0
    if flags["mandate"]:
        factor *= 0.35
    if flags["scope"]:
        factor *= 0.35
    if flags["side"]:
        factor *= 0.6
    return factor


def _referenced_candidate(source: Dict[str, Any], reference: str) -> Optional[Dict[str, Any]]:
    """The candidate for a cross-reference such as 'Article 62(c)' or 'Article 36(1), point (m)'."""
    request = rt.parse_unit_request(re.sub(r",?\s*points?\s*", "", reference))
    if not request or request["kind"] != "article":
        return None
    index = source["index"]
    for unit_position in rt.find_units(index, request):
        members = source["by_unit"].get(unit_position, [])
        if not members:
            continue
        passages = index["passages"]
        paragraph, point = request.get("paragraph"), request.get("point")
        pick = [i for i in members if paragraph and str(passages[i].get("paragraph") or "") == paragraph]
        if point:
            pick = [i for i in (pick or members) if point in (passages[i].get("points") or [])] or pick
        return _candidate_for(source, (pick or members)[0])
    return None


def _retrieve(note: Dict[str, Any], interpretation: Dict[str, Any], sources: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """The candidate provisions for one note, best first, with their retrieval signals."""
    queries: List[Tuple[str, float]] = []
    for query in interpretation.get("queries", [])[:6]:
        if str(query).strip():
            queries.append((str(query).strip()[:300], 1.0))
    if interpretation.get("summary_en"):
        queries.append((str(interpretation["summary_en"])[:400], 0.8))
    queries.append((note["text"][:600], 1.0))
    terms = [str(item.get("term") or "") for item in interpretation.get("concepts", []) if item.get("term")]
    if terms:
        queries.append((" ; ".join(terms)[:300], 0.7))
    queries = list({query: (query, weight) for query, weight in queries}.values())
    vectors = _embed_queries([query for query, _ in queries]) if any(source["embeddings"] is not None for source in sources) else {}

    scores: Dict[Tuple[Any, ...], float] = {}
    found: Dict[Tuple[Any, ...], Dict[str, Any]] = {}
    signals: Dict[Tuple[Any, ...], Dict[str, Any]] = {}

    def add(candidate: Dict[str, Any], rank: int, weight: float, signal: Dict[str, Any]) -> None:
        key = candidate["key"]
        found.setdefault(key, candidate)
        scores[key] = scores.get(key, 0.0) + weight / (RANK_K + rank)
        entry = signals.setdefault(key, {
            "best_rank": rank, "keyword_rank": None, "similarity": None, "queries": 0, "via_concept": False,
            "concept_rank": None, "via_reference": None,
        })
        entry["best_rank"] = min(entry["best_rank"], rank)
        for name in ("keyword_rank",):
            if signal.get(name) is not None:
                entry[name] = signal[name] if entry[name] is None else min(entry[name], signal[name])
        if signal.get("semantic_similarity") is not None:
            entry["similarity"] = max(entry["similarity"] or 0.0, signal["semantic_similarity"])
        if signal.get("via_concept"):
            entry["via_concept"] = True
            entry["concept_rank"] = rank if entry["concept_rank"] is None else min(entry["concept_rank"], rank)
        if signal.get("via_reference") and not entry.get("via_reference"):
            entry["via_reference"] = signal["via_reference"]

    for query, weight in queries:
        groups = []
        for source in sources:
            candidates, _ = rt.search_candidates(
                source["index"], source["search"], query, source["embeddings"], vectors.get(query), depth=QUERY_DEPTH,
            )
            groups.append((source, candidates))
        seen: Set[Tuple[Any, ...]] = set()
        rank = 0
        for source, passage_position, _, signal in rt.fuse_candidates(groups)[: QUERY_DEPTH * 2]:
            candidate = _candidate_for(source, passage_position)
            if candidate is None or candidate["key"] in seen:
                continue
            seen.add(candidate["key"])
            rank += 1
            add(candidate, rank, weight, signal)
            if candidate["key"] in signals:
                signals[candidate["key"]]["queries"] += 1
            if rank >= QUERY_DEPTH:
                break

    # Provisions whose title names a concept of the note (e.g. 'Treatment of credit risk adjustment').
    own = rt.concepts(note["full_text"])
    about = own | rt.concepts(" ".join(terms))
    per_concept: List[List[Tuple[Dict[str, Any], int]]] = []
    for position in sorted(about, key=lambda item: (item not in own, item)):
        phrases = [set(rt.tokens(phrase)) for phrase in rt.GLOSSARY[position] if len(rt.tokens(phrase)) >= 2]
        ranked = []
        for source in sources:
            index = source["index"]
            for item in rt.concept_evidence(index, {position}, limit=25):
                unit = index["units"][item["unit"]]
                if runtime._MANDATE_CLAUSE.search(item["sentence"]):
                    continue
                density = item["mentions"] / max(1.0, (unit.get("chars") or 1000) / 1000)
                side = bool(runtime._SIDE_PROVISION.match(unit.get("title") or ""))
                title_words = set(rt.tokens(unit.get("title") or ""))
                title_match = int(any(phrase <= title_words for phrase in phrases))
                passage_position = next(
                    (i for i in source["by_unit"].get(item["unit"], []) if item["sentence"][:60] in index["passages"][i]["text"]),
                    None,
                )
                if passage_position is not None:
                    ranked.append(((side, -title_match, -density), source, passage_position))
        ranked.sort(key=lambda entry: entry[0])
        per_concept.append([(source, passage_position) for _, source, passage_position in ranked[:8]])
    rank = 0
    for depth in range(8):
        for ranked_list in per_concept:
            if depth < len(ranked_list):
                source, passage_position = ranked_list[depth]
                candidate = _candidate_for(source, passage_position)
                if candidate is not None:
                    rank += 1
                    add(candidate, rank, 1.2, {"via_concept": True})

    # Empowerment clauses, provisions for other exposures and disclosure/reporting articles rank lower.
    note_scope_text = note["text"]
    for key, candidate in found.items():
        candidate["flags"] = _flags(candidate, note_scope_text)
        scores[key] *= _demotion(candidate["flags"])

    ordered = sorted(found, key=lambda key: -scores[key])
    # Follow the cross-references of the strongest candidates: 'shall treat general credit risk adjustments in
    # accordance with Article 62(c)' makes Article 62(c) a candidate even when its own wording matches no query.
    source_by_id = {source["id"]: source for source in sources}
    for key in ordered[:6]:
        referring = found[key]
        source = source_by_id[referring["regulation_id"]]
        for reference in rt.references(referring["text"], referring["label"])[:8]:
            candidate = _referenced_candidate(source, reference)
            if candidate is None or candidate["unit"] == referring["unit"]:
                continue
            flags = found[candidate["key"]]["flags"] if candidate["key"] in found else _flags(candidate, note_scope_text)
            candidate["flags"] = flags
            add(candidate, 1, 0.35 * _demotion(flags), {"via_reference": referring["reference"]})
    ordered = sorted(found, key=lambda key: -scores[key])

    chosen: List[Dict[str, Any]] = []
    per_unit: Dict[Tuple[str, int], int] = {}

    def take(key: Tuple[Any, ...]) -> None:
        candidate = found[key]
        unit_key = (candidate["regulation_id"], candidate["unit"])
        if any(item["key"] == key for item in chosen) or per_unit.get(unit_key, 0) >= CANDIDATES_PER_UNIT:
            return
        per_unit[unit_key] = per_unit.get(unit_key, 0) + 1
        candidate = dict(candidate)
        candidate["id"] = f"C{len(chosen) + 1}"
        candidate["score"] = round(scores[key], 4)
        candidate["retrieval"] = {**signals[key], "rank": len(chosen) + 1}
        chosen.append(candidate)

    # Most slots go by fused score; some are reserved for the provisions that deal most with the note's concepts
    # (e.g. 'Exposures in default', whose risk weights depend on specific credit risk adjustments) and for the
    # provisions the strongest candidates refer to.
    for key in ordered:
        if len(chosen) >= CANDIDATE_LIMIT - REFERENCE_SLOTS - CONCEPT_SLOTS:
            break
        take(key)
    dense = sorted(
        (key for key in found if signals[key].get("concept_rank") and not found[key]["flags"]["mandate"] and not found[key]["flags"]["scope"]),
        key=lambda key: signals[key]["concept_rank"],
    )
    for key in dense[:CONCEPT_SLOTS]:
        take(key)
    for key in [key for key in ordered if signals[key].get("via_reference")][:REFERENCE_SLOTS]:
        take(key)
    for key in ordered:
        if len(chosen) >= CANDIDATE_LIMIT:
            break
        take(key)
    stats = {
        "queries": [query for query, _ in queries],
        "semantic": bool(vectors),
        "pool": len(found),
        "concepts": sorted(rt.concept_name(position) for position in about),
    }
    return chosen, stats


# --------------------------------------------------------------------------------------------------
# Verification
# --------------------------------------------------------------------------------------------------

_GAP = re.compile(r"\s*(?:\.\.\.|…|\[\.\.\.\]|\[…\])\s*")


def _fragments(quote: str) -> List[str]:
    return [part.strip() for part in _GAP.split(quote or "") if part.strip()]


def _exact_spans(quote: str, original: str) -> List[str]:
    """The exact source wording of each '…'-separated fragment of a quote (for highlighting)."""
    spans = []
    for fragment in _fragments(quote):
        if fragment in original:
            spans.append(fragment)
            continue
        ratio, window = rt.closest_original(fragment, original)
        if ratio >= 0.9 and window:
            spans.append(window)
    return spans


def _check_quote(quote: str, original: str) -> Tuple[str, str, float]:
    """('verified' | 'corrected' | 'failed', exact quote, similarity) for a quote against one source text."""
    quote = (quote or "").strip().strip('"“”„«»')
    if len(quote) < 8:
        return "failed", quote, 0.0
    if runtime._quote_in(quote, rt.normalise_for_match(original)):
        spans = _exact_spans(quote, original)
        return "verified", " … ".join(spans) if spans and not _GAP.search(quote) and len(spans) == 1 else quote, 1.0
    if _GAP.search(quote):
        return "failed", quote, 0.0
    ratio, window = rt.closest_original(quote, original)
    if ratio >= QUOTE_REPAIR_SIMILARITY and window:
        return "corrected", window, ratio
    return "failed", quote, ratio


def _point_of(candidate: Dict[str, Any], quote: str) -> Optional[str]:
    """The point (e.g. '95' or 'b') of a candidate that holds the quote, when the candidate lists points."""
    if not candidate["points"]:
        return None
    text = candidate["text"]
    found = None
    cursor = 0
    # 'Tier 2 items shall consist of the following: … (c) for institutions …' cites point (c): the last fragment
    # that lies inside a point decides, the lead-in before the points does not.
    for fragment in _fragments(quote):
        at = text.find(fragment, cursor)
        if at < 0:
            _, window = rt.closest_original(fragment, text)
            at = text.find(window, cursor) if window else -1
        if at < 0:
            continue
        cursor = at + len(fragment)
        best = None
        for point in candidate["points"]:
            # The window includes the fragment's own start, so a fragment beginning with its marker '(c) …' counts too.
            for match in re.finditer(rf"(?:^|\s)\({re.escape(point)}\)\s", text[: at + len(point) + 4]):
                if best is None or match.start() > best[0]:
                    best = (match.start(), point)
        if best:
            found = best[1]
    return found


def _verify(
    decision: Dict[str, Any], candidates: List[Dict[str, Any]], note: Dict[str, Any], sources: List[Dict[str, Any]], final: bool,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[str]]:
    """Accepted links, rejected links and problems (feedback for a correction round)."""
    by_id = {candidate["id"]: candidate for candidate in candidates}
    source_by_id = {source["id"]: source for source in sources}
    note_concepts = rt.concepts(note["full_text"])
    accepted: List[Dict[str, Any]] = []
    rejected: List[Dict[str, Any]] = []
    problems: List[Dict[str, Any]] = []
    seen: Set[str] = set()
    for raw in (decision.get("links") or [])[: MAX_LINKS + 2]:
        candidate = by_id.get(str(raw.get("candidate") or "").strip().strip("[]"))
        if candidate is None:
            problems.append(_problem("unknown_candidate", candidate=raw.get("candidate")))
            continue
        if candidate["id"] in seen:
            continue
        seen.add(candidate["id"])
        label = f"{candidate['id']} ({candidate['reference']})"
        checks: List[Dict[str, Any]] = []
        failed: List[Dict[str, Any]] = []

        status, quote, similarity = _check_quote(str(raw.get("regulation_quote") or ""), candidate["text"])
        target = candidate
        if status == "failed":
            # The quote may come from another paragraph of the same provision or another candidate.
            located = None
            source = source_by_id[candidate["regulation_id"]]
            result = rt.verify_quote(source["index"], str(raw.get("regulation_quote") or ""), [candidate["unit"]])
            if result.get("status") == "verified":
                located = f"{candidate['label']}({result['paragraphs'][0]})" if result.get("paragraphs") else candidate["label"]
            elif result.get("status") == "found_in_other_unit":
                located = source["index"]["units"][result["unit"]]["label"]
            if located and final:
                other = next((c for c in candidates if c["reference"] == located or (c["label"] == located)), None)
                if other is not None and other["id"] not in seen:
                    status, quote, similarity = _check_quote(str(raw.get("regulation_quote") or ""), other["text"])
                    if status != "failed":
                        target = other
                        seen.add(other["id"])
                        checks.append({"id": "citation_retargeted", "status": "corrected", "params": {"from": candidate["reference"], "to": other["reference"]}})
            if status == "failed":
                if located:
                    failed.append(_problem("quote_other_unit", label, located=located, reference=candidate["reference"]))
                else:
                    closest = rt.closest_original(str(raw.get("regulation_quote") or ""), candidate["text"])[1]
                    failed.append(_problem("quote_not_verbatim", label, closest=closest[:300]))
        checks.insert(0, {
            "id": "regulation_quote",
            "status": {"verified": "pass", "corrected": "corrected", "failed": "fail"}[status],
            "params": {"page": target["page_label"], "similarity": round(similarity, 3)},
        })
        regulation_quote = quote

        note_status, note_quote, note_similarity = _check_quote(str(raw.get("release_note_quote") or ""), note["full_text"])
        if note_status == "failed":
            for field in note["fields"]:
                field_status, field_quote, field_similarity = _check_quote(str(raw.get("release_note_quote") or ""), field["value"])
                if field_status != "failed":
                    note_status, note_quote, note_similarity = field_status, field_quote, field_similarity
                    break
        if note_status == "failed":
            failed.append(_problem("note_quote_not_verbatim", label))
        checks.append({
            "id": "note_quote",
            "status": {"verified": "pass", "corrected": "corrected", "failed": "fail"}[note_status],
            "params": {"similarity": round(note_similarity, 3)},
        })

        conflict = runtime._scope_conflict(note["text"], [{"title": target["title"]}])
        # Class ids travel with the labels so the page can show them in its own language.
        note_classes = sorted(rt.scope_classes(note["text"]))
        if conflict:
            provision_classes = sorted(rt.scope_classes(target["title"]))
            failed.append(_problem(
                "scope_conflict", label, note=conflict[1], provision=conflict[2], title=conflict[0],
                note_classes=note_classes, provision_classes=provision_classes,
            ))
            checks.append({"id": "scope", "status": "fail", "params": {
                "note": conflict[1], "provision": conflict[2], "title": conflict[0],
                "note_classes": note_classes, "provision_classes": provision_classes,
            }})
        else:
            note_scope = [runtime._SCOPE_LABELS.get(name, name) for name in note_classes]
            checks.append({"id": "scope", "status": "pass" if note_scope else "info", "params": {"note": ", ".join(note_scope), "note_classes": note_classes}})

        mandate = runtime._MANDATE_CLAUSE.search(regulation_quote) or (
            runtime._MANDATE_CLAUSE.search(target["text"]) and len(target["text"]) < 1500
        )
        if mandate:
            failed.append(_problem("mandate", label))
        checks.append({"id": "substantive", "status": "fail" if mandate else "pass", "params": {}})

        shared = sorted(rt.concept_name(position) for position in note_concepts & rt.concepts(target["text"]))
        if note_concepts:
            wanted = sorted(rt.concept_name(p) for p in note_concepts)
            if not shared:
                failed.append(_problem("concept_missing", label, concepts=wanted))
            checks.append({"id": "concept", "status": "pass" if shared else "fail", "params": {"concepts": shared, "note": wanted}})
        else:
            checks.append({"id": "concept", "status": "info", "params": {"concepts": shared, "note": []}})
        retrieval = target["retrieval"]
        checks.append({"id": "retrieval", "status": "info", "params": {
            "rank": retrieval["rank"], "keyword_rank": retrieval.get("keyword_rank"), "similarity": retrieval.get("similarity"),
            "queries": retrieval.get("queries", 0), "via_concept": retrieval.get("via_concept", False),
        }})

        reference = target["reference"]
        point = _point_of(target, regulation_quote) if status != "failed" else None
        if point and target["kind"] == "article" and not reference.endswith(f"({point})"):
            # Cite the point that holds the quote: 'Article 62' -> 'Article 62(c)', 'Article 4(1)(94)–(96)' -> 'Article 4(1)(95)'.
            reference = re.sub(r"\(\w+\)–\(\w+\)$", "", reference)
            reference = reference if reference.endswith(f"({point})") else f"{reference}({point})"

        model_confidence = max(0, min(100, int(raw.get("confidence") or 0)))
        evidence = (
            0.4 * (1.0 if status == "verified" else 0.85 if status == "corrected" else 0.0)
            + 0.15 * (1.0 if note_status == "verified" else 0.85 if note_status == "corrected" else 0.0)
            + 0.25 * (1.0 if shared else 0.55 if not note_concepts else 0.25)
            + 0.2 * max(0.2, 1.0 - (retrieval["rank"] - 1) * 0.08)
        )
        confidence = round(0.6 * model_confidence + 40 * evidence)
        link = {
            "candidate": target["id"],
            "regulation_id": target["regulation_id"],
            "regulation": target["regulation"],
            "regulation_file": target["regulation_file"],
            "unit": target["unit"],
            "label": target["label"],
            "reference": reference,
            "title": target["title"],
            "path": target["path"],
            "pages": target["pages"],
            "page": target["page"],
            "page_label": target["page_label"],
            "paragraph": target["paragraph"],
            "point": point,
            "link_type": "indirect" if raw.get("link_type") == "indirect" else "direct",
            "approach": raw.get("approach") or "not_specific",
            "model_confidence": model_confidence,
            "confidence": confidence,
            "band": "high" if confidence >= 80 else "medium" if confidence >= 60 else "low",
            "affected_element": str(raw.get("affected_element") or "").strip(),
            "rationale": str(raw.get("rationale") or "").strip(),
            "regulation_quote": regulation_quote,
            "regulation_highlights": _exact_spans(regulation_quote, target["text"]) if status != "failed" else [],
            "note_quote": note_quote,
            "note_highlights": _exact_spans(note_quote, note["full_text"]) if note_status != "failed" else [],
            "provision_text": target["text"],
            "formula": target["formula"],
            "concepts": shared,
            "checks": checks,
            "references": rt.references(target["text"], target["label"])[:12],
        }
        if failed:
            problems.extend(failed)
            if final:
                rejected.append({
                    "candidate": target["id"], "reference": target["reference"], "title": target["title"],
                    "regulation": target["regulation"], "page_label": target["page_label"], "source": "verification",
                    "reason": " ".join(item["bodies"]["en"] for item in failed),
                    "reason_i18n": {language: " ".join(item["bodies"][language] for item in failed) for language in TEXT_LANGUAGES},
                })
            continue
        accepted.append(link)
    accepted.sort(key=lambda link: (link["link_type"] != "direct", -link["confidence"]))
    for raw in decision.get("rejected") or []:
        candidate = by_id.get(str(raw.get("candidate") or "").strip().strip("[]"))
        if candidate is None or candidate["id"] in seen or any(item["candidate"] == candidate["id"] for item in rejected):
            continue
        rejected.append({
            "candidate": candidate["id"], "reference": candidate["reference"], "title": candidate["title"],
            "regulation": candidate["regulation"], "page_label": candidate["page_label"], "source": "model",
            "reason": str(raw.get("reason") or "").strip(),
        })
    return accepted, rejected, problems


def _reference_gaps(
    accepted: List[Dict[str, Any]], decision: Dict[str, Any], candidates: List[Dict[str, Any]], note: Dict[str, Any],
    sources: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """Reference-chain completeness: a quoted rule that applies 'in accordance with Article 62(c)' must have that
    provision linked or explicitly dismissed when it deals with the note's concept."""
    note_concepts = rt.concepts(note["full_text"])
    if not note_concepts:
        return []
    linked = {link["candidate"] for link in accepted}
    dismissed = {str(item.get("candidate") or "").strip("[] ") for item in decision.get("rejected") or []}
    source_by_id = {source["id"]: source for source in sources}
    gaps = []
    for link in accepted:
        source = source_by_id[link["regulation_id"]]
        for reference in rt.references(link["regulation_quote"], link["label"]):
            target = _referenced_candidate(source, reference)
            if target is None:
                continue
            match = next((c for c in candidates if c["key"] == target["key"]), None) or next(
                (c for c in candidates if c["regulation_id"] == target["regulation_id"] and c["unit"] == target["unit"]), None,
            )
            if match is None or match["id"] in linked or match["id"] in dismissed or not note_concepts & rt.concepts(match["text"]):
                continue
            gaps.append(_problem(
                "reference_gap", source=link["candidate"], source_reference=link["reference"], reference=reference,
                target=match["id"], target_reference=match["reference"],
            ))
    return list({gap["text"]: gap for gap in gaps}.values())


def _no_link_challenge(note: Dict[str, Any], candidates: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """A reminder when the model found no link although candidates deal with the note's own concepts."""
    about = rt.concepts(note["full_text"])
    if not about:
        return None
    fitting = [
        candidate for candidate in candidates
        if about & rt.concepts(candidate["text"]) and not candidate["flags"]["mandate"] and not candidate["flags"]["scope"]
    ][:4]
    if not fitting:
        return None
    names = sorted(rt.concept_name(position) for position in about)
    listed = "; ".join(f"[{c['id']}] {c['reference']} ({c['title']})" for c in fitting)
    return _problem("no_link_challenge", concepts=names, listed=listed)


# --------------------------------------------------------------------------------------------------
# Translation of the AI-written texts
# --------------------------------------------------------------------------------------------------

TRANSLATE_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["items"],
    "properties": {
        "items": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["id", "text"],
                "properties": {"id": {"type": "string"}, "text": {"type": "string"}},
            },
        },
    },
}

_TRANSLATE_SYSTEM = (
    "You translate the texts of a regulatory traceability analysis for display in another language. A bank's release "
    "notes (changes to its regulatory calculation engine) were linked to provisions of EU banking regulation (CRR, CRD and "
    "others); the texts are the analyst's interpretation of each change, the rationales of the links, the reviewer's "
    "reasons and short labels.\n\n"
    "Translate the text of every item into the language named in its 'to' field and return every item with its id.\n"
    "Rules:\n"
    "- Use the terminology of the official English and German language versions of EU banking regulation, e.g. exposure "
    "value = Risikopositionswert, risk weight = Risikogewicht, risk-weighted exposure amount = risikogewichteter "
    "Positionsbetrag, specific / general credit risk adjustments = spezifische / allgemeine Kreditrisikoanpassungen, "
    "carrying amount / accounting value = Buchwert, accrued interest = aufgelaufene Zinsen, own funds = Eigenmittel, credit "
    "risk mitigation = Kreditrisikominderung, retail exposures = Risikopositionen aus dem Mengengeschäft, standardised "
    "approach = Standardansatz, off-balance-sheet items = außerbilanzielle Posten, conversion factor = Umrechnungsfaktor.\n"
    "- Keep the numbering of provision references exactly as written (e.g. 'Article 111(1)', 'Article 4(1)(95)', 'Annex I'; "
    "in German the words become 'Artikel' and 'Anhang'), and keep all numbers, "
    "Jira IDs, abbreviations (CRR, CRD, IRB, EAD, RWA, CCF, EWB, PWB, IReF) and any text in quotation marks – those are "
    "verbatim quotes from the release note or the regulation.\n"
    "- Keep the meaning, the level of detail and the concise, professional tone; do not add, drop or soften anything.\n"
    "- A short label stays a short label, a single term is translated as the term.\n"
    "- When an item has a 'note', an earlier translation of it was rejected for that reason – fix it."
)
_TRANSLATION_TOKEN = re.compile(r"\b\d+[a-z]?\b(?:\(\w{1,5}\))*")


def _translation_items(result: Dict[str, Any], language: str) -> List[Dict[str, str]]:
    """The AI-written texts of a finished note, each with its target language: the texts written in the run language go
    into the other language; the regulation terms and affected items, which the model always writes in English, into
    German."""
    others = [code for code in TEXT_LANGUAGES if code != language]
    german = [code for code in TEXT_LANGUAGES if code != "en"]
    items: List[Dict[str, str]] = []

    def add(field: str, text: Any, targets: List[str]) -> None:
        text = str(text or "").strip()
        if text:
            items.extend({"id": f"{target}|{field}", "to": target, "text": text} for target in targets)

    interpretation = result.get("interpretation") or {}
    add("summary", interpretation.get("summary"), others)
    for index, concept in enumerate(interpretation.get("concepts") or []):
        add(f"effect.{index}", concept.get("effect"), others)
        add(f"term.{index}", concept.get("term"), german)
    for index, item in enumerate(interpretation.get("affected_items") or []):
        add(f"affected.{index}", item, german)
    if result.get("assessment") != interpretation.get("summary"):
        add("assessment", result.get("assessment"), others)
    if not result.get("no_link_code"):
        add("no_link_reason", result.get("no_link_reason"), others)
    for index, link in enumerate(result.get("links") or []):
        add(f"rationale.{index}", link.get("rationale"), others)
        add(f"element.{index}", link.get("affected_element"), others)
        review = next((check for check in link.get("checks") or [] if check.get("id") == "review"), None)
        if review:
            add(f"review.{index}", (review.get("params") or {}).get("reason"), others)
    for index, item in enumerate(result.get("rejected") or []):
        if item.get("source") != "verification":  # verification reasons come from the templates in both languages
            add(f"rejected.{index}", item.get("reason"), others)
    return items


def _translation_problem(source: str, text: str) -> Optional[str]:
    if not text:
        return "The translation is missing."
    missing = set(_TRANSLATION_TOKEN.findall(source)) - set(_TRANSLATION_TOKEN.findall(text))
    if missing:
        return f"Keep these numbers and references exactly as in the source: {', '.join(sorted(missing))}."
    return None


_TRANSLATION_FIELDS = {"rationale": "rationale", "element": "affected_element", "review": "review_reason"}


def _translate_note(llm: "_LLM", result: Dict[str, Any], language: str) -> Tuple[Dict[str, Dict[str, Any]], int, int]:
    """The note's AI-written texts in the other language(s): ({language: fields}, texts translated, texts kept in the
    original because their translation failed the checks twice)."""
    items = _translation_items(result, language)
    if not items:
        return {}, 0, 0
    translated: Dict[str, str] = {}
    pending = items
    notes: Dict[str, str] = {}
    for _ in range(2):
        payload = []
        for item in pending:
            entry = {"id": item["id"], "to": LANGUAGE_NAMES[item["to"]], "text": item["text"]}
            if item["id"] in notes:
                entry["note"] = notes[item["id"]]
            payload.append(entry)
        answer = llm.json(
            [{"role": "system", "content": _TRANSLATE_SYSTEM}, {"role": "user", "content": json.dumps({"items": payload}, ensure_ascii=False)}],
            "translation", TRANSLATE_SCHEMA,
        )
        returned = {str(entry.get("id") or ""): str(entry.get("text") or "").strip() for entry in answer.get("items") or []}
        notes = {}
        for item in pending:
            text = returned.get(item["id"], "")
            problem = _translation_problem(item["text"], text)
            if problem:
                notes[item["id"]] = problem
            else:
                translated[item["id"]] = text
        pending = [item for item in pending if item["id"] in notes]
        if not pending:
            break
    fields: Dict[str, Dict[str, Any]] = {}
    for item in items:
        text = translated.get(item["id"])
        if text is None:
            continue
        target, field = item["id"].split("|", 1)
        bucket = fields.setdefault(target, {})
        kind, _, key = field.partition(".")
        if kind in ("summary", "assessment", "no_link_reason"):
            bucket[kind] = text
        elif kind == "effect":
            bucket.setdefault("effects", {})[key] = text
        elif kind == "term":
            bucket.setdefault("terms", {})[key] = text
        elif kind == "affected":
            bucket.setdefault("affected_items", {})[key] = text
        elif kind in _TRANSLATION_FIELDS:
            bucket.setdefault("links", {}).setdefault(key, {})[_TRANSLATION_FIELDS[kind]] = text
        elif kind == "rejected":
            bucket.setdefault("rejected", {})[key] = text
    return fields, len(translated), len(items) - len(translated)


def _localized_note(note: Dict[str, Any], language: str, run_language: str) -> Dict[str, Any]:
    """A copy of a note with its texts in `language` where a translation exists (the export's view of the note)."""
    note = copy.deepcopy(note)
    fields = (note.get("translations") or {}).get(language) or {}
    interpretation = note.get("interpretation")
    original_summary = (interpretation or {}).get("summary")
    if interpretation:
        if fields.get("summary"):
            interpretation["summary"] = fields["summary"]
        for index, concept in enumerate(interpretation.get("concepts") or []):
            concept["effect"] = (fields.get("effects") or {}).get(str(index), concept.get("effect"))
            concept["term"] = (fields.get("terms") or {}).get(str(index), concept.get("term"))
        interpretation["affected_items"] = [
            (fields.get("affected_items") or {}).get(str(index), item) for index, item in enumerate(interpretation.get("affected_items") or [])
        ]
    if note.get("assessment") and note.get("assessment") == original_summary and interpretation:
        note["assessment"] = interpretation["summary"]
    elif fields.get("assessment"):
        note["assessment"] = fields["assessment"]
    if fields.get("no_link_reason"):
        note["no_link_reason"] = fields["no_link_reason"]
    for index, link in enumerate(note.get("links") or []):
        texts = (fields.get("links") or {}).get(str(index)) or {}
        link["rationale"] = texts.get("rationale", link.get("rationale"))
        link["affected_element"] = texts.get("affected_element", link.get("affected_element"))
        for check in link.get("checks") or []:
            if check.get("id") == "review" and texts.get("review_reason"):
                check["params"] = {**(check.get("params") or {}), "reason": texts["review_reason"]}
    for index, item in enumerate(note.get("rejected") or []):
        item["reason"] = ((item.get("reason_i18n") or {}).get(language)
                          or (fields.get("rejected") or {}).get(str(index))
                          or item.get("reason"))
    if note.get("corrections_i18n", {}).get(language):
        note["corrections"] = note["corrections_i18n"][language]
    return note


# --------------------------------------------------------------------------------------------------
# Jobs
# --------------------------------------------------------------------------------------------------

def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


class _Job:
    def __init__(self, state: Dict[str, Any]):
        self.state = state
        self.lock = threading.Lock()
        self.cancel = threading.Event()
        self.started = time.monotonic()
        self.last_save = 0.0

    def log(self, level: str, message: str, note: Optional[str] = None) -> None:
        with self.lock:
            self.state["log"].append({
                "t": round(time.monotonic() - self.started, 2), "level": level, "message": message, "note": note,
            })

    def say(self, level: str, code: str, note: Optional[str] = None, **params: Any) -> None:
        """A coded log line in English (`message`) and German (`i18n`)."""
        english, german = _log_texts(code, params)
        with self.lock:
            self.state["log"].append({
                "t": round(time.monotonic() - self.started, 2), "level": level, "message": english, "note": note,
                "code": code, "i18n": {"de": german},
            })

    def count_translation(self, texts: int, kept: int) -> None:
        with self.lock:
            stats = self.state.setdefault("translation", {"notes": 0, "texts": 0, "kept": 0})
            stats["notes"] += 1 if texts or kept else 0
            stats["texts"] += texts
            stats["kept"] += kept

    def set_model(self, model: str) -> None:
        with self.lock:
            self.state["model"] = model

    def count_call(self, prompt: int, completion: int) -> None:
        with self.lock:
            usage = self.state["usage"]
            usage["llm_calls"] += 1
            usage["prompt_tokens"] += prompt
            usage["completion_tokens"] += completion

    def note(self, key: str, **changes: Any) -> None:
        with self.lock:
            for note in self.state["notes"]:
                if note["key"] == key:
                    note.update(changes)
                    break

    def stage(self, stage: str) -> None:
        with self.lock:
            self.state["stage"] = stage

    def snapshot(self, lite: bool = False) -> Dict[str, Any]:
        with self.lock:
            if not lite:
                return copy.deepcopy(self.state)
            state = {key: value for key, value in self.state.items() if key not in ("notes", "provisions")}
            state = copy.deepcopy(state)
            state["notes"] = [
                {key: note.get(key) for key in ("key", "jira_id", "file", "sheet", "phase", "status", "link_count", "error")}
                for note in self.state["notes"]
            ]
            return state

    def save(self, force: bool = False) -> None:
        now = time.monotonic()
        if not force and now - self.last_save < 2.0:
            return
        self.last_save = now
        _write_run(self.snapshot())


_JOBS: Dict[str, _Job] = {}
_JOBS_LOCK = threading.Lock()


def _run_path(run_id: str) -> Path:
    if not _ID.match(run_id or ""):
        raise HTTPException(404, "Run not found")
    return RUNS_DIR / f"{run_id}.json"


def _write_run(state: Dict[str, Any]) -> None:
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    path = _run_path(state["id"])
    temporary = path.with_name(f"{path.stem}.{uuid.uuid4().hex[:8]}.tmp")
    temporary.write_text(json.dumps(state, ensure_ascii=False, default=str), encoding="utf-8")
    for attempt in range(20):
        try:
            os.replace(temporary, path)
            return
        except PermissionError:
            if attempt == 19:
                temporary.unlink(missing_ok=True)
                raise
            time.sleep(0.05)


def _read_run(run_id: str) -> Dict[str, Any]:
    path = _run_path(run_id)
    if not path.exists():
        raise HTTPException(404, "Run not found")
    for attempt in range(10):
        try:
            state = json.loads(path.read_text(encoding="utf-8"))
            break
        except PermissionError:
            if attempt == 9:
                raise
            time.sleep(0.05)
    if state.get("status") == "running" and run_id not in _JOBS:
        state["status"] = "interrupted"  # the server restarted while it ran
    return state


def _process_note(job: _Job, llm: _LLM, note: Dict[str, Any], sources: List[Dict[str, Any]], language: str) -> None:
    key = note["key"]
    label = note["jira_id"] or f"{note['sheet']} #{note['ordinal']}"
    started = time.monotonic()
    try:
        job.note(key, phase="interpreting")
        interpretation = llm.json(_interpret_messages(note, language), "interpretation", INTERPRET_SCHEMA)
        concepts = [item for item in interpretation.get("concepts", []) if isinstance(item, dict)][:10]
        normalised_note = rt.normalise_for_match(note["full_text"])
        for item in concepts:
            phrase = str(item.get("source_phrase") or "").strip()
            item["verbatim"] = bool(phrase) and runtime._quote_in(phrase, normalised_note)
        interpretation["concepts"] = concepts
        job.say("info", "interpreted", label, terms=[str(item.get("term")) for item in concepts])

        job.note(key, phase="retrieving", interpretation=interpretation)
        candidates, stats = _retrieve(note, interpretation, sources)
        job.say("info", "retrieved", label, candidates=len(candidates), pool=stats["pool"], queries=len(stats["queries"]), semantic=bool(stats["semantic"]))
        if not candidates:
            _finish_note(job, llm, key, label, interpretation, language, started, {
                "status": "no_link", "link_count": 0, "links": [], "rejected": [], "candidates": [], "retrieval": stats,
                "assessment": "", "no_link_reason": _NO_CANDIDATES.get(language, _NO_CANDIDATES["en"]), "no_link_code": "no_candidates",
                "corrections": [],
            })
            return

        job.note(key, phase="adjudicating")
        messages = _decision_messages(note, interpretation, candidates, language)
        decision = llm.json(messages, "link_decision", DECISION_SCHEMA)
        job.note(key, phase="verifying")
        accepted, rejected, problems = _verify(decision, candidates, note, sources, final=False)
        problems += _reference_gaps(accepted, decision, candidates, note, sources)
        challenge = _no_link_challenge(note, candidates) if not decision.get("links") else None
        feedback: List[Dict[str, Any]] = []
        if problems or challenge:
            feedback = problems + ([challenge] if challenge else [])
            job.say("warn", "correction_round", label, texts={language: [item["texts"][language] for item in feedback] for language in TEXT_LANGUAGES})
            messages = messages + [
                {"role": "assistant", "content": json.dumps(decision, ensure_ascii=False)},
                {"role": "user", "content": (
                    "Automatic verification of your answer found these problems:\n- " + "\n- ".join(item["text"] for item in feedback) +
                    "\nFix or drop each affected link, keep the correct ones unchanged, and return the complete JSON again."
                )},
            ]
            job.note(key, phase="adjudicating")
            decision = llm.json(messages, "link_decision", DECISION_SCHEMA)
            job.note(key, phase="verifying")
        accepted, rejected, problems = _verify(decision, candidates, note, sources, final=True)
        failed_links = sum(1 for item in rejected if item["source"] == "verification")
        if failed_links:
            job.say("warn", "verification_rejected", label, count=failed_links)
        escalated = False
        if accepted:
            job.note(key, phase="reviewing")
            accepted, dismissed, escalated = _review(llm, note, interpretation, accepted, language)
            rejected = dismissed + rejected
            changes = [item["reference"] for item in dismissed]
            downgraded = [link["reference"] for link in accepted if any(check["id"] == "review" and check["status"] == "corrected" for check in link["checks"])]
            job.say(
                "warn" if (changes or downgraded or escalated) else "info", "review", label,
                escalated=escalated, rejected=changes, downgraded=downgraded, confirmed=len(accepted),
            )
        linked_ids = {link["candidate"] for link in accepted}
        rejected_ids = {item["candidate"] for item in rejected}
        candidate_rows = [
            {
                "id": candidate["id"], "reference": candidate["reference"], "title": candidate["title"],
                "regulation": candidate["regulation"], "regulation_id": candidate["regulation_id"], "unit": candidate["unit"],
                "page_label": candidate["page_label"], "page": candidate["page"], "path": candidate["path"],
                "score": candidate["score"], "retrieval": candidate["retrieval"], "flags": candidate["flags"],
                "decision": "linked" if candidate["id"] in linked_ids else "rejected" if candidate["id"] in rejected_ids else "not_selected",
                "excerpt": candidate["text"][:420],
            }
            for candidate in candidates
        ]
        if not accepted:
            status = "no_link"
        elif escalated or max(link["confidence"] for link in accepted) < 60:
            status = "review"
        else:
            status = "linked"
        job.say("success" if accepted else "info", "result", label, links=[
            {"reference": link["reference"], "link_type": link["link_type"]} for link in accepted
        ])
        _finish_note(job, llm, key, label, interpretation, language, started, {
            "status": status,
            "link_count": len(accepted),
            "links": accepted,
            "rejected": rejected,
            "candidates": candidate_rows,
            "retrieval": stats,
            "assessment": str(decision.get("assessment") or "").strip(),
            "no_link_reason": str(decision.get("no_link_reason") or "").strip() if not accepted else "",
            "corrections": [item["text"] for item in feedback],
            "corrections_i18n": {language: [item["texts"][language] for item in feedback] for language in TEXT_LANGUAGES} if feedback else {},
        })
    except _Cancelled:
        job.note(key, phase="cancelled")
    except Exception as exc:
        traceback.print_exc()
        message = runtime._friendly_llm_error(exc, "en") if getattr(exc, "status_code", None) else str(exc)
        job.note(key, phase="failed", status="failed", link_count=0, error=message[:400])
        job.say("error", "note_failed", label, error=message[:300])
    finally:
        job.save()


def _finish_note(
    job: _Job, llm: _LLM, key: str, label: str, interpretation: Dict[str, Any], language: str, started: float, result: Dict[str, Any],
) -> None:
    """Translate the note's AI-written texts into the other language, then publish the result. A failed or cancelled
    translation leaves the texts in their original language; the result itself is never lost."""
    job.note(key, phase="translating")
    translations: Dict[str, Dict[str, Any]] = {}
    texts = kept = 0
    try:
        translations, texts, kept = _translate_note(llm, {**result, "interpretation": interpretation}, language)
    except _Cancelled:
        pass
    except Exception as exc:
        traceback.print_exc()
        job.say("warn", "translation_failed", label, error=str(exc)[:200])
    if result.get("no_link_code") == "no_candidates":
        for other in TEXT_LANGUAGES:
            if other != language:
                translations.setdefault(other, {})["no_link_reason"] = _NO_CANDIDATES[other]
    job.count_translation(texts, kept)
    job.note(key, phase="done", translations=translations, duration_ms=int((time.monotonic() - started) * 1000), **result)


def _provision_key(link: Dict[str, Any]) -> str:
    return f"{link['regulation_id']}:{link['reference']}"


def _sort_key(item: Dict[str, Any]) -> Tuple[Any, ...]:
    number = str(item.get("number") or "")
    match = re.match(r"(\d+)([a-z]*)", number)
    article = (int(match.group(1)), match.group(2)) if match else (10**6, number)
    paragraph = re.findall(r"\((\w+)\)", item["reference"].split(" ", 1)[-1])
    return (item["regulation"], item["kind"] != "article", article, [p.zfill(4) for p in paragraph])


def _assemble(job: _Job) -> None:
    with job.lock:
        notes = job.state["notes"]
        provisions: Dict[str, Dict[str, Any]] = {}
        for note in notes:
            for link in note.get("links") or []:
                key = _provision_key(link)
                entry = provisions.setdefault(key, {
                    "key": key, "regulation": link["regulation"], "regulation_id": link["regulation_id"],
                    "reference": link["reference"], "label": link["label"], "title": link["title"], "path": link["path"],
                    "unit": link["unit"], "page_label": link["page_label"], "page": link["page"],
                    "number": None, "kind": "article" if link["label"].startswith(("Article", "Artikel")) else "annex",
                    "notes": [], "direct": 0, "indirect": 0,
                })
                entry["number"] = re.sub(r"^\D+", "", link["label"]) or None
                entry["notes"].append({"key": note["key"], "jira_id": note["jira_id"], "link_type": link["link_type"], "confidence": link["confidence"]})
                entry[link["link_type"]] += 1
                link["provision_key"] = key
        ordered = sorted(provisions.values(), key=_sort_key)
        links = [link for note in notes for link in note.get("links") or []]
        quote_checks = [check for link in links for check in link["checks"] if check["id"] in ("regulation_quote", "note_quote")]
        done = [note for note in notes if note.get("phase") == "done"]
        summary = {
            "notes": len(notes),
            "analysed": len(done),
            "linked": sum(1 for note in notes if note.get("status") == "linked"),
            "review": sum(1 for note in notes if note.get("status") == "review"),
            "no_link": sum(1 for note in notes if note.get("status") == "no_link"),
            "failed": sum(1 for note in notes if note.get("status") == "failed"),
            "links": len(links),
            "direct": sum(1 for link in links if link["link_type"] == "direct"),
            "indirect": sum(1 for link in links if link["link_type"] == "indirect"),
            "provisions": len(ordered),
            "articles": len({(item["regulation_id"], item["label"]) for item in ordered}),
            "quotes_total": len(quote_checks),
            "quotes_verified": sum(1 for check in quote_checks if check["status"] in ("pass", "corrected")),
            "quotes_corrected": sum(1 for check in quote_checks if check["status"] == "corrected"),
            "average_confidence": round(sum(link["confidence"] for link in links) / len(links)) if links else None,
            "candidates_reviewed": sum(len(note.get("candidates") or []) for note in notes),
            "rejected": sum(len(note.get("rejected") or []) for note in notes),
            "corrections": sum(1 for note in notes if note.get("corrections")),
        }
        summary["coverage"] = round(100 * (summary["linked"] + summary["review"]) / summary["notes"]) if summary["notes"] else 0
        job.state["provisions"] = ordered
        job.state["summary"] = summary


def _run_job(job: _Job, notes: List[Dict[str, Any]], regulation_ids: List[str], language: str) -> None:
    # Let the start request's response go out first: parsing a large index holds the interpreter lock for seconds.
    time.sleep(0.3)
    try:
        job.stage("load")
        job.say("info", "load_indexes")
        sources = _regulation_sources(regulation_ids)
        semantic = [source["short"] for source in sources if source["embeddings"] is not None]
        job.say("info", "indexes_loaded", sources=[
            {"short": source["short"], "articles": source["index"]["quality"]["articles"], "passages": source["index"]["quality"]["passages"]}
            for source in sources
        ], semantic=semantic)
        llm = _LLM(job)
        job.set_model(llm.model)
        job.say("info", "model", model=llm.model, workers=WORKERS)
        job.stage("analyse")
        with ThreadPoolExecutor(max_workers=WORKERS) as pool:
            futures = [pool.submit(_process_note, job, llm, note, sources, language) for note in notes]
            for future in futures:
                future.result()
        job.stage("assemble")
        _assemble(job)
        summary = job.state["summary"]
        translation = job.state.get("translation") or {}
        if translation.get("texts"):
            job.say("info", "translated", notes=translation["notes"], texts=translation["texts"], kept=translation["kept"])
        with job.lock:
            job.state["status"] = "cancelled" if job.cancel.is_set() else "completed"
        job.say("success" if not job.cancel.is_set() else "warn", "run_cancelled" if job.cancel.is_set() else "run_completed", summary=summary)
    except HTTPException as exc:
        with job.lock:
            job.state.update(status="failed", error=str(exc.detail))
        job.say("error", "run_error", error=str(exc.detail))
    except Exception as exc:
        traceback.print_exc()
        with job.lock:
            job.state.update(status="failed", error=str(exc)[:500])
        job.say("error", "run_failed", error=str(exc)[:300])
    finally:
        with job.lock:
            job.state["finished_at"] = _now()
            job.state["duration_ms"] = int((time.monotonic() - job.started) * 1000)
            job.state["stage"] = "finished"
        job.save(force=True)
        with _JOBS_LOCK:
            _JOBS.pop(job.state["id"], None)


# --------------------------------------------------------------------------------------------------
# API
# --------------------------------------------------------------------------------------------------

class RunRequest(BaseModel):
    release_note_file_ids: List[str]
    regulation_ids: List[str]
    language: str = "en"


@router.get("/sources")
def list_sources():
    files = []
    for context, notes in _release_note_files():
        files.append({
            "id": str(context.get("id")),
            "filename": context.get("filename"),
            "kind": context.get("kind") or "excel",
            "notes": len(notes),
            "jira_ids": [note["jira_id"] for note in notes if note["jira_id"]][:12],
        })
    workbooks = {str(item.get("id")): item for item in runtime._platform("list_release_note_workbooks")()}
    for entry in files:
        entry["upload_date"] = workbooks.get(entry["id"], {}).get("upload_date")
    regulations = []
    for document in runtime._platform("list_regulation_documents")():
        regulations.append({
            "id": str(document.get("id")),
            "filename": document.get("filename"),
            "short": _short_name(document),
            "title": document.get("document_title"),
            "status": document.get("index_status"),
            "articles": document.get("article_count"),
            "annexes": document.get("annex_count"),
            "passages": document.get("passage_count"),
            "pages": document.get("page_count"),
            "semantic": bool(document.get("semantic")),
            "language": document.get("language"),
            "upload_date": document.get("upload_date"),
        })
    model = os.environ.get("REGULATION_MATCHER_MODEL", "").strip() or (
        DEFAULT_OPENAI_MODEL if os.environ.get("OPENAI_API_KEY") else os.environ.get("AZURE_OPENAI_DEPLOYMENT_NAME", "")
    )
    return {"release_note_files": files, "regulations": regulations, "model": model}


@router.post("/runs")
def start_run(request: RunRequest):
    file_ids = list(dict.fromkeys(str(item) for item in request.release_note_file_ids if str(item).strip()))
    regulation_ids = list(dict.fromkeys(str(item) for item in request.regulation_ids if str(item).strip()))
    if not file_ids:
        raise HTTPException(400, "Select at least one release-note file.")
    if not regulation_ids:
        raise HTTPException(400, "Select at least one regulation.")
    with _JOBS_LOCK:
        if _JOBS:
            raise HTTPException(409, "A matching run is already in progress. Wait for it to finish or cancel it.")
    _ready_regulations(regulation_ids)  # fail fast when a regulation is not ready; the job loads the indexes
    # Fail fast when no LLM is configured (checked from the environment; creating the client imports openai).
    if not os.environ.get("OPENAI_API_KEY") and not (os.environ.get("AZURE_OPENAI_API_KEY") and os.environ.get("AZURE_OPENAI_ENDPOINT")):
        raise HTTPException(503, "OpenAI is not configured. Set OPENAI_API_KEY, or AZURE_OPENAI_API_KEY and AZURE_OPENAI_ENDPOINT.")
    files = _release_note_files(file_ids)
    if not files:
        raise HTTPException(404, "The selected release-note files no longer exist.")
    notes = [note for _, file_notes in files for note in file_notes]
    if not notes:
        raise HTTPException(400, "The selected files contain no release notes with text.")
    if len(notes) > MAX_NOTES:
        raise HTTPException(400, f"The selected files hold {len(notes)} release notes; a run takes at most {MAX_NOTES}. Select fewer files.")
    language = request.language if request.language in LANGUAGE_NAMES else "en"
    documents = {str(item.get("id")): item for item in runtime._platform("list_regulation_documents")()}
    run_id = str(uuid.uuid4())
    state = {
        "id": run_id,
        "status": "running",
        "stage": "queued",
        "created_at": _now(),
        "finished_at": None,
        "duration_ms": None,
        "language": language,
        "model": None,
        "usage": {"llm_calls": 0, "prompt_tokens": 0, "completion_tokens": 0},
        "release_note_files": [
            {"id": str(context.get("id")), "filename": context.get("filename"), "kind": context.get("kind") or "excel", "notes": len(file_notes)}
            for context, file_notes in files
        ],
        "regulations": [
            {
                "id": regulation_id, "filename": documents[regulation_id].get("filename"), "short": _short_name(documents[regulation_id]),
                "title": documents[regulation_id].get("document_title"), "articles": documents[regulation_id].get("article_count"),
                "pages": documents[regulation_id].get("page_count"),
            }
            for regulation_id in regulation_ids
        ],
        "settings": {"candidates_per_note": CANDIDATE_LIMIT, "max_links": MAX_LINKS, "workers": WORKERS, "quote_repair_similarity": QUOTE_REPAIR_SIMILARITY},
        "notes": [{**note, "phase": "queued", "status": None, "link_count": 0} for note in notes],
        "provisions": [],
        "summary": None,
        "log": [],
        "error": None,
        # The page shows the names of the glossary concepts in the reader's language.
        "concept_names": {"de": CONCEPT_NAMES_DE},
    }
    job = _Job(state)
    job.say("info", "run_started", notes=len(notes), files=len(files), regulations=len(regulation_ids))
    with _JOBS_LOCK:
        _JOBS[run_id] = job
    job.save(force=True)
    threading.Thread(target=_run_job, args=(job, notes, regulation_ids, language), name=f"matcher-{run_id[:8]}", daemon=True).start()
    return job.snapshot(lite=True)


@router.get("/runs")
def list_runs():
    runs = []
    if RUNS_DIR.exists():
        for path in sorted(RUNS_DIR.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True)[:40]:
            try:
                state = _read_run(path.stem)
            except (HTTPException, OSError, json.JSONDecodeError):
                continue
            runs.append({key: state.get(key) for key in (
                "id", "status", "created_at", "finished_at", "duration_ms", "model", "language", "release_note_files", "regulations", "summary",
            )})
    runs.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"runs": runs}


@router.get("/runs/{run_id}")
def get_run(run_id: str, lite: bool = False):
    job = _JOBS.get(run_id)
    if job is not None:
        return job.snapshot(lite=lite)
    state = _read_run(run_id)
    if lite:
        state["notes"] = [
            {key: note.get(key) for key in ("key", "jira_id", "file", "sheet", "phase", "status", "link_count", "error")}
            for note in state.get("notes", [])
        ]
        state.pop("provisions", None)
    return state


@router.post("/runs/{run_id}/cancel")
def cancel_run(run_id: str):
    job = _JOBS.get(run_id)
    if job is None:
        raise HTTPException(404, "This run is not in progress.")
    job.cancel.set()
    job.say("warn", "cancel_requested")
    return {"id": run_id, "status": "cancelling"}


@router.delete("/runs/{run_id}")
def delete_run(run_id: str):
    if run_id in _JOBS:
        raise HTTPException(409, "Cancel the run before deleting it.")
    path = _run_path(run_id)
    if not path.exists():
        raise HTTPException(404, "Run not found")
    path.unlink()
    return {"id": run_id, "deleted": True}


@router.get("/provision")
def get_provision(regulation_id: str, unit: int):
    documents = {str(item.get("id")): item for item in runtime._platform("list_regulation_documents")()}
    document = documents.get(regulation_id)
    if not document:
        raise HTTPException(404, "Regulation not found")
    try:
        index = runtime._regulation_index_for(document)
    except runtime.ToolError as exc:
        raise HTTPException(409, str(exc)) from exc
    if not 0 <= unit < len(index["units"]):
        raise HTTPException(404, "Provision not found")
    record = index["units"][unit]
    passages = [
        {
            "paragraph": passage.get("paragraph"), "points": passage.get("points") or [], "page": passage["page"],
            "page_end": passage.get("page_end"), "kind": passage["kind"], "footnote": passage.get("footnote"),
            "text": passage["text"], "context": passage.get("context"), "formula": bool(passage.get("formula")),
        }
        for passage in index["passages"] if passage["unit"] == unit
    ]
    body = " ".join(passage["text"] for passage in passages if passage["kind"] != "footnote")
    return {
        "regulation_id": regulation_id,
        "regulation": _short_name(document),
        "regulation_file": document.get("filename"),
        "document_title": index.get("document_title"),
        "label": record["label"],
        "title": record.get("title"),
        "path": record.get("path") or [],
        "page_start": record.get("page_start"),
        "page_end": record.get("page_end"),
        "passages": passages,
        "references": rt.references(body, record["label"]),
    }


# --------------------------------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------------------------------

_EXPORT_TEXT: Dict[str, Dict[str, Any]] = {
    "en": {
        "sheets": ("Summary", "Traceability matrix", "Coverage grid", "Rejected candidates", "Audit log"),
        "title": "Regulation–Release Note Traceability",
        "run": "Run {id}",
        "summary": {
            "status": "Status", "started": "Started (UTC)", "finished": "Finished (UTC)", "duration": "Duration (s)",
            "model": "Model", "calls": "LLM calls", "language": "Language of the texts", "run_language": "Run language",
            "files": "Release-note files", "regulations": "Regulations", "notes": "Release notes analysed",
            "linked": "Linked", "review": "Needs review", "no_link": "No link found", "failed": "Failed",
            "coverage": "Coverage (%)", "links": "Links (direct / indirect)", "provisions": "Distinct provisions",
            "quotes": "Quotes verified", "quotes_value": "{verified} of {total}", "confidence": "Average confidence",
            "candidates": "Candidate provisions reviewed",
        },
        "columns": (
            "Jira ID", "Release-note file", "Sheet", "Status", "Regulation", "Provision", "Title", "Location", "Page", "Link",
            "Approach", "Confidence", "Affected element", "Rationale", "Regulation quote (verified)", "Release-note quote", "Checks",
        ),
        "no_link": "no link",
        "grid": "Jira ID \\ Provision",
        "rejected": ("Jira ID", "Regulation", "Provision", "Title", "Page", "Decided by", "Reason"),
        "log": ("Seconds", "Level", "Release note", "Message"),
        "languages": {"en": "English", "de": "German", "es": "Spanish"},
        "run_status": {"running": "Running", "completed": "Completed", "failed": "Failed", "cancelled": "Cancelled", "interrupted": "Interrupted"},
        "note_status": {"linked": "Linked", "review": "Needs review", "no_link": "No link", "failed": "Failed"},
        "link_type": {"direct": "Direct", "indirect": "Indirect"},
        "approach": {"standardised": "Standardised approach", "irb": "IRB approach", "both": "SA and IRB", "not_specific": "Not approach-specific"},
        "source": {"model": "Model judgement", "verification": "Failed verification", "review": "Four-eyes review"},
        "check_status": {"pass": "pass", "corrected": "corrected", "fail": "fail", "warn": "warning", "info": "info"},
        "checks": {
            "regulation_quote": "Regulation quote verbatim", "note_quote": "Release-note quote verbatim", "scope": "Scope fits",
            "substantive": "Substantive provision", "concept": "Shared regulatory concept", "citation_retargeted": "Citation corrected",
            "review": "Four-eyes review",
        },
    },
    "de": {
        "sheets": ("Zusammenfassung", "Rückverfolgbarkeitsmatrix", "Abdeckungsraster", "Verworfene Kandidaten", "Laufprotokoll"),
        "title": "Rückverfolgbarkeit Regulierung – Release Notes",
        "run": "Lauf {id}",
        "summary": {
            "status": "Status", "started": "Gestartet (UTC)", "finished": "Beendet (UTC)", "duration": "Laufzeit (s)",
            "model": "Modell", "calls": "LLM-Aufrufe", "language": "Sprache der Texte", "run_language": "Sprache des Laufs",
            "files": "Release-Note-Dateien", "regulations": "Regulierungen", "notes": "Analysierte Release Notes",
            "linked": "Verknüpft", "review": "Prüfen", "no_link": "Kein Bezug gefunden", "failed": "Fehlgeschlagen",
            "coverage": "Abdeckung (%)", "links": "Bezüge (direkt / indirekt)", "provisions": "Unterschiedliche Vorschriften",
            "quotes": "Verifizierte Zitate", "quotes_value": "{verified} von {total}", "confidence": "Durchschnittliche Konfidenz",
            "candidates": "Geprüfte Kandidatenvorschriften",
        },
        "columns": (
            "Jira-ID", "Release-Note-Datei", "Blatt", "Status", "Regulierung", "Vorschrift", "Titel", "Fundstelle", "Seite", "Bezug",
            "Ansatz", "Konfidenz", "Betroffenes Element", "Begründung", "Regulierungszitat (verifiziert)", "Release-Note-Zitat", "Prüfungen",
        ),
        "no_link": "kein Bezug",
        "grid": "Jira-ID \\ Vorschrift",
        "rejected": ("Jira-ID", "Regulierung", "Vorschrift", "Titel", "Seite", "Entschieden durch", "Begründung"),
        "log": ("Sekunden", "Stufe", "Release Note", "Meldung"),
        "languages": {"en": "Englisch", "de": "Deutsch", "es": "Spanisch"},
        "run_status": {"running": "Läuft", "completed": "Abgeschlossen", "failed": "Fehlgeschlagen", "cancelled": "Abgebrochen", "interrupted": "Unterbrochen"},
        "note_status": {"linked": "Verknüpft", "review": "Prüfen", "no_link": "Kein Bezug", "failed": "Fehlgeschlagen"},
        "link_type": {"direct": "Direkt", "indirect": "Indirekt"},
        "approach": {"standardised": "Standardansatz", "irb": "IRB-Ansatz", "both": "KSA und IRB", "not_specific": "Ansatzunabhängig"},
        "source": {"model": "Modellurteil", "verification": "Prüfung nicht bestanden", "review": "Vier-Augen-Prüfung"},
        "check_status": {"pass": "bestanden", "corrected": "korrigiert", "fail": "nicht bestanden", "warn": "Warnung", "info": "Info"},
        "checks": {
            "regulation_quote": "Regulierungszitat wörtlich", "note_quote": "Release-Note-Zitat wörtlich", "scope": "Anwendungsbereich passt",
            "substantive": "Materielle Vorschrift", "concept": "Gemeinsamer regulatorischer Begriff", "citation_retargeted": "Fundstelle korrigiert",
            "review": "Vier-Augen-Prüfung",
        },
    },
}


@router.get("/runs/{run_id}/export")
def export_run(run_id: str, language: str = "en"):
    """Excel audit workbook in English or German: labels, the AI-written texts (translated where the run has a
    translation) and the log."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    language = language if language in _EXPORT_TEXT else "en"
    text = _EXPORT_TEXT[language]
    state = _JOBS[run_id].snapshot() if run_id in _JOBS else _read_run(run_id)
    if state.get("status") == "running":
        raise HTTPException(409, "Wait until the run has finished.")
    run_language = state.get("language") or "en"
    notes = [_localized_note(note, language, run_language) for note in state.get("notes", [])]
    workbook = Workbook()
    dark = PatternFill("solid", fgColor="111827")
    accent = PatternFill("solid", fgColor="F5C400")
    direct_fill = PatternFill("solid", fgColor="D1FAE5")
    indirect_fill = PatternFill("solid", fgColor="DBEAFE")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(sheet: Any, row: int, values: List[str], widths: List[int]) -> None:
        for column, (value, width) in enumerate(zip(values, widths), start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = dark
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = border
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = sheet.cell(row=row + 1, column=1)

    def label(group: str, value: Any) -> Any:
        return text[group].get(value, value) if value is not None else None

    summary = state.get("summary") or {}
    names = text["summary"]
    sheet = workbook.active
    sheet.title = text["sheets"][0]
    sheet["A1"] = text["title"]
    sheet["A1"].font = Font(bold=True, size=16)
    sheet["A2"] = text["run"].format(id=state["id"])
    sheet["A2"].font = Font(color="6B7280")
    rows = [
        (names["status"], label("run_status", state.get("status"))),
        (names["started"], state.get("created_at")),
        (names["finished"], state.get("finished_at")),
        (names["duration"], round((state.get("duration_ms") or 0) / 1000, 1)),
        (names["model"], state.get("model")),
        (names["calls"], (state.get("usage") or {}).get("llm_calls")),
        (names["language"], text["languages"].get(language, language)),
        (names["run_language"], text["languages"].get(run_language, run_language)),
        (names["files"], ", ".join(item["filename"] for item in state.get("release_note_files", []))),
        (names["regulations"], ", ".join(f"{item['short']} ({item['filename']})" for item in state.get("regulations", []))),
        (names["notes"], summary.get("notes")),
        (names["linked"], summary.get("linked")),
        (names["review"], summary.get("review")),
        (names["no_link"], summary.get("no_link")),
        (names["failed"], summary.get("failed")),
        (names["coverage"], summary.get("coverage")),
        (names["links"], f"{summary.get('links')} ({summary.get('direct')} / {summary.get('indirect')})"),
        (names["provisions"], summary.get("provisions")),
        (names["quotes"], names["quotes_value"].format(verified=summary.get("quotes_verified"), total=summary.get("quotes_total"))),
        (names["confidence"], summary.get("average_confidence")),
        (names["candidates"], summary.get("candidates_reviewed")),
    ]
    for offset, (name, value) in enumerate(rows, start=4):
        sheet.cell(row=offset, column=1, value=name).font = Font(bold=True)
        sheet.cell(row=offset, column=2, value=value)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 90
    sheet.row_dimensions[3].height = 4
    sheet["A3"].fill = accent
    sheet["B3"].fill = accent

    links_sheet = workbook.create_sheet(text["sheets"][1])
    widths = [12, 24, 16, 12, 12, 18, 30, 40, 8, 10, 14, 11, 34, 70, 70, 50, 44]
    header(links_sheet, 1, list(text["columns"]), widths)
    row = 2
    for note in notes:
        links = note.get("links") or []
        if not links:
            values = [note.get("jira_id"), note.get("file"), note.get("sheet"), label("note_status", note.get("status")), "", "", "", "", "",
                      text["no_link"], "", "", "", note.get("no_link_reason") or note.get("error") or "", "", "", ""]
            for column, value in enumerate(values, start=1):
                cell = links_sheet.cell(row=row, column=column, value=value)
                cell.alignment = wrap
                cell.border = border
            row += 1
            continue
        for link in links:
            checks = "; ".join(
                f"{text['checks'][check['id']]}: {label('check_status', check['status'])}" for check in link["checks"] if check["id"] in text["checks"]
            )
            values = [
                note.get("jira_id"), note.get("file"), note.get("sheet"), label("note_status", note.get("status")), link["regulation"], link["reference"],
                link["title"], " › ".join(link["path"]), link["page_label"], label("link_type", link["link_type"]), label("approach", link["approach"]),
                link["confidence"], link["affected_element"], link["rationale"], link["regulation_quote"], link["note_quote"], checks,
            ]
            for column, value in enumerate(values, start=1):
                cell = links_sheet.cell(row=row, column=column, value=value)
                cell.alignment = wrap
                cell.border = border
            links_sheet.cell(row=row, column=10).fill = direct_fill if link["link_type"] == "direct" else indirect_fill
            row += 1
    links_sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{max(1, row - 1)}"

    grid = workbook.create_sheet(text["sheets"][2])
    provisions = state.get("provisions") or []
    header(grid, 1, [text["grid"]] + [f"{item['regulation']} {item['reference']}" for item in provisions], [18] + [14] * len(provisions))
    for offset, note in enumerate(notes, start=2):
        grid.cell(row=offset, column=1, value=note.get("jira_id") or f"{note.get('sheet')} #{note.get('ordinal')}").font = Font(bold=True)
        by_key = {link.get("provision_key"): link for link in note.get("links") or []}
        for column, provision in enumerate(provisions, start=2):
            link = by_key.get(provision["key"])
            cell = grid.cell(row=offset, column=column, value=(f"{'D' if link['link_type'] == 'direct' else 'I'} · {link['confidence']}" if link else ""))
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
            if link:
                cell.fill = direct_fill if link["link_type"] == "direct" else indirect_fill

    rejected_sheet = workbook.create_sheet(text["sheets"][3])
    header(rejected_sheet, 1, list(text["rejected"]), [12, 12, 18, 34, 8, 18, 90])
    row = 2
    for note in notes:
        for item in note.get("rejected") or []:
            values = [note.get("jira_id"), item["regulation"], item["reference"], item["title"], item["page_label"], label("source", item["source"]), item["reason"]]
            for column, value in enumerate(values, start=1):
                cell = rejected_sheet.cell(row=row, column=column, value=value)
                cell.alignment = wrap
                cell.border = border
            row += 1

    log_sheet = workbook.create_sheet(text["sheets"][4])
    header(log_sheet, 1, list(text["log"]), [10, 10, 16, 120])
    for offset, entry in enumerate(state.get("log", []), start=2):
        message = (entry.get("i18n") or {}).get(language) or entry.get("message")
        for column, value in enumerate([entry.get("t"), entry.get("level"), entry.get("note"), message], start=1):
            log_sheet.cell(row=offset, column=column, value=value).alignment = wrap

    buffer = BytesIO()
    workbook.save(buffer)
    stamp = (state.get("created_at") or "")[:10]
    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="regulation-traceability-{stamp}-{run_id[:8]}-{language}.xlsx"'},
    )
